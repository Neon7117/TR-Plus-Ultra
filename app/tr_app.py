# -*- coding: utf-8 -*-
# TR_APP_MARKER — ห้ามลบบรรทัดนี้ ตัวเปิดใช้เช็กว่าโหลดไฟล์ถูกตัว
"""
TR Plus Ultra — โปรแกรมหลัก
===========================
V0.9.0 : แท็บใหม่ WR Master — สร้าง Item Code จากชีทให้อัตโนมัติ

ไฟล์นี้อยู่บน GitHub ตัวเปิด (.exe) จะโหลดมารันทุกครั้ง
แก้ไฟล์นี้แล้ว push = ทุกคนได้ของใหม่ทันที ไม่ต้อง build .exe ใหม่
"""
import os
import re
import csv
import shutil
import json
import asyncio
import queue
import threading
import traceback
from datetime import datetime

import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog

from playwright.async_api import async_playwright

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    XLSX_OK = True
except Exception:
    XLSX_OK = False


# ============================================================================
#  [1] CONFIG — ทุกอย่างที่ผูกกับหน้าตาเว็บอยู่ตรงนี้ก้อนเดียว
#      เว็บเปลี่ยนเมื่อไหร่ แก้แค่บล็อกนี้
#      (ค่าทั้งหมดยืนยันกับเว็บจริงแล้ว 2026-09-07)
# ============================================================================
BASE = 'https://aztek-tools-v2.thehof.gg'
ITEM_LIST_URL = BASE + '/hof/talesrunner/shop/items'

SEL = {
    # หน้า list
    'search_input': 'input[placeholder*="ค้นหา Item"]',
    'search_button': 'button:has-text("ค้นหา")',
    'filter_option': '#filter-option',      # ItemOption
    'filter_duration': '#filter-duration',  # DurationIndex
    'row': 'table tbody tr',
    'page_input': 'input[type="number"]',
    'next_text': 'Next',
    # หน้ารายละเอียด
    'd_name': '#item-name',
    'd_kind': '#item-game-item-id',
    'd_price': '#item-option',
    'd_duration': '#duration-index',
    'd_qty': '#display-quantity',
    'd_trade_label': 'แลกเปลี่ยน',
}

# ---- หน้าสร้างไอเทม ----
ITEM_CREATE_URL = BASE + '/hof/talesrunner/shop/items/create'
SEL_CREATE = {
    # ช่องกรอก (id ของเว็บ v2)
    'name':     '#item-name',            # ชื่อไอเทม
    'kind':     '#item-game-item-id',    # game_item_id (ItemKind / fdItemNum)
    'desc':     '#item-detail',          # คำอธิบาย
    'price':    '#item-option',          # Price
    'duration': '#duration-index',       # ระยะเวลาไอเทม (วัน)
    'mail':     '#mail-title',           # หัวข้อจดหมาย
    'qty':      '#display-quantity',     # จำนวน (ส่วนแสดงบนเว็บ)
    # ตัวที่ไม่มี id แน่นอน -> หาโดยอ้าง "ข้อความที่อยู่ข้างๆ"
    'type_label':  'ประเภทไอเทม',
    'type_ph':     'เลือกประเภท',
    'web_label':   'เปิดใช้งานการแสดงผลบนเว็บ',
    'trade_label': 'แลกเปลี่ยนได้',
    'submit':      'สร้าง Item',
    'img_label':   'รูปภาพไอเทม',        # กล่องอัปโหลดรูปทางขวาของฟอร์ม
}
# กติกาไฟล์รูปตามที่เว็บเขียนไว้ใต้หัวข้อ "รูปภาพไอเทม": .png/.jpg/.webp ไม่เกิน 5 MB
IMG_EXTS = ('.png', '.jpg', '.jpeg', '.webp')
IMG_MAX_MB = 5
DEFAULT_TYPE = 'GENERAL'
DEFAULT_SUFFIX = '(Gift)'

# ---- หน้าสร้างบันเดิล ----
BUNDLE_CREATE_URL = BASE + '/hof/talesrunner/shop/bundles/create'
TIERS = ['General', 'A', 'S', 'SS', 'SS+', 'SSS']
DEFAULT_TIER = 'General'
BUNDLE_TYPES = ['FIXED', 'CHOICE', 'RANDOM', 'GACHAPON', 'GACHAPON_LIMIT']
DEFAULT_BUNDLE_TYPE = 'FIXED'
BUNDLE_MAX_PAGES = 40      # ID สั้นๆ เจอผลหลายหน้า ไล่หาได้ถึงหน้านี้

# famepoint ต้องเป็น "เครดิตเรียลไทม์" ไม่ใช่ "เครดิต" ธรรมดา (เว็บตั้งค่าเริ่มต้นเป็นแบบธรรมดา)
FAME_CREDIT_TYPE = 'WALLET_REALTIME_CREDIT'
FAME_CREDIT_LABEL = 'เครดิตเรียลไทม์ (WALLET_REALTIME_CREDIT)'
FAME_OPTION = 'hof-fame-point'          # ตัวเลือกในดรอปดาวน์แท็บ Credit
FAME_NAME = 'Fame Point'
EXP_NAME = 'Player Experience'

SEL_BUNDLE = {
    'name':        'ชื่อ Bundle',
    'type':        'ประเภท Bundle',
    'deliver':     'ส่งทันที',
    'tab_item':    'Item',
    'tab_credit':  'Credit',
    'tab_debit':   'Debit',
    'tab_mileage': 'Mileage',
    'tab_exp':     'Player Exp.',
    'search_ph':   'ค้นหา',
    'add_btn':     'เพิ่ม',
    'qty_label':   'จำนวน (Quantity)',
    'tier_label':  'Tier',
    'kind_label':  'ประเภท',
    'submit':      'สร้าง Bundle',
}
# ป๊อปอัปยืนยันหลังกดสร้าง — คำที่ถือว่า "ยืนยัน" กับคำที่ห้ามกด
BUNDLE_YES = ['ยืนยัน', 'ตกลง', 'ยืนยันการสร้าง', 'Confirm', 'OK', 'Yes', 'ใช่']
BUNDLE_NO = ['ยกเลิก', 'ปิด', 'Cancel', 'Close', 'ไม่', 'No']
# ป๊อปอัปยืนยันหน้าตาเดียวกันนี้เด้งทั้งตอนสร้าง Item และสร้าง Bundle — ใช้ชุดคำเดียวกัน
CONFIRM_YES = BUNDLE_YES
CONFIRM_NO = BUNDLE_NO
# ---- หน้าสร้าง Item Code (แถบ WR Master) ----
ITEMCODE_CREATE_URL = BASE + '/hof/talesrunner/itemcodes/create'
WR_SPARE = 2               # ชีทบอก 550 -> ใส่ 552 เผื่อไว้เทสเอง
SEL_CODE = {
    # ฝั่งซ้าย
    'name_th':   'ชื่อ Item Code (ไทย)',
    'name_en':   'ชื่อ Item Code (อังกฤษ)',
    'kind':      'ประเภท',
    'kind_val':  'ALL',
    'per_user':  'จำนวนการใช้งานต่อ 1 User',
    'start':     'เวลาเริ่มใช้งาน',
    'end':       'เวลาสิ้นสุด',
    'limit_sw':  'จำกัดจำนวน',
    'limit_max': 'จำนวนครั้งที่สามารถใช้งานได้',
    'limit_left': 'จำนวนคงเหลือ',
    # ฝั่งขวา (ของรางวัล)
    'reward':    'ของรางวัล',
    'rw_th':     'ชื่อรางวัล (ไทย)',
    'rw_en':     'ชื่อรางวัล (อังกฤษ)',
    'rw_sw':     'จำกัดจำนวน Code',
    'rw_max':    'จำนวนซอง',
    'rw_left':   'จำนวนคงเหลือ',
    'code_type': 'ประเภทของ Code',
    'code_fix':  'Fix Codes',
    'code_list': 'รายการ Code',
    'bundle':    'Bundle',
    'bundle_btn': 'เลือก bundle',
    'submit':    'สร้าง Item Code',
}

# ลำดับคอลัมน์ในตาราง: Aztek Item Id | ชื่อ | ประเภท | ItemKind | Actions
COL = {'id': 0, 'name': 1, 'type': 2, 'kind': 3}

TMPL_HEADERS = ['ItemKind', 'ItemName', 'duration', 'trade', 'qty']

# ---- ชื่อคอลัมน์ที่ยอมรับได้ (ชีทแต่ละใบใช้ชื่อไม่เหมือนกัน) ----
# เรียงตามลำดับความสำคัญ ตัวแรกที่เจอจะถูกเลือกเป็นค่าเริ่มต้น
ALIAS = {
    'kind': ['itemkind', 'fditemnum', 'item kind', 'fditemid', 'itemid', 'fditemkind'],
    'name': ['itemname', 'display name', 'displayname', 'ชื่อไอเทม', 'ชื่อ', 'item name'],
    'dur':  ['duration', 'ระยะเวลาของขวัญ', 'ระยะเวลา', 'durationindex', 'วัน'],
    'trade': ['trade', 'แลกเปลี่ยน', 'แลกเปลี่ยนได้', 'tradeable'],
    'qty':  ['qty', 'quantity', 'จำนวน'],
}
# แถวที่ถือว่าเป็น "หัวตาราง" ต้องมีคำใดคำหนึ่งนี้
HEADER_MARKERS = ('itemkind', 'fditemnum', 'fditemkind')
SCAN_ROWS = 400        # ลึกสุดที่จะไล่หาหัวตารางในแต่ละชีท

APP_VERSION = globals().get('TRPU_VERSION') or 'dev'
DATA_DIR = globals().get('TRPU_DATA_DIR') or os.path.join(
    os.environ.get('LOCALAPPDATA') or os.path.expanduser('~'), 'TRPlusUltra')
os.makedirs(DATA_DIR, exist_ok=True)
CHROME_PROFILE = os.path.join(DATA_DIR, 'chrome_profile')
PREF_FILE = os.path.join(DATA_DIR, 'prefs.json')

# ---- บันทึกการใช้งาน ----
# เก็บไว้ในโฟลเดอร์ข้อมูลโปรแกรมเป็นค่าเริ่มต้น แต่ย้ายที่ได้จากในโปรแกรม
LOCAL_LOG_DIR = DATA_DIR
LOG_MAX_EVENTS = 20000          # เกินนี้จะตัดของเก่าทิ้ง

# โฟลเดอร์กลางของทีม — เว้นว่าง = ปิดอยู่ (เก็บเฉพาะในเครื่อง)
# พร้อมเมื่อไหร่ ใส่ path ที่ทุกคนเขียนได้ เช่น
#   CENTRAL_LOG_DIR = r'C:\Users\B_A_N\OneDrive\TR-Logs'
# แล้วโปรแกรมจะเขียน log ซ้ำอีกชุดเป็นไฟล์ของแต่ละเครื่องลงไปให้เอง
CENTRAL_LOG_DIR = ''

TR_USER = os.environ.get('USERNAME') or os.environ.get('USER') or 'unknown'
TR_MACHINE = os.environ.get('COMPUTERNAME') or 'unknown'

# ---- แจ้งบั๊ก ----
# กดปุ่มแล้วเปิดชีทกลางให้เลย ไปกรอกในชีทเอง
BUG_SHEET_URL = ('https://docs.google.com/spreadsheets/d/'
                 '1X2wLu4CQCMVkX9ewOcimtqudjbaSRB6cVOTN12AKHkw/edit?usp=sharing')

C = {
    'bg': '#11161f', 'card': '#151c28', 'line': '#263041', 'input': '#0c1119',
    'fg': '#dfe6f2', 'dim': '#8b98ad', 'accent': '#4f6bed',
    'ok': '#5ed09a', 'warn': '#e3b341', 'err': '#f0736a',
}
FM = ('Segoe UI', 10)
FB = ('Segoe UI', 10, 'bold')


# ============================================================================
#  [2] helper
# ============================================================================
# ---- ไอคอนโปรแกรม ----
# ฝังไว้ในไฟล์นี้เลย จะได้ติดไปกับการอัปเดตจาก GitHub ไม่ต้องแจกไฟล์รูปเพิ่ม
# (ไอคอนของตัว .exe เป็นคนละส่วน อยู่ในขั้นตอน build)
APP_ICON_B64 = (
    'iVBORw0KGgoAAAANSUhEUgAAAEAAAABACAYAAACqaXHeAAAYw0lEQVR42u1be1iVZbb/rfe77AtsNldRFEEQEAWFTE1D0fBC'
    'VpYWdtFyzEbHbs5Y51SaMTRZnjrHHJssS8tGpymZLk7WmGhKapmigimpqAhoilw2t82+fd+7zh9IOaV1OnN6cprzPs/38Dzs'
    'zeZdv3et31rr964N/P/6//VjLvqXNl4IgbVr1yqX8ibFj/GheXl5CgCeNWvWzU1NTVczs5qfny/+VU5eAKD58+cnJiUlnX3g'
    'gQdeys/Pv+PHBPySc31d15GTk7MuJSWFg4KCDCHEbedeU37WIZCXl6cQEScmJuYMGDDguqysLMPtditRUVHyX4IDCgsLQURo'
    'bW2dOnnyZFIURRIRTNP8+ZNgfn6+ICJz/PjxPTMyMiZlZmbC7XYLZgYA/tkDUFBQIACgrKxs1NChQ0N0XTctFgsBwLXXXjtQ'
    'CPGzDwEWQsDhcIy47LLLGAA7nU4AQHBwsPNn7wHMLJkZ0dHRfXv16kUAqEuXLgCA+vr6S5YE1P+DUpfPxb/My8vr7nQ6k2w2'
    'GwCIxMRECQA1NTXaz9UDGACtX79eYWYqLCwMdjgcDpvNhkAggPj4eAghUFZWpkkpf34AbNmyJVTXdd67d29AURQGIIKDg0lR'
    'FPj9fiQmJsLpdCIkJCSGmZVLMRv8rwDorOtXrlx5Rb9+/d5LTk6+zjRN7eGHH+4eHR2tmqYp/X4/hYeHU0xMDBRFSQJgASCZ'
    'mf7pASgoKGAAOHr0aLHFYhk0adKkv1533XWfvfDCC3c0NTVJRVHI6/UCAGVlZaGmpiaosLAwBACIfj4dsiKEwKBBgx7atm0b'
    't7S08CuvvMLvvPMO19bW8okTJ5iZ+aOPPpJWq5WnT58++rxO8Z9/dYZBTk5O8vTp091ut9tkZoOZubq6mg8fPsx+v5+Z2ejd'
    'uzenpaUt0TTtZ9cRKswsMjIylm/bto1dLpfR0tLC1dXV3NzczK2trXzixAkzJSWFnU7nmSNHjkRdakrRP3IaBIBVVZWBQGDj'
    'u+++CyklWltb4ff7MW/ePGRkZKBfv37i8OHD7PV6o2fNmpV5qekC4gcaLM77CWZW4uPjH4mMjPx9ZWUlnzlzRgDAvffei+ef'
    'fx7VVSegqSrb7DaOioo6Mnjw4EOdwP1TCRznCxmdCDCzkpmRuf7e++7jyspKbmxs5NraWt752U4GwLE9Yjml/+Wyd1KKTExI'
    'rFq4cOGVqqqe/xGdJTT9SB5BAMQ/KsUpAKASwMxq94lTesAZHY+Bw7vFp/Z7tE9yEvdLTfX3Te1r3nTTTTxv3jyeNGkSh4WF'
    'c3rGIJ6RKXhkV3BkVIwZ2z3m9JicUQvsNuv5JErnpUZx/u9+QA9CAEReXp6SnZ2tntuzuMB7fjDvCIUIS5YsjA5Pz3y0S/fY'
    'PX1DHM2X62rrIJulqbvVKnunD5ChMd1laIiDJ0+ezF6vlz/b+Skn9knjh4Za+bEROj86fx77/T7++JOdHBzRjaN6Jq1bt+nT'
    'aABQFAUTJ07sBiDOYrFcyOvUizx03vP1HxFB13UQEZhZz87OTluwYMHQ73OTC568Jsgccv0N08uLty+8xu/qdnuwgYF2gtNG'
    'ME2gwUvY5DLxqmFHsVBQuOpVtLW2Imv4cLz71l/w+pMP44Z5S/Dov83BhuKdeO13v+aJIWXmiUZTfaI09MDcgmemhAdOGRu3'
    'frKViKzHjh17x+PxLDtx4sQeVVUNZoaUEkT0d8UTEcFisSAQCCAQCKCoqKj76NGj7Tk5OSn79u2LZ+bucXFxGQ0NDTEAxI03'
    '3nhLSkqKX9f1phkzZtR9k4Pom65CRKouREBJ6nNfQs2JpUuCvBjdXRhwkoDNJFMhEBOElwkeQDYCT50Eng3pgmeeegqbi4qQ'
    'npYGNcgJw90AaMHY/dKjePpyNxJiFEAVxroDAfWJ48kn29SgimWLnx41Kmc0iouLsWXLFhw8eHCfaZp1e/fu3VdVVdV4zhPO'
    'J83g4cOHZ1ZWVtpdLpc1PT09UVVVZ7du3fSUlBSkpaXhwIEDWLRoEeLj4w+5XK7aK664on3ZsmV3xcXFnWZmEBFfqB0mIiIC'
    'Ail5t+TF/G390tecHhkdKxAIlyqFASKIAZ0BCZhewGwBFAthvlWF/cBpvLJ6NSZfdw36D8hE9shs7Ni+DR9+vBu/6OlBQk87'
    'PA1+QJrq9fGK7GU/3mPCe2aPQ8erOWu4D9nZ2TI7O1sBkHnq1CmcPXt2bF1dHerr69Hc3PyVF4SEhMDhcKBLly4IDQ2Fw+FA'
    'aGgo7HY7A5AvvPACli5dyoZhyOPHj/dJTU01c3JyxsTFxZ0+ZyNfKATOETsrqbm5uQm7Sl5eQw1dw3opHIiWQu0CIIoBBwE6'
    'AGaQm8BNAJ8lBM4yLKd13HjYB2PKFMQ5nbA5HIiKCEPpF8dgPVCESbYTGB0L6DYN7JegEI3/vNvL7ybdIV74/WI0Nrpgt9s5'
    'KCiInU4nf0+qJADwer3wer1CCIGKigoqKCjAe++9JwEIi8WC2NjY1yoqKn5NRE35+fmioKBAXkgQIQDEzProq65a4a05OXVe'
    'axPC4jX2hUihRzDQhYBoAoUBbANIAmghQGeQZAgfYLpN/JtTwbRdJai06hg7dizuvPOXuOeeu7Hww/0oKtqEh557CHlaBcKC'
    'NKQKUH876K3WZpAQ0DUdHo+HPB4PNTQ0QAgBIsL5WqKUEqZpgpmhaRpsNhvq6uqwfv16vPjii1xZWSmJSHE6nb5hw4Y9WVRU'
    '9DgR4WLGdwIghBBmcnLyg3dOmzb16vAI48A9s5WhDkEUBMBBICfAkQB3BcxgBvsJsBIUSRAGoLk7AMl0MkJOVsOfnoYZ06fj'
    '3XXvor6hAQ/cOwsrV65ATO9kvHTLUAyxt+PBUhXBBkEfboEpGaqmwhHigMfjgc/nAzOjkwg7M4au69B1HVJKVFZW4sMPP8SH'
    'H36II0eOoLm5mYKCgpSYmJg948aNm7Ns2bIdUkrBHcQmv0sSMx9//PFuq1evnnPrrbfKk61u8XyIlWYJPxQdIAuDgwTgYFAU'
    'Q41TAIOAagbcEqeOMgJeQBpArCYQ7vejrN4FAqOluQU35eUhxOHAqldfw6zZs7GyRypmKLvQxUb4xU7GfSl9oBKh2eOBEAJh'
    'YWEAAI/HA6/XC8MwIKWEy+VCVVUVduzYga1bt2Lf3n2QzLDb7ez1emGa5tlu3bq/XFFRsaiiosIdERHhuO222/xE5P9OTZCI'
    'sHbt2tGDBg2K6tajh9lD05RFA6/EvsrNyEjUYEKeqyoIbj9wtMxEaSlh62YFFcdCEBYVg6M1tbA0erHcaUI1AxjQNQJdY2Ph'
    '83mwd28Jxowdh+rqatz7wL+j/dgR6P003Oj0409xYbj6hhtgGH4Q0Ven73Q6vwKivb0dPp8PRUVFePLJp5CUlIyxY8chP38B'
    'EhLioWoKSAgKsgX5goKcvfbu3f/Qhg1FuxYtKtj33HPPNUZHR9tHjBgRKCwsDFyIV1RFUbB///7Q3Nxc1nWdAeD2uQ/iP27Z'
    'jDcCBOkDyMMQAQt2bfOiTM5HY2s73PJZHG8SeOWJAaiXw/H6PfMQIn3QVAXxbafx/AvL0Xy2Du+sLYSq6Ridk4OSvftwTFpQ'
    'fKwJOxqBxLt/g+EDMxEI+AFS0NbWBiklGhoa4HK5oGkaOk5X4voJN+COabfCogd/tXlvG+D3gUgADid6QmBKdvaVyM6+wpg7'
    '9549u3eX/vmaa8a+W1hYWD9mzBitqKjIA0B+o3giZGRk5ISEhGzavHmzBCAURcFt0+7EuC2vYtp4HYGQANBNgdpForgxC02p'
    'z0DWvoiAuxrhCVOx/oNDuPOd/0AvvxXT20xk2wJwPPsGyj7ZjgOHDmHNH1fjiYVPQJomJt06BYsXLECfjDTMfWQ+bJpARESH'
    'fN7Q0ADDMOD1elFXVweXy4W4uF5ITu4FQEHDGWDnRz4cKg3g7CnA4wZMAyACrHZwWJSQ8ckCl4+wKCkZHQmurq6hZM2aPz0x'
    'd+6cjyZMmIC//vWv7QDMv2sYmJmdTmfRqlWrciZOnGgaRkBpamvHxNGjsUCWYOyVFvgsAagRAk1eA2+c6Iv2pGmwhndBbcUB'
    '9PxgCWYKgT9Xq/ib4cHQKTdjxitvIMYZiueefw4lZWUYdPkgnKypwac7tmP162+gqbkJbBgwTANRUdGw221oaGjAqlWr0L9/'
    'f9hsNvTt2w/h4WGo+5Lxtzf8KCk20NYCKCqgaQShdBgPBqQEDAMwAgxFARJSSebeonPGMFUxzIB/x/ZdT44cmfWH3NzcwIYN'
    'G9ydIFBeXp5SWFhoTpgwYdSZM2c2b9q0iW02G6mqQpWnzmDqteNxn68Utwy1AboBqTFgGDhcDbiagehWIBEq+CRjiVtDqsOP'
    'PROmwZrQG7vnz0efhYsw49Y8mKaE1zSR2qcPXlmxEsNHDMfcBx7AM888g/DwcISFhUFVVYwcORJudzu2bfsYVqsVxev9eHuF'
    'gbYWhj2IoKgA89fP167ccZydgHjbGaYJDL5KMe/4jVWx2oFdO0uWDBk66He5ubnGhg0b2gBIpby8nAEoNTU1x4UQ3Y8ePXr5'
    '9ddfL71en4gKD0PupIlYUPQ5tm85hIFCwhlQQa0KokwNsQFCeLsK72kT7CJc5mCkRhAiy/YiesdHuD1Cw5O7SqEkpcHDjAUP'
    'PoCWkzUoPXoUh774Qr711lucnp7Offv2Jb/fj0AggNDQMMyefTe6do3GG8u8eHuFAVUj2Ox0QcM7jGcwE3De65pO0K2E4+VS'
    'HCgxOPUyIVP69hg2btx45ZFHHt6WnZ2tVlVVBZRzrSlt2rQJjY2NZTNnzpzicrmCx40bx4ZhkDMoCOOvzsXnli54esshHC53'
    'IbjeRGSDAa0eQANDbVWgaAo0J8PvN9FVE+hhEbCDMMbXiu2Fb6N89WvIq6vCZ0w44mrGgdJS0jSNiIhycnLg8XjgcrkwZPAV'
    '6NYtGmtf9OGD102ERtC5IugiJSExDEODoph/19p0AmELItSfYSovMenybEX2TuqRFReXULF06ZKD06ZNEwoAFBcXc15enjJ5'
    '8uSmqVOnHlmxYsWtdXV1cty4cSSEQi3NzRg+OBNX3zYNFd0SsbrRj7frfNhZ145St4nDUqLaMGE3JcJVFdIvEQhISIMQJgRG'
    '2VWMD9awn1Ve7jaoe2Rk/cSbblpTU1MTRETb0tPTU4UQHOoMo8ioCHz8vh+FLxkIjSSYxnfJ0iba2oMx9PIiSKmiqTkSqmr8'
    'PRASsNoIjWeZTh6TGJKjUXJS0pAdO7avP3z4cOtXSk9nKOzfv//QmDFjzNdeey2nuLhYDhgwAEnJyeTx+rC3ZBeGZWZi5t2z'
    '0XvMOIRcfwOMsWPROnoMTgy9AmvcJqq+qEQWAEXoUCRgmAzhJ3zqA6a2s/RIidi4noc+++yzCQcOHHh19erVR2w2211paelI'
    'TEwgVx3hxcf9UDX6Hk3eRJvbgf59d2PK5CXY+NHNME0NQshvdfksAaudUF3BFOyEmZphDenfP5Oeemrh5m9p9KZpKocPHy4e'
    'P368uXXr1pzly5dTbW2tUd/QgPvvn0PDs4fjsQWPISk2FjeNvxopPXtiYFo/DBs0GCOuvwF/021YW34IgYZWWL0S5Afelwqm'
    'BAiNJstgu1W5Mmt4/r59+/YuXrw4kJOT015RUTFpct7NkeER4fIvL/uoYr+EzU4XdXshJHx+KyLCa/GbOfejrSUCRVsmQ1UD'
    'uOjFEwOqSqiqMCnrah0x3cN6bd788SblAhITCgoKRFVVVfGkSZM+r62tHVJcXBz+/vvvk9VqleXl5WZxcTGPGTuWJIMCpgQp'
    'AgyB8JBgZPZNRnFLKyoGZOD9IDuWtXvwe7ebPSwNQVJLHzDgLxs3bswvKCjgmTNnqqtWrWrv3r3nFTNnTk/3tFrkmt/7hKbT'
    't4ju/JiXUoEQJu6fNQ9RvQ7i2BeD8dme0VBV/3cqX4oGNDeAHKHg5P6WoC6R0fXqt/8BMQD2+/3izTfffHv//v077r777nuO'
    'HTt2a1NTU+9du3YJIsKMGTMQFBSEmJgYMzIiAklJSRiRnU3Lly+nXnE96c3Cv2BTcTFunzaNUSfZZrVqKamphbt37556rj6n'
    'gQMHAoB+331zKoOCQrBhfQDuFsARCkjz4hfSPp8Vv7rzt4hP2gf2BKO1zQlT0rls8B36oezIDru2GBg3WeeBgzLGfdd8gASg'
    '9O/fvxbAY8z8n3feeeeIzz//PMvlcg1ta2tLc7lc4RUVFUpFRQU+3bkTf1y9+qtIunv2bCpcu5bqGxthtVhERmbm05988smj'
    'RBTobE+//PJLBYCWnp4aCwa+2BeAol5cCVCEiZY2B/ImvIyBQ9cj0BwJzdkAV1MUCBJE8ivhhFlcQEAFdAtQexLUUMsU3TUy'
    '5fsGJDpzi0JELQDWCyHW67qOwsLC7m+//Xavo0ePDvB4PD1tNltaRUVFrMfjSd29e7e6e/duAEBoaChGjhw5e926dS92CE6g'
    'zt68vLycANgjI0NTQEDtSSZVu7D7K4qBltYQDB+yEVdftwLSHdqR+qTA6do4+PwCms8OKQWkVGCxeCCE/BYnCAVwtzDVHDUR'
    'Ea3bfohUTHl5eaKwsLDTO77W1VQVQgj4fD7LL3/5y97V1dW3lJeXT2xqakpOS0ubW1JS8gfDMFRmNjslqXP6owVA93Zvy0YE'
    'HAkPTXVLaZL45gWyEBJt7mAkxh/C3Dn3QRUm6BwPkGKipiYJ7vYQKMKEaaqwBDfhi4ND8d6Hd0D7BjEKBWh1Mab+WsdVN2g/'
    'aESGCwsLzfMF1N/+9re0detWUVxczACYiHwADhLRAinlwttvvz1lzZo1Zed0euMiV+OqIkj3+AAjACjKt0nPH9AR1+MoHpxz'
    'H/TwWsBnBRQT8NmAgI7YXgcBYQKGBgQ34+QXg7F1+wSIi+ggDMDn/QdnhDrJ8hvtJeXn51NBQQERkRdA2bkQkt/DNVII4GKT'
    'dCwFukbXYMsn18FoD4aiGvD5LchM34GevQ5CeoMgTQVqcBOO7x+OxX9YDMNUoWu+C3JBpyf8XwxJfWuvncMT54EhLwZgdnY2'
    'iouL2z3tvqaQ0OCeVhvQ7gZU9et6n5mg6z7sKc3Gp7tGgURHwxMIAOmpuwFhQsoO48982RsvrnwchqnCovsg5UWMJyA0nH70'
    'W1q+mPGda+TIkQDgdrlaq4mAsCjiC5W+zASLxYOQkFY4Q1pg0f3o1+cAEhL2w2x3QA3qMP7Z555Fc2sYLLr3osYzA5qF0DVW'
    'ADDlT3pN3a9fPwnAW1X15ecAEN9HwAgwSFwYBCkVSEnw+nQkJ5aCdR8Uq/sr45tawmG1eCClcpGwBQw/4Ixg7pEg0NzScuqn'
    'vqc3AWDduve2A9IcMkoXioLvLmaYoGsBpPUpAelenP4y8X9kPACQAHxeRt/LFFZU4NAXldt+UgAmT54sBw4cKJ59dlFpQ4Or'
    'tHeaQEJfYfra+YKESMQIGBZEhJ9FQsoe1NckY8nz//U/Mr6zEtQtQM5EnQDTv3btG+/+1B7Ad9xxhwnAtXlz8Z8A0DVTNBjG'
    'hUt6IgmvV8NlAz5GwFTx3LKn4WqOhM36/cYrKuBuZWRmKWaPBEGnTtXuWLz4mW0/+ajKnDlz/P379xe/+tVdb9XXN5SlD1aV'
    'oWOEbHVxR1l8fr5kAavVix4xx/GH55bg5OlesNvcMM3vNl4IwO8DQsLBN96lwzSkfO2VNUujoqLaLomRtREjRtCePXt8Dkdo'
    '1bBhw27se5kqyveaqD/DZLURWHYCoMBubcfxE31Rfao3guxt33vyQgCmBAI+YOajutmrj6qW7i1dcfNteS8NGjSILwkAysvL'
    '5ZgxY7SVK1+uGjlilJmS2mtU38uEPFhiUuNZJqu9Q+8jdMhf7R5HR6r7nu0rSke9EPADU+eoxpCrdPXLU7W7UlKT7h82bFh7'
    't27dPJfM0OLx48fNCRMm6PMfnbdnfO61enLfHsMuz1a45qjk6mNMqkpQtQ5uEMK8aIVH9HWV19bCCA4h3DVPM4bk6GrtmfrS'
    'W2+b/AtN02pVVfVu2rTJuJSmNvnw4cNmbm6u8vAjD20fn3utSEiKyRo6ViObHUb1MZOaGkAAQSgEITrS2vkP0NFPeNwMImBg'
    'tmLOfsxGvfooyunTZz+9+ZabZnz88cdVgwcP9m/bts1/SQ0snu+5ubm5QRs2bBCvv/7mpGuvvbrA4XD08HuAzev8smSrgTMn'
    'WXjbO9IafzVpBWgWIDQCnJqp8KgJVhGXDACGLCs79HJGRvoiAI25ubm+DRs2+C710TyRnZ0dDCA4Kyurf+m+A8/7fO2n+dw6'
    'e4p573aDN78T4A9e9/OGN/28c1OAKw9JlkbnuwKempovNy5d+vwNAJwRERGOgQMHavgnWpSXl2cD4ADgHJczbugH6zc+efLk'
    'yRKTvW3nvqLzjRUINDU1VO7+rOSP+fkLbwLQDUBwXl5eMC7ypc3/BuDjbyRYiWeLAAAAAElFTkSuQmCC'
)


def set_app_icon(win):
    """ติดไอคอนให้หน้าต่าง — ทำไม่ได้ก็ไม่เป็นไร โปรแกรมต้องไม่พังเพราะเรื่องนี้"""
    try:
        img = tk.PhotoImage(data=APP_ICON_B64)
        win.iconphoto(True, img)
        win._trpu_icon = img          # กัน Python เก็บกวาดรูปทิ้ง ไอคอนจะได้ไม่หาย
        return True
    except Exception:
        return False


def find_chrome_exe():
    cands = [
        os.path.join(os.environ.get('PROGRAMFILES', ''), 'Google/Chrome/Application/chrome.exe'),
        os.path.join(os.environ.get('PROGRAMFILES(X86)', ''), 'Google/Chrome/Application/chrome.exe'),
        os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Google/Chrome/Application/chrome.exe'),
    ]
    for p in cands:
        if p and os.path.exists(p):
            return p
    return None


def launch_kwargs(headless=False):
    kw = {
        'user_data_dir': CHROME_PROFILE,
        'headless': headless,
        'args': ['--disable-blink-features=AutomationControlled', '--start-maximized'],
        'no_viewport': True,
    }
    exe = find_chrome_exe()
    if exe:
        kw['executable_path'] = exe
    return kw


def load_prefs():
    try:
        with open(PREF_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def save_prefs(d):
    try:
        with open(PREF_FILE, 'w', encoding='utf-8') as f:
            json.dump(d, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def _dig_id(obj, depth=0):
    """ขุดหา "เลข Bundle" จากคำตอบของเว็บ — คีย์ชื่ออะไรก็ได้ ขอให้สื่อว่าเป็น id"""
    if obj is None or depth > 6:
        return ''
    if isinstance(obj, dict):
        for k in ('bundleId', 'bundle_id', 'id', 'seq', 'no'):
            v = obj.get(k)
            if isinstance(v, (int, str)) and str(v).strip().isdigit():
                return str(v).strip()
        for v in obj.values():
            got = _dig_id(v, depth + 1)
            if got:
                return got
    elif isinstance(obj, list):
        for v in obj:
            got = _dig_id(v, depth + 1)
            if got:
                return got
    return ''


def norm_name(s):
    return re.sub(r'[\s()（）\[\]]', '', str(s or '')).lower()


def norm_bool(v):
    s = str(v or '').strip().lower()
    if s in ('yes', 'y', '1', 'true'):
        return 'yes'
    if s in ('no', 'n', '0', 'false'):
        return 'no'
    return 'any'


# ============================================================================
#  [2.5] บันทึกการใช้งาน (history)
#        เขียนทีละบรรทัดเป็น JSON ลง history.jsonl — อ่านย้อนหลังได้ทั้งหมด
#        ทุกบรรทัดมีชื่อผู้ใช้/ชื่อเครื่องติดไปด้วย เผื่อวันหน้าเอามารวมกันทั้งทีม
#        ถ้าเขียนไม่ได้ก็เงียบไป ห้ามทำให้โปรแกรมหลักพัง
# ============================================================================
SESSION_ID = datetime.now().strftime('%y%m%d%H%M%S')


def history_file():
    return os.path.join(LOCAL_LOG_DIR, 'history.jsonl')


def stats_file():
    return os.path.join(LOCAL_LOG_DIR, 'stats.json')


def set_local_log_dir(path, move=True):
    """ย้ายที่เก็บ log ในเครื่อง — ย้ายไฟล์เดิมตามไปด้วย จะได้ไม่เสียประวัติ
       คืนค่า (จำนวนไฟล์ที่ย้าย, ข้อความผิดพลาดถ้ามี)"""
    global LOCAL_LOG_DIR
    new_dir = (path or '').strip() or DATA_DIR
    old_dir = LOCAL_LOG_DIR
    moved, err = 0, ''
    try:
        os.makedirs(new_dir, exist_ok=True)
        if move and os.path.abspath(new_dir) != os.path.abspath(old_dir):
            for fn in ('history.jsonl', 'stats.json'):
                src = os.path.join(old_dir, fn)
                dst = os.path.join(new_dir, fn)
                if not os.path.exists(src):
                    continue
                if os.path.exists(dst):
                    if fn.endswith('.jsonl'):     # มีไฟล์อยู่แล้ว → ต่อท้ายไม่ให้ของเก่าหาย
                        with open(src, 'r', encoding='utf-8') as f, \
                                open(dst, 'a', encoding='utf-8') as g:
                            g.write(f.read())
                        moved += 1
                else:
                    shutil.copy2(src, dst)
                    moved += 1
    except Exception as ex:
        err = str(ex)[:150]
        new_dir = old_dir
    LOCAL_LOG_DIR = new_dir
    try:
        pr = load_prefs()
        pr['local_log_dir'] = '' if os.path.abspath(new_dir) == os.path.abspath(DATA_DIR) else new_dir
        save_prefs(pr)
    except Exception:
        pass
    return moved, err


def load_local_log_dir():
    global LOCAL_LOG_DIR
    try:
        saved = (load_prefs().get('local_log_dir') or '').strip()
        if saved and os.path.isdir(saved):
            LOCAL_LOG_DIR = saved
    except Exception:
        pass
    return LOCAL_LOG_DIR


def set_central_dir(path):
    """ตั้ง/ล้างโฟลเดอร์กลางของทีม — เก็บลง prefs ไม่ต้องแก้โค้ด"""
    global CENTRAL_LOG_DIR
    CENTRAL_LOG_DIR = (path or '').strip()
    if CENTRAL_LOG_DIR:
        try:
            os.makedirs(CENTRAL_LOG_DIR, exist_ok=True)
        except Exception:
            pass
    try:
        pr = load_prefs()
        pr['central_log_dir'] = CENTRAL_LOG_DIR
        save_prefs(pr)
    except Exception:
        pass
    return CENTRAL_LOG_DIR


def load_central_dir():
    """อ่านค่าที่เคยตั้งไว้ตอนเปิดโปรแกรม"""
    global CENTRAL_LOG_DIR
    try:
        saved = (load_prefs().get('central_log_dir') or '').strip()
        if saved:
            CENTRAL_LOG_DIR = saved
    except Exception:
        pass
    return CENTRAL_LOG_DIR


def pending_items():
    """สิ่งที่ยังตั้งค่าไม่ครบ — เอาไว้เตือนบนหน้าจอ จะได้ไม่ลืม"""
    todo = []
    if not CENTRAL_LOG_DIR:
        todo.append('ยังไม่ได้ตั้งโฟลเดอร์กลางของทีม — log ตอนนี้เก็บเฉพาะในเครื่องนี้ '
                    'ถ้าอยากรวมสถิติทั้งทีม กดปุ่ม "เลือกโฟลเดอร์" ข้างบน')
    elif not os.path.isdir(CENTRAL_LOG_DIR):
        todo.append(f'โฟลเดอร์กลางที่ตั้งไว้หายไป: {CENTRAL_LOG_DIR}')
    return todo


def _central_path():
    """ไม่สร้างโฟลเดอร์เอง — ถ้าโฟลเดอร์หาย (ไดรฟ์หลุด/ยังไม่ sync)
       ให้ข้ามไปเฉยๆ แล้วให้ pending_items() เตือนแทน
       จะได้ไม่ไปสร้างโฟลเดอร์ว่างทิ้งไว้แล้ว log กองผิดที่โดยไม่รู้ตัว"""
    if not CENTRAL_LOG_DIR or not os.path.isdir(CENTRAL_LOG_DIR):
        return None
    try:
        safe = re.sub(r'[^A-Za-z0-9._@-]', '_', f'{TR_USER}@{TR_MACHINE}')
        return os.path.join(CENTRAL_LOG_DIR, f'log_{safe}.jsonl')
    except Exception:
        return None


def log_event(_ev, **fields):
    """บันทึก 1 เหตุการณ์

    ชื่อพารามิเตอร์แรกขึ้นต้นด้วย _ เพราะฟิลด์ที่ส่งเข้ามามีชื่อ kind ด้วย
    (kind = ItemKind ของไอเทม) ถ้าตั้งชื่อซ้ำจะชนกันแล้วโปรแกรมล้ม
    """
    ev = {
        'ts': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'kind': _ev,
        'user': TR_USER,
        'machine': TR_MACHINE,
        'session': SESSION_ID,
        'app': APP_VERSION,
    }
    ev.update(fields)
    line = json.dumps(ev, ensure_ascii=False)
    for path in (history_file(), _central_path()):
        if not path:
            continue
        try:
            with open(path, 'a', encoding='utf-8') as f:
                f.write(line + '\n')
        except Exception:
            pass
    _bump_stats(_ev, fields)
    return ev


def _bump_stats(kind, fields):
    """ตัวนับรวม เอาไว้โชว์เร็วๆ โดยไม่ต้องอ่านไฟล์ทั้งก้อน"""
    try:
        st = read_stats()
        st['total_events'] = st.get('total_events', 0) + 1
        st['last_used'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        st.setdefault('first_used', st['last_used'])
        st['by_kind'] = st.get('by_kind', {})
        st['by_kind'][kind] = st['by_kind'].get(kind, 0) + 1
        day = datetime.now().strftime('%Y-%m-%d')
        st['by_day'] = st.get('by_day', {})
        st['by_day'][day] = st['by_day'].get(day, 0) + 1
        if kind == 'search_done':
            st['items_found'] = st.get('items_found', 0) + int(fields.get('found') or 0)
        if kind == 'error':
            st['errors'] = st.get('errors', 0) + 1
        with open(stats_file(), 'w', encoding='utf-8') as f:
            json.dump(st, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def read_stats():
    try:
        with open(stats_file(), 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def read_events(limit=None, kinds=None):
    """อ่าน history ย้อนหลัง (ใหม่สุดอยู่ท้ายไฟล์)"""
    out = []
    try:
        with open(history_file(), 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    ev = json.loads(line)
                except Exception:
                    continue
                if kinds and ev.get('kind') not in kinds:
                    continue
                out.append(ev)
    except Exception:
        return []
    return out[-limit:] if limit else out


def trim_history():
    """ตัด log เก่าทิ้งถ้ายาวเกิน LOG_MAX_EVENTS"""
    try:
        hf = history_file()
        if not os.path.exists(hf):
            return
        with open(hf, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        if len(lines) <= LOG_MAX_EVENTS:
            return
        with open(hf, 'w', encoding='utf-8') as f:
            f.writelines(lines[-LOG_MAX_EVENTS:])
    except Exception:
        pass


# ============================================================================
#  [3] Template Excel
# ============================================================================
def build_template(path):
    notes = [
        'ItemKind = รหัสไอเท็ม (ช่องค้นหาหลัก)   |   ItemName = ใส่เพื่อกรองชื่อซ้ำอีกชั้น (ไม่บังคับ)',
        'duration : เว้นว่าง = เฉพาะไอเท็มถาวร (ระยะเวลา = 0)  |  ใส่ตัวเลข = จำนวนวันนั้น (เช่น 15)  |  any = ไม่กรอง',
        'trade : Any / Yes / No        qty : เว้นว่าง = ไม่กรอง หรือใส่ค่าที่ต้องตรงเป๊ะ',
        'กรอกรายการตั้งแต่แถวที่ 5 ลงมา',
    ]
    if XLSX_OK and path.lower().endswith(('.xlsx', '.xlsm')):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Items'
        for r, line in enumerate(notes, 1):
            c = ws.cell(row=r, column=1, value=line)
            c.font = Font(italic=True, color='8B949E', size=9)
            c.fill = PatternFill('solid', fgColor='21262D')
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(TMPL_HEADERS))
        blue = PatternFill('solid', fgColor='1F6FEB')
        orange = PatternFill('solid', fgColor='D29922')
        for col, h in enumerate(TMPL_HEADERS, 1):
            cell = ws.cell(row=5, column=col, value=h)
            search_col = h in ('ItemKind', 'ItemName')
            cell.fill = blue if search_col else orange
            cell.font = Font(color='FFFFFF' if search_col else '0D1117', bold=True)
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[cell.column_letter].width = [14, 30, 12, 10, 10][col - 1]
        ws.append(['40852', '', '', 'No', '1'])
        ws.append(['40852', '', '15', 'Yes', ''])
        ws.append(['40852', '', 'any', 'Any', ''])
        ws.add_data_validation(DataValidation(
            type='list', formula1='"Any,Yes,No"', allow_blank=True, sqref='D6:D1000'))
        wb.save(path)
        return
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(TMPL_HEADERS)
        w.writerow(['40852', '', '', 'No', '1'])


def read_template(path):
    """อ่านไฟล์ CSV แบบ template ง่ายๆ (ไฟล์ Excel ใช้หน้าต่างนำเข้าแทน)"""
    rows = []
    with open(path, 'r', encoding='utf-8-sig', errors='replace') as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames or []
        cmap = {k: guess_column(headers, k) for k in ('kind', 'name', 'dur', 'trade', 'qty')}
        for d in reader:
            cells = [str(d.get(h, '') or '') for h in headers]
            row = cells_to_row(cells, cmap)
            if row:
                rows.append(row)
    return rows


# ============================================================================
#  [3.5] อ่านไฟล์ Excel ต้นฉบับ (ชีทเยอะ หัวตารางอยู่ลึก และมีหลายบล็อกต่อชีท)
# ============================================================================
# อีโมจิ/สัญลักษณ์รูปภาพที่ชีทชอบใส่นำหน้าหัวข้อ (🎨 ✅ 🎁 ⚙️ 💛 ✨ 🖌️ 🐾 ...)
# ไม่รวม ™ © ® และ • · — ที่อาจเป็นส่วนหนึ่งของชื่อจริง
_EMOJI_RE = re.compile(
    '['
    '\U0001F000-\U0001FAFF'      # อีโมจิทั้งก้อน (หน้ายิ้ม ของใช้ ธง ฯลฯ)
    '☀-➿'              # สัญลักษณ์เบ็ดเตล็ด + dingbats (✅ ✨ ⚙ ✔ ❌ ➡)
    '⬀-⯿'              # ดาว/ลูกศรแบบอีโมจิ (⭐ ⬅)
    '⌀-⏿'              # สัญลักษณ์เทคนิค (⏰ ⌛ ⏳)
    '■-◿'              # รูปทรงเรขาคณิต (◼ ▶ ●)
    '︀-️'              # variation selector (ตัวที่ทำให้กลายเป็นอีโมจิสี)
    '‍⃣⃠'         # ZWJ / keycap
    ']+')
_TRIM_CHARS = ' \t\n\r ​•·|'


def clean_text(v):
    """เอาเฉพาะตัวหนังสือจากข้อความในชีท — ตัดอีโมจิ/สัญลักษณ์รูปภาพออกให้หมด

    ชีทของทีมชอบใส่อีโมจินำหน้าหัวข้อ เช่น "🎨 Digital Grand Winner (1 รางวัล)"
    อีโมจิพวกนี้ไม่ควรติดไปถึงเว็บหรือชื่อบันเดิล — เอาแต่ตัวหนังสือพอ
    ใช้กับทุกข้อความที่อ่านมาจากไฟล์ต้นฉบับ (ชื่อบันเดิล / ชื่อไอเทม / ชื่อชีท)
    """
    s = '' if v is None else str(v)
    if not s:
        return ''
    s = _EMOJI_RE.sub(' ', s)
    s = re.sub(r'[ \t ]+', ' ', s)
    s = re.sub(r' *\n *', '\n', s)
    return s.strip(_TRIM_CHARS)


def num_str(v):
    """แปลงค่าจากเซลล์เป็นข้อความ — ตัด .0 ที่ Excel ติดมากับตัวเลข"""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    m = re.fullmatch(r'(\d+)\.0+', s)
    return m.group(1) if m else s


def dur_from_text(v):
    """กติกาเดียวกับเครื่องมือเดิมเป๊ะ

         เว้นว่าง   ->  ''     = เฉพาะไอเท็มถาวร (บนเว็บช่องระยะเวลาไอเท็ม = 0)
         'ถาวร'     ->  ''     = ความหมายเดียวกัน (ชีทต้นฉบับเขียนคำนี้)
         'any'      ->  'any'  = ไม่กรองระยะเวลา
         ตัวเลข     ->  ตัวเลขนั้น เช่น '15'
    """
    s = num_str(v).strip()
    if not s:
        return ''
    if 'ถาวร' in s or 'permanent' in s.lower():
        return ''
    if s.lower() == 'any':
        return 'any'
    m = re.search(r'(\d+)', s)
    return m.group(1) if m else 'any'


def _norm_h(s):
    return re.sub(r'[\s_\-]', '', str(s or '')).strip().lower()


def _is_marker_cell(c):
    """เซลล์นี้เป็น 'ชื่อคอลัมน์' จริงไหม — ต้องสั้นและขึ้นต้นด้วยคำหลัก
       (กันบรรทัดคำอธิบายยาวๆ ที่บังเอิญมีคำว่า ItemKind อยู่ในประโยค)"""
    h = _norm_h(c)
    if not h or len(h) > 24:
        return False
    return any(h == m or h.startswith(m) for m in HEADER_MARKERS)


def is_header_row(cells):
    if sum(1 for c in cells if _norm_h(c)) < 2:      # หัวตารางต้องมีอย่างน้อย 2 ช่อง
        return False
    return any(_is_marker_cell(c) for c in cells)


def guess_column(headers, key):
    """เดาว่าคอลัมน์ไหนคือ key ('kind'/'name'/…) คืน index หรือ -1 ถ้าไม่เจอ"""
    hs = [_norm_h(h) for h in headers]
    for want in ALIAS.get(key, []):
        w = _norm_h(want)
        for i, h in enumerate(hs):
            if h == w:
                return i
    for want in ALIAS.get(key, []):
        w = _norm_h(want)
        for i, h in enumerate(hs):
            if h and (w in h or h in w):
                return i
    return -1


def cells_to_row(cells, cmap):
    """แปลงข้อมูล 1 แถว + การจับคู่คอลัมน์ → เงื่อนไขค้นหา 1 รายการ"""
    def get(k):
        i = cmap.get(k, -1)
        return str(cells[i]).strip() if 0 <= i < len(cells) else ''

    kind = num_str(get('kind'))
    name = get('name')
    if cmap.get('kind', -1) >= 0:
        if not re.fullmatch(r'\d+', kind):      # ตัดแถวหมายเหตุ/ยอดรวมทิ้ง
            return None
    elif not name:
        return None
    return {
        'kind': kind,
        'name': name,
        'dur': dur_from_text(get('dur')) if cmap.get('dur', -1) >= 0 else 'any',
        'trade': norm_bool(get('trade')) if cmap.get('trade', -1) >= 0 else 'any',
        'qty': get('qty') if cmap.get('qty', -1) >= 0 else '',
    }


def open_workbook(path):
    """เปิดไฟล์ Excel ครั้งเดียวแล้วใช้ซ้ำ — ไฟล์ใหญ่การเปิดใหม่ทุกครั้งช้ามาก"""
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def scan_sheets_wb(wb, progress=None):
    """คืน [(ชื่อชีท, จำนวนบล็อกตารางไอเทมที่เจอ)]"""
    out = []
    names = wb.sheetnames
    for k, name in enumerate(names):
        n = 0
        try:
            for row in wb[name].iter_rows(min_row=1, max_row=SCAN_ROWS, values_only=True):
                if row and is_header_row(row):
                    n += 1
        except Exception:
            n = 0
        out.append((name, n))
        if progress:
            progress(k + 1, len(names), name)
    return out


def scan_sheets(path):
    wb = open_workbook(path)
    try:
        return scan_sheets_wb(wb)
    finally:
        wb.close()


def parse_sheet_wb(wb, sheet):
    """อ่านชีทเดียว คืน (headers, blocks)
       blocks = [{'row': แถวหัวตาราง, 'headers': [...], 'data': [[cell, …], …]}]"""
    rows = [list(r) if r else []
            for r in wb[sheet].iter_rows(min_row=1, max_row=SCAN_ROWS, values_only=True)]

    blocks, i = [], 0
    while i < len(rows):
        if rows[i] and is_header_row(rows[i]):
            headers = [num_str(c) for c in rows[i]]
            data, j = [], i + 1
            while j < len(rows) and len(data) < 500:
                r = rows[j]
                if not r or is_header_row(r) or not any(num_str(c) for c in r):
                    break
                data.append([num_str(c) for c in r])
                j += 1
            blocks.append({'row': i + 1, 'headers': headers, 'data': data})
            i = max(j, i + 1)
        else:
            i += 1
    return (blocks[0]['headers'] if blocks else []), blocks


def parse_sheet(path, sheet):
    wb = open_workbook(path)
    try:
        return parse_sheet_wb(wb, sheet)
    finally:
        wb.close()


def blocks_to_rows(blocks, cmap):
    out = []
    for b in blocks:
        for cells in b['data']:
            row = cells_to_row(cells, cmap)
            if row:
                out.append(row)
    return out


# ============================================================================
#  [4] ตรรกะ deep check (ยกมาจากเครื่องมือเดิม)
# ============================================================================
def has_deep(c):
    return c.get('trade', 'any') != 'any' or str(c.get('qty', '')).strip() != '' \
        or str(c.get('dur', 'any')).strip().lower() != 'any'


def duration_only_numeric(c):
    """กรองแค่ระยะเวลาเป็นตัวเลข → ใช้ฟิลเตอร์บนหน้า list ได้ ไม่ต้องเปิดรายละเอียด"""
    return c.get('trade', 'any') == 'any' and not str(c.get('qty', '')).strip() \
        and re.fullmatch(r'\d+', str(c.get('dur', '')).strip() or '')


def match_deep(detail, crit):
    notes = []
    want = str(crit.get('dur', 'any')).strip().lower()
    actual = str(detail.get('duration') or '').strip()
    if want != 'any':
        if want == '':
            if actual not in ('', '0'):
                return False, f'ไม่ใช่ถาวร (dur={actual})'
            notes.append('ถาวร')
        elif actual != want:
            return False, f'dur {actual or "-"} ≠ {want}'
        else:
            notes.append(f'{want} วัน')
    elif actual:
        notes.append(f'{actual} วัน')

    trade = detail.get('trade')
    if crit.get('trade', 'any') != 'any':
        if trade is None:
            return False, 'อ่านค่าแลกเปลี่ยนไม่ได้'
        if trade != (crit['trade'] == 'yes'):
            return False, 'trade=' + ('Y' if trade else 'N')
    if trade is not None:
        notes.append('trade=' + ('Y' if trade else 'N'))

    q = str(crit.get('qty', '')).strip()
    if q:
        if str(detail.get('qty') or '').strip() != q:
            return False, f'qty {detail.get("qty") or "-"} ≠ {q}'
        notes.append('qty=' + q)
    return True, ' · '.join(notes)


# ============================================================================
#  [5] JS ที่ยิงเข้าไปอ่านหน้าเว็บ
# ============================================================================
JS_READ_ROWS = """
(col) => {
  const out = [];
  document.querySelectorAll('table tbody tr').forEach(tr => {
    const td = [...tr.querySelectorAll('td')];
    if (td.length < 4) return;
    const id = (td[col.id].innerText || '').trim();
    if (!/^\\d+$/.test(id)) return;
    out.push({ id,
      name: (td[col.name].innerText || '').trim(),
      type: (td[col.type].innerText || '').trim(),
      kind: (td[col.kind].innerText || '').trim() });
  });
  return out;
}"""

JS_TOTAL_PAGES = """
() => {
  const i = document.querySelector('input[type="number"]');
  if (i && i.max) return parseInt(i.max, 10) || 1;
  const m = (document.body.innerText || '').match(/of\\s+(\\d+)/);
  return m ? parseInt(m[1], 10) : 1;
}"""

JS_READ_DETAIL = """
(sel) => {
  const v = s => { const e = document.querySelector(s); return e ? (e.value || '').trim() : null; };
  function near(el, n = 5) {
    let c = el;
    for (let i = 0; i < n && c.parentElement; i++) {
      c = c.parentElement;
      const t = (c.innerText || '').trim();
      if (t && t.length < 160) return t;
    }
    return '';
  }
  let trade = null;
  const boxes = [...document.querySelectorAll('input[type="checkbox"]'),
                 ...document.querySelectorAll('[role="checkbox"]'),
                 ...document.querySelectorAll('[role="switch"]')];
  for (const cb of boxes) {
    if (near(cb).includes(sel.tradeLabel)) {
      trade = cb.type === 'checkbox' ? !!cb.checked
        : (cb.getAttribute('data-state') === 'checked' || cb.getAttribute('aria-checked') === 'true');
      break;
    }
  }
  return { name: v(sel.name), kind: v(sel.kind), price: v(sel.price),
           duration: v(sel.duration), qty: v(sel.qty), trade };
}"""

JS_CLICK_ROW = """
(id) => {
  for (const tr of document.querySelectorAll('table tbody tr')) {
    const td = tr.querySelectorAll('td');
    if (td.length && (td[0].innerText || '').trim() === String(id)) {
      const a = tr.querySelector('a');
      if (a) { a.click(); return true; }
    }
  }
  return false;
}"""

# ---------------------------------------------------------------------------
#  JS สำหรับหน้า "สร้างไอเทม"
#  เว็บเป็น React — กรอกด้วย .value เฉยๆ ไม่พอ ต้องใช้ native setter
#  แล้วยิง event input/change เอง ไม่งั้น React ไม่รู้ว่าค่าเปลี่ยน
#  (วิธีเดียวกับเครื่องมือเดิม)
# ---------------------------------------------------------------------------
JS_SET_VALUE = """
([sel, val]) => {
  const el = document.querySelector(sel);
  if (!el) return 'ไม่เจอช่อง';
  el.scrollIntoView({block:'center'});
  const proto = el.tagName === 'TEXTAREA'
      ? HTMLTextAreaElement.prototype : HTMLInputElement.prototype;
  const d = Object.getOwnPropertyDescriptor(proto, 'value');
  if (d && d.set) { d.set.call(el, String(val)); } else { el.value = String(val); }
  el.dispatchEvent(new Event('input',  {bubbles:true}));
  el.dispatchEvent(new Event('change', {bubbles:true}));
  return 'ok';
}"""

# กรอกโดยอ้างข้อความที่อยู่ข้างๆ — ใช้เป็นตัวสำรองเวลา id เปลี่ยน
JS_SET_BY_LABEL = """
([lbl, val]) => {
  const vis = el => { if(!el) return false; const s=getComputedStyle(el);
    if(s.display==='none'||s.visibility==='hidden') return false;
    const r=el.getBoundingClientRect(); return r.width>0 && r.height>0; };
  const hit = t => { t=(t||'').trim();
    return t===lbl || t.replace(/\\s*\\*\\s*$/,'')===lbl ||
           (t.length<=lbl.length+3 && t.indexOf(lbl)===0); };
  const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
  let node, target=null;
  while ((node = w.nextNode()) && !target) {
    if (!hit(node.textContent)) continue;
    let c = node.parentElement;
    for (let i=0; i<5 && c && !target; i++) {
      for (const el of c.querySelectorAll('input,textarea')) {
        if (vis(el) && el.type!=='checkbox' && el.type!=='file') { target = el; break; }
      }
      c = c.parentElement;
    }
  }
  if (!target) return 'ไม่เจอช่อง';
  target.scrollIntoView({block:'center'});
  const proto = target.tagName === 'TEXTAREA'
      ? HTMLTextAreaElement.prototype : HTMLInputElement.prototype;
  Object.getOwnPropertyDescriptor(proto,'value').set.call(target, String(val));
  target.dispatchEvent(new Event('input',  {bubbles:true}));
  target.dispatchEvent(new Event('change', {bubbles:true}));
  return 'ok';
}"""

# ติ๊ก/ปลดติ๊ก checkbox หรือสวิตช์ ที่อยู่ใกล้ข้อความที่ระบุ
# เดินหาจาก "ข้อความ" ไม่ใช่ตัวอิลิเมนต์ใบสุดท้าย เพราะของจริงป้ายมักมีคำอธิบายต่อท้าย
# อยู่ในก้อนเดียวกัน เช่น  แลกเปลี่ยนได้ (ไม่เลือก = "ผูกมัดไอดี")  -> ป้ายไม่ใช่ใบ
# คืนค่าแยกให้ด้วยว่า "เปลี่ยนให้แล้ว" หรือ "เป็นแบบนี้อยู่แล้ว" จะได้ไม่ขึ้น Log เหมือนพัง
JS_SWITCH = """
([lbl, act]) => {
  const vis = el => { if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden' || s.opacity === '0') return false;
    const r = el.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  // อ่านสถานะ: เอา aria-checked / data-state ก่อน เพราะเว็บสมัยใหม่ใช้ปุ่มวาดเอง
  const rd = el => {
    const a = el.getAttribute('aria-checked');
    if (a === 'true' || a === 'false') return a === 'true';
    const d = el.getAttribute('data-state');
    if (d === 'checked' || d === 'unchecked') return d === 'checked';
    return !!el.checked;
  };
  // เลือกตัวที่ "มองเห็นได้" ก่อนเสมอ
  // เว็บสมัยใหม่ซ่อน <input type=checkbox> ตัวจริงไว้ แล้ววาดปุ่มสวยๆ ทับ
  // ถ้าเผลอไปกดตัวที่ซ่อนอยู่ หน้าจอจะไม่เปลี่ยนตาม แล้วอ่านค่ากลับมาก็ไม่ตรง
  const pick = box => {
    const all = [...box.querySelectorAll(
      'input[type="checkbox"],[role="checkbox"],[role="switch"]')];
    if (!all.length) return null;
    return all.filter(vis)[0] || all[0];
  };
  const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
  let node, cb = null;
  while ((node = w.nextNode()) && !cb) {
    if ((node.textContent || '').trim().indexOf(lbl) !== 0) continue;
    let n = node.parentElement;
    for (let i = 0; i < 6 && n && !cb; i++) { cb = pick(n); n = n.parentElement; }
  }
  if (!cb) return 'ไม่เจอสวิตช์';
  if (act === 'read') return 'ok|' + (rd(cb) ? '1' : '0');
  cb.scrollIntoView({block: 'center'});
  if (act === 'label') {                 // กดที่ตัวควบคุมไม่ได้ผล ลองกดที่ป้ายแทน
    let lab = cb.closest ? cb.closest('label') : null;
    if (!lab && cb.id) lab = document.querySelector('label[for="' + cb.id + '"]');
    if (!lab) return 'ไม่มีป้ายให้กด';
    lab.click();
    return 'ok';
  }
  cb.click();
  return 'ok';
}"""

# อ่าน "ประเภทไอเทม (game_item_type)" ที่หน้าเว็บตั้งไว้ตอนนี้
# หน้าสร้างไอเทมล็อกค่านี้เป็น GENERAL ไว้ กดเลือกไม่ได้ (คนกดเองก็ไม่ได้)
# เลยต้องอ่านค่าปัจจุบันมาเทียบ ถ้าตรงกับที่ต้องการอยู่แล้วก็จบ ไม่ใช่ข้อผิดพลาด
JS_READ_TYPE = """
(labels) => {
  const tx = el => (el.innerText || el.textContent || '').trim();
  const hit = (box, lbl) => {
    if (!box) return '';
    const s = (box.tagName === 'SELECT') ? box : box.querySelector('select');
    if (s && s.selectedIndex >= 0) return s.options[s.selectedIndex].text.trim();
    const q = 'input:not([type=checkbox]):not([type=file]):not([type=radio])';
    const inp = (box.tagName === 'INPUT' && box.type !== 'checkbox')
        ? box : box.querySelector(q);
    if (inp && String(inp.value || '').trim()) return String(inp.value).trim();
    const t = tx(box);
    if (t && t.length <= 40 && t.indexOf(lbl) !== 0) return t;
    return '';
  };
  for (const lbl of labels) {
    const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
    let node;
    while ((node = w.nextNode())) {
      const t = (node.textContent || '').trim();
      if (!t || t.indexOf(lbl) !== 0) continue;
      const own = node.parentElement;
      if (own && own.getAttribute && own.getAttribute('for')) {
        const v = hit(document.getElementById(own.getAttribute('for')), lbl);
        if (v) return v;
      }
      // ค่าของช่องอยู่ "ถัดจากป้าย" เสมอ — ไล่จากพี่น้องที่ตามมา ไม่ใช่กวาดทั้งกล่อง
      // (กวาดทั้งกล่องจะไปเจอหัวข้อการ์ดที่อยู่ก่อนหน้าป้าย)
      let el = own;
      for (let up = 0; up < 3 && el; up++) {
        let sib = el.nextElementSibling;
        for (let k = 0; k < 3 && sib; k++) {
          const v = hit(sib, lbl);
          if (v) return v;
          sib = sib.nextElementSibling;
        }
        el = el.parentElement;
      }
    }
  }
  return '';
}"""

# เลือกประเภทไอเทม (game_item_type) — รองรับทั้ง <select> และ combobox ของ React
JS_PICK_TYPE = """
([val, ph]) => {
  for (const s of document.querySelectorAll('select')) {
    const opts = [...s.options];
    let i = opts.findIndex(o => o.text.trim().toLowerCase() === val.toLowerCase());
    if (i < 0) i = opts.findIndex(o => o.text.trim().toLowerCase().includes(val.toLowerCase()));
    if (i >= 0) {
      const d = Object.getOwnPropertyDescriptor(HTMLSelectElement.prototype, 'value');
      d.set.call(s, opts[i].value);
      s.dispatchEvent(new Event('input',  {bubbles:true}));
      s.dispatchEvent(new Event('change', {bubbles:true}));
      return 'ok';
    }
  }
  return 'combobox';
}"""

# หาช่องอัปโหลดรูป แล้วติดป้ายไว้ให้ playwright จับถูกตัว
# อ้างจากข้อความ "รูปภาพไอเทม" ไม่อ้างตำแหน่ง — ฟอร์มสลับที่เมื่อไหร่ก็ยังหาเจอ
JS_MARK_FILE = """
([lbl]) => {
  for (const el of document.querySelectorAll('[data-trpu-pic]'))
    el.removeAttribute('data-trpu-pic');
  const all = [...document.querySelectorAll('input[type=file]')];
  if (!all.length) return 'ไม่เจอช่องอัปโหลดรูปในหน้านี้';
  let target = null;
  const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
  let node;
  while ((node = w.nextNode()) && !target) {
    if ((node.textContent || '').indexOf(lbl) < 0) continue;
    let c = node.parentElement;
    for (let i = 0; i < 6 && c && !target; i++) {
      const f = c.querySelector('input[type=file]');
      if (f) target = f;
      c = c.parentElement;
    }
  }
  if (!target && all.length === 1) target = all[0];
  if (!target) return 'เจอช่องอัปโหลด ' + all.length + ' ช่อง แต่ไม่รู้ว่าช่องไหนคือ ' + lbl;
  target.setAttribute('data-trpu-pic', '1');
  target.scrollIntoView({block: 'center'});
  return 'ok';
}"""

# อ่านสภาพกล่องอัปโหลดรูป ณ ตอนนี้
# สำคัญ: เว็บจริงอัปโหลดไฟล์ทันทีแล้ว "ล้างช่อง input ทิ้ง" (React คุมค่าเอง)
# เพราะงั้นจะดูแค่ el.files.length ไม่ได้ ต้องดูหลักฐานฝั่งหน้าเว็บด้วย
# (รูปตัวอย่างที่ขึ้นมา / ข้อความแจ้งอัปโหลดสำเร็จ) ไม่งั้นจะหาว่าเว็บไม่รับไฟล์ทั้งที่รับแล้ว
JS_PIC_STATE = """
([lbl]) => {
  const el = document.querySelector('input[type=file][data-trpu-pic]');
  const n = (el && el.files) ? el.files.length : (el ? 0 : -1);
  let box = null;
  const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
  let node;
  while ((node = w.nextNode()) && !box) {
    if ((node.textContent || '').indexOf(lbl) < 0) continue;
    let c = node.parentElement;
    for (let i = 0; i < 6 && c && !box; i++) {
      if (c.querySelector('input[type=file]')) box = c;
      c = c.parentElement;
    }
  }
  if (!box && el) {
    box = el;
    for (let i = 0; i < 4 && box.parentElement; i++) box = box.parentElement;
  }
  const imgs = box ? [...box.querySelectorAll('img')] : [];
  const body = document.body ? (document.body.innerText || '') : '';
  const low = body.toLowerCase();
  const toast = (low.indexOf('upload') >= 0 || low.indexOf('อัปโหลด') >= 0) &&
                (body.indexOf('สำเร็จ') >= 0 || low.indexOf('success') >= 0);
  return {n: n, name: (n > 0) ? el.files[0].name : '',
          imgs: imgs.length,
          src: imgs.length ? String(imgs[0].src || '').slice(-80) : '',
          toast: toast};
}"""

# อ่านค่าที่อยู่ในฟอร์มตอนนี้กลับมา — ใช้ตรวจว่ากรอกเข้าไปจริงไหม
JS_READ_FORM = """
(sel) => {
  const v = s => { const e = document.querySelector(s); return e ? e.value : null; };
  return { name: v(sel.name), kind: v(sel.kind), price: v(sel.price),
           duration: v(sel.duration), qty: v(sel.qty), mail: v(sel.mail),
           desc: v(sel.desc) };
}"""

# ---------------------------------------------------------------------------
#  JS สำหรับหน้า "สร้างบันเดิล"
#  การ์ดแต่ละใบในรายการไอเท็มมีช่อง "จำนวน (Quantity)" กับดรอปดาวน์ Tier
#  หาโดยไล่จากข้อความ แล้ววิ่งขึ้นไปหา container ของการ์ดใบนั้น
# ---------------------------------------------------------------------------
# --- ตัวช่วย: หา "กล่องเพิ่มของเข้า Bundle" ให้เจอก่อน แล้วค่อยหาช่องข้างใน ---
#     สำคัญมาก: ถ้าไม่จำกัดขอบเขต จะไปเจอช่อง "ชื่อ Bundle" ที่แผงขวาแทน
_JS_BOX = """
  function addBox(){
    // กล่องจริงต้องมีช่องกรอกอยู่ข้างใน ไม่ใช่แค่ป้ายข้อความที่เขียนว่า "เพิ่มของเข้า Bundle"
    let best = null, bl = 1e9;
    for (const el of document.querySelectorAll('div,section,form')) {
      const t = el.textContent || '';
      if (t.indexOf('เพิ่มของเข้า Bundle') < 0) continue;
      if (!el.querySelector('input')) continue;
      if (t.length < bl) { bl = t.length; best = el; }
    }
    return best;
  }
  function vis(el){
    if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden') return false;
    const r = el.getBoundingClientRect();
    return r.width > 0 && r.height > 0;
  }
  function setVal(el, v){
    el.scrollIntoView({block:'center'});
    el.focus();
    const p = el.tagName === 'TEXTAREA' ? HTMLTextAreaElement.prototype : HTMLInputElement.prototype;
    Object.getOwnPropertyDescriptor(p, 'value').set.call(el, String(v));
    el.dispatchEvent(new Event('input', {bubbles:true}));
    el.dispatchEvent(new Event('change', {bubbles:true}));
  }
"""

JS_BUNDLE_SEARCH = """
(val) => {
""" + _JS_BOX + """
  let el = document.querySelector('input[placeholder*="ค้นหาชื่อ Item"]');
  if (!el) {
    const box = addBox();
    if (!box) return 'ไม่เจอกล่องเพิ่มของเข้า Bundle';
    el = box.querySelector('input[placeholder*="ค้นหา"]')
      || [...box.querySelectorAll('input')].filter(
           e => vis(e) && e.type !== 'checkbox' && e.type !== 'file' &&
                e.getBoundingClientRect().width > 200)[0];
  }
  if (!el) return 'ไม่เจอช่องค้นหาในกล่อง';
  setVal(el, val);
  return 'ok';
}"""

# เลือกแถวผลค้นหาที่ "ID: <เลข>" ตรงเป๊ะ แล้วกดปุ่มเพิ่มของแถวนั้น
# ห้ามกดปุ่มสุดท้ายมั่วๆ เพราะลิสต์ที่ยังไม่กรองมีเป็นร้อยหน้า
JS_BUNDLE_PICK_ROW = """
([id]) => {
""" + _JS_BOX + """
  // อ่านเลข ID จากข้อความของแถว โดยไม่ใช้ regex เลย
  // (เคยพลาดเพราะ backslash ใน regex หายตอนส่งเข้าเบราว์เซอร์ จนไม่แมตช์อะไรเลย)
  function idOf(t) {
    const i = t.indexOf('ID:');
    if (i < 0) return null;
    let j = i + 3, out = '';
    while (j < t.length && t.charCodeAt(j) <= 32) j++;
    while (j < t.length && t.charAt(j) >= '0' && t.charAt(j) <= '9') { out += t.charAt(j); j++; }
    return out || null;
  }
  const box = addBox();
  if (!box) return 'ไม่เจอกล่องเพิ่มของเข้า Bundle';
  const want = String(id);
  let hit = null, hl = 1e9;
  for (const el of box.querySelectorAll('div,li,tr')) {
    const t = el.textContent || '';
    if (t.indexOf('เพิ่ม') < 0) continue;
    if (idOf(t) !== want) continue;
    if (t.length < hl) { hl = t.length; hit = el; }
  }
  if (!hit) return 'ไม่เจอแถวที่ ID ตรง';
  const btn = [...hit.querySelectorAll('button')].filter(
    b => (b.textContent || '').indexOf('เพิ่ม') >= 0 && vis(b))[0];
  if (!btn) return 'เจอแถวแล้วแต่ไม่เจอปุ่มเพิ่ม';
  btn.scrollIntoView({block: 'center'});
  btn.click();
  let seen = (hit.textContent || '');
  let clean = '';
  for (let k = 0; k < seen.length && clean.length < 90; k++) {
    const c = seen.charAt(k);
    clean += (seen.charCodeAt(k) <= 32) ? ' ' : c;
  }
  return 'ok|' + clean;
}"""

JS_BUNDLE_NAME = """
(val) => {
  const el = document.querySelector('input[placeholder*="ชื่อ Bundle"]');
  if (!el) return 'ไม่เจอช่องชื่อ Bundle';
  el.scrollIntoView({block:'center'});
  el.focus();
  Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value')
        .set.call(el, String(val));
  el.dispatchEvent(new Event('input', {bubbles:true}));
  el.dispatchEvent(new Event('change', {bubbles:true}));
  return 'ok';
}"""

JS_BUNDLE_NAME_GET = """
() => {
  const el = document.querySelector('input[placeholder*="ชื่อ Bundle"]');
  return el ? el.value : null;
}"""

JS_BUNDLE_WALLET_QTY = """
(val) => {
""" + _JS_BOX + """
  const box = addBox();
  if (!box) return 'ไม่เจอกล่องเพิ่มของเข้า Bundle';
  const cands = [...box.querySelectorAll('input')].filter(
    e => vis(e) && e.type !== 'checkbox' && e.type !== 'file' &&
         e.getBoundingClientRect().width < 200);
  const el = cands[cands.length - 1];
  if (!el) return 'ไม่เจอช่องจำนวนในกล่อง';
  setVal(el, val);
  return 'ok';
}"""

# กด "ขยายทั้งหมด" — การ์ดที่ย่ออยู่จะไม่มีช่องจำนวน/Tier ให้กรอก
JS_BUNDLE_EXPAND = """
() => {
  for (const b of document.querySelectorAll('button,[role="button"]')) {
    if ((b.textContent || '').indexOf('ขยายทั้งหมด') >= 0) { b.click(); return 'ok'; }
  }
  return 'ไม่เจอปุ่มขยายทั้งหมด';
}"""

# --- ค้นหา/เลือกแถว: ไม่ยึดตำแหน่ง ไม่ยึดเวลา ---
# เว็บโหลดผลช้าบ้างเร็วบ้าง และ ID สั้นๆ (เช่น 134) จะเจอผลเป็นร้อย กระจายหลายหน้า
# เพราะงั้นต้อง "รอจนกว่าจะเจอแถวที่ ID ตรงจริงๆ" และเปิดหน้าถัดไปหาต่อได้
JS_BUNDLE_ROW_STATE = """
([id]) => {
""" + _JS_BOX + """
  function idAfter(t, key) {
    const i = t.indexOf(key);
    if (i < 0) return null;
    let j = i + key.length, out = '';
    while (j < t.length && t.charCodeAt(j) <= 32) j++;
    while (j < t.length && t.charAt(j) >= '0' && t.charAt(j) <= '9') { out += t.charAt(j); j++; }
    return out || null;
  }
  const box = addBox();
  if (!box) return {state: 'nobox'};
  const txt = box.textContent || '';
  const loading = txt.indexOf('กำลังโหลด') >= 0;
  let rows = 0, hit = false;
  for (const el of box.querySelectorAll('div,li,tr')) {
    const t = el.textContent || '';
    if (t.indexOf('เพิ่ม') < 0) continue;
    const got = idAfter(t, 'ID:');
    if (got === null) continue;
    rows++;
    if (got === String(id)) hit = true;
  }
  return {state: loading ? 'loading' : 'ready', rows: rows, hit: hit};
}"""

# ป๊อปอัป "ยืนยันการสร้าง" ที่เด้งหลังกดปุ่มสร้าง — ต้องกดยืนยันให้ด้วย
# ห้ามกดโดนปุ่มยกเลิกเด็ดขาด เลยเทียบชื่อปุ่มแบบตรงๆ ไม่เดา
JS_BUNDLE_CONFIRM = """
([words, nos]) => {
  function vis(el){
    if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden' || s.opacity === '0') return false;
    const r = el.getBoundingClientRect();
    return r.width > 0 && r.height > 0;
  }
  // กล่องที่เด้งทับหน้าจอ (dialog / alertdialog / modal)
  const boxes = [...document.querySelectorAll(
    '[role="dialog"],[role="alertdialog"],dialog,[data-state="open"]')].filter(vis);
  if (!boxes.length) return 'ไม่มีป๊อปอัป';
  const box = boxes[boxes.length - 1];
  const btns = [...box.querySelectorAll('button,[role="button"],a')].filter(vis);
  const seen = [];
  for (const b of btns) {
    const t = (b.textContent || '').trim();
    seen.push(t);
    if (nos.indexOf(t) >= 0) continue;          // ยกเลิก/ปิด — ห้ามกด
    if (words.indexOf(t) >= 0) { b.click(); return 'ok|' + t; }
  }
  return 'มีป๊อปอัปแต่ไม่เจอปุ่มยืนยัน (มี: ' + seen.join(', ') + ')';
}"""

# ป๊อปอัปยืนยันตัวเดียวกันนี้ใช้ได้ทั้งหน้าสร้าง Item และหน้าสร้าง Bundle
JS_CONFIRM_POPUP = JS_BUNDLE_CONFIRM

# หาเลข Bundle ที่เพิ่งสร้าง — จาก URL ก่อน ไม่มีค่อยหาในหน้า
JS_BUNDLE_MADE_ID = """
() => {
  function digitsAfter(t, key) {
    const i = t.indexOf(key);
    if (i < 0) return null;
    let j = i + key.length, out = '';
    while (j < t.length && (t.charAt(j) === ' ' || t.charAt(j) === ':' ||
                            t.charAt(j) === '#' || t.charAt(j) === '=')) j++;
    while (j < t.length && t.charAt(j) >= '0' && t.charAt(j) <= '9') { out += t.charAt(j); j++; }
    return out || null;
  }
  const u = location.pathname;
  const parts = u.split('/').filter(x => x);
  const last = parts[parts.length - 1] || '';
  let onlyNum = last.length > 0;
  for (let i = 0; i < last.length; i++)
    if (last.charAt(i) < '0' || last.charAt(i) > '9') onlyNum = false;
  if (onlyNum && u.indexOf('bundle') >= 0) return last;
  const t = document.body ? (document.body.innerText || '') : '';
  return digitsAfter(t, 'Bundle ID') || digitsAfter(t, 'bundleId') ||
         digitsAfter(t, 'รหัส Bundle') || '';
}"""

# อ่าน "Aztek Item Id" ของไอเทมที่เพิ่งสร้าง จากหน้าเว็บ (เผื่อ API ไม่ส่งเลขกลับมา)
# สร้างเสร็จเว็บมักเด้งไปหน้า .../items/<เลข>/edit  หรือโชว์เลขไว้ในหน้า
# ============================================================================
#  JS ของหน้าสร้าง Item Code (WR Master)
#  หน้านี้มีป้ายชื่อ "ซ้ำกันเป๊ะ" สองฝั่ง เช่น "จำนวนการใช้งานต่อ 1 User"
#  และ "จำนวนคงเหลือ" -> ทุกคำสั่งต้องบอกด้วยว่าจะเอาฝั่งไหน
#  ฝั่งขวา = กล่องที่มีหัวข้อ "ของรางวัล" · ฝั่งซ้าย = ทุกอย่างที่ไม่ได้อยู่ในกล่องนั้น
# ============================================================================
_JS_CODE_LIB = """
  const vis = el => { if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden') return false;
    const r = el.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const rdsw = el => {
    const a = el.getAttribute('aria-checked');
    if (a === 'true' || a === 'false') return a === 'true';
    const d = el.getAttribute('data-state');
    if (d === 'checked' || d === 'unchecked') return d === 'checked';
    return !!el.checked;
  };
  // กล่องของรางวัล = กล่องที่เล็กที่สุดที่ครอบข้อความ "ของรางวัล" ไว้
  function rewardBox() {
    const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
    let node;
    while ((node = w.nextNode())) {
      if ((node.textContent || '').trim().indexOf('ของรางวัล') !== 0) continue;
      let c = node.parentElement;
      for (let i = 0; i < 6 && c; i++) {
        if (c.querySelector('textarea') || c.querySelectorAll('input').length >= 3)
          return c;
        c = c.parentElement;
      }
    }
    return null;
  }
  function inSide(el, side) {
    const rw = rewardBox();
    if (!rw) return true;                       // หาไม่เจอ ก็ไม่ต้องกรอง
    const inside = rw.contains(el);
    return side === 'right' ? inside : !inside;
  }
  // หาช่องกรอก/ปุ่มที่อยู่ "ใต้ป้าย" ที่ระบุ และอยู่ฝั่งที่ต้องการ
  function findBy(lbl, side, pick) {
    const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
    let node;
    while ((node = w.nextNode())) {
      const t = (node.textContent || '').trim().replace(/\\s*\\*\\s*$/, '');
      if (t !== lbl && t.indexOf(lbl) !== 0) continue;
      const own = node.parentElement;
      if (!own || !inSide(own, side)) continue;
      if (own.getAttribute && own.getAttribute('for')) {
        const g = pick(document.getElementById(own.getAttribute('for')));
        if (g) return g;
      }
      let el = own;
      for (let up = 0; up < 4 && el; up++) {
        let sib = el.nextElementSibling;
        for (let k = 0; k < 3 && sib; k++) {
          const g = pick(sib);
          if (g && inSide(g, side)) return g;
          sib = sib.nextElementSibling;
        }
        const g2 = pick(el.parentElement);
        if (g2 && inSide(g2, side)) return g2;
        el = el.parentElement;
      }
    }
    return null;
  }
  const pickText = box => {
    if (!box) return null;
    const q = 'input:not([type=checkbox]):not([type=radio]):not([type=file]),textarea';
    if (box.matches && box.matches(q) && vis(box)) return box;
    const all = [...box.querySelectorAll(q)].filter(vis);
    return all[0] || null;
  };
  const pickSel = box => {
    if (!box) return null;
    if (box.tagName === 'SELECT' && vis(box)) return box;
    const all = [...box.querySelectorAll('select')].filter(vis);
    return all[0] || null;
  };
  const pickSw = box => {
    if (!box) return null;
    const q = 'input[type=checkbox],[role=checkbox],[role=switch]';
    if (box.matches && box.matches(q)) return vis(box) ? box : null;
    const all = [...box.querySelectorAll(q)];
    return all.filter(vis)[0] || null;
  };
  function setVal(el, val) {
    el.scrollIntoView({block: 'center'});
    const proto = el.tagName === 'TEXTAREA'
        ? HTMLTextAreaElement.prototype : HTMLInputElement.prototype;
    const d = Object.getOwnPropertyDescriptor(proto, 'value');
    if (d && d.set) d.set.call(el, String(val)); else el.value = String(val);
    el.dispatchEvent(new Event('input', {bubbles: true}));
    el.dispatchEvent(new Event('change', {bubbles: true}));
  }
"""

# กรอกช่องข้อความ / อ่านค่ากลับ — ต้องบอกฝั่ง (left/right)
JS_CODE_TEXT = """
([lbl, side, act, val]) => {
""" + _JS_CODE_LIB + """
  const el = findBy(lbl, side, pickText);
  if (!el) return 'ไม่เจอช่อง';
  if (act === 'read') return 'ok|' + String(el.value == null ? '' : el.value);
  setVal(el, val);
  return 'ok';
}"""

# เลือกค่าในดรอปดาวน์ (ประเภท / ประเภทของ Code)
JS_CODE_SELECT = """
([lbl, side, act, val]) => {
""" + _JS_CODE_LIB + """
  const s = findBy(lbl, side, pickSel);
  if (!s) return 'ไม่เจอดรอปดาวน์';
  if (act === 'read')
    return 'ok|' + (s.selectedIndex >= 0 ? s.options[s.selectedIndex].text.trim() : '');
  const opts = [...s.options];
  const low = String(val).toLowerCase();
  let i = opts.findIndex(o => o.text.trim().toLowerCase() === low);
  if (i < 0) i = opts.findIndex(o => o.text.trim().toLowerCase().indexOf(low) === 0);
  if (i < 0) i = opts.findIndex(o => o.text.trim().toLowerCase().indexOf(low) >= 0);
  if (i < 0) return 'ไม่มีตัวเลือกนี้ (มี: ' + opts.map(o => o.text.trim()).join(', ') + ')';
  const d = Object.getOwnPropertyDescriptor(HTMLSelectElement.prototype, 'value');
  d.set.call(s, opts[i].value);
  s.dispatchEvent(new Event('input', {bubbles: true}));
  s.dispatchEvent(new Event('change', {bubbles: true}));
  return 'ok';
}"""

# สวิตช์ของหน้านี้ — แยกฝั่งได้ และอ่าน/กดแยกกัน (รอ React วาดใหม่ฝั่ง Python)
JS_CODE_SWITCH = """
([lbl, side, act]) => {
""" + _JS_CODE_LIB + """
  const cb = findBy(lbl, side, pickSw);
  if (!cb) return 'ไม่เจอสวิตช์';
  if (act === 'read') return 'ok|' + (rdsw(cb) ? '1' : '0');
  cb.scrollIntoView({block: 'center'});
  cb.click();
  return 'ok';
}"""

# ปุ่ม "เลือก bundle" / ช่องค้นหา / แถวผลลัพธ์ในป๊อปอัป
JS_CODE_BUNDLE = """
([act, val]) => {
""" + _JS_CODE_LIB + """
  if (act === 'open') {
    const btns = [...document.querySelectorAll('button,[role=button],a')].filter(vis);
    const hit = btns.find(b => {
      const t = (b.textContent || '').trim();
      return t === 'เลือก bundle' || t === 'เปลี่ยน' || t.indexOf('เลือก bundle') === 0;
    });
    if (!hit) return 'ไม่เจอปุ่มเลือก bundle';
    hit.click();
    return 'ok';
  }
  if (act === 'search') {
    const ins = [...document.querySelectorAll('input[type=text],input:not([type])')]
        .filter(vis);
    const box = ins.find(i => {
      const p = (i.getAttribute('placeholder') || '');
      return p.indexOf('ค้นหา') >= 0 || p.toLowerCase().indexOf('search') >= 0;
    });
    if (!box) return 'ไม่เจอช่องค้นหา bundle';
    setVal(box, val);
    return 'ok';
  }
  if (act === 'pick') {
    // แถวที่ "เลข id ตรงเป๊ะ" เท่านั้น กันไปคลิกตัวที่เลขคล้ายกัน
    const want = String(val);
    const all = [...document.querySelectorAll('[data-id],li,tr,div')].filter(vis);
    for (const el of all) {
      if (el.querySelector && el.querySelector('[data-id],li,tr')) continue;
      const t = (el.textContent || '').trim();
      if (!t || t.length > 160) continue;
      const byAttr = el.getAttribute && el.getAttribute('data-id') === want;
      const byText = new RegExp('(^|[^0-9])(ID|Id|id)?\\\\s*:?\\\\s*#?' +
                                want + '([^0-9]|$)').test(t);
      if (byAttr || byText) { el.click(); return 'ok|' + t.slice(0, 90); }
    }
    return 'ไม่เจอ bundle เลข ' + want + ' ในรายการ';
  }
  if (act === 'read') {
    const rw = rewardBox() || document.body;
    const w = document.createTreeWalker(rw, NodeFilter.SHOW_TEXT, null, false);
    let node;
    while ((node = w.nextNode())) {
      const t = (node.textContent || '').trim();
      if (t.indexOf('#') === 0 && t.length > 1) return 'ok|' + t.slice(0, 90);
    }
    return 'ok|';
  }
  return 'ไม่รู้จักคำสั่ง';
}"""

JS_ITEM_MADE_ID = """
() => {
  function digitsAfter(t, key) {
    const i = t.indexOf(key);
    if (i < 0) return null;
    let j = i + key.length, out = '';
    while (j < t.length && (t.charAt(j) === ' ' || t.charAt(j) === ':' ||
                            t.charAt(j) === '#' || t.charAt(j) === '=')) j++;
    while (j < t.length && t.charAt(j) >= '0' && t.charAt(j) <= '9') { out += t.charAt(j); j++; }
    return out || null;
  }
  const u = location.pathname;
  if (u.indexOf('item') >= 0) {
    const parts = u.split('/').filter(x => x);
    for (let i = parts.length - 1; i >= 0; i--) {
      const p = parts[i];
      let num = p.length > 0;
      for (let j = 0; j < p.length; j++)
        if (p.charAt(j) < '0' || p.charAt(j) > '9') num = false;
      if (num) return p;
    }
  }
  const t = document.body ? (document.body.innerText || '') : '';
  return digitsAfter(t, 'Aztek Item Id') || digitsAfter(t, 'Item ID') ||
         digitsAfter(t, 'ID') || '';
}"""

# อ่าน "หน้า X / Y" ของผลค้นหา — จะได้รู้ว่ามีกี่หน้า และอยู่หน้าไหนแล้ว
JS_BUNDLE_PAGES = """
() => {
""" + _JS_BOX + """
  const box = addBox();
  if (!box) return {page: 0, total: 0};
  const t = box.textContent || '';
  const i = t.indexOf('หน้า');
  if (i < 0) return {page: 0, total: 0};
  const nums = [];
  let cur = '';
  for (let j = i; j < t.length && nums.length < 2; j++) {
    const c = t.charAt(j);
    if (c >= '0' && c <= '9') cur += c;
    else if (cur) { nums.push(parseInt(cur, 10)); cur = ''; }
  }
  if (cur && nums.length < 2) nums.push(parseInt(cur, 10));
  return {page: nums[0] || 0, total: nums[1] || 0};
}"""

# กดหน้าถัดไปของผลค้นหา (เฉพาะในกล่องเพิ่มของ)
JS_BUNDLE_NEXT_PAGE = """
() => {
""" + _JS_BOX + """
  const box = addBox();
  if (!box) return 'nobox';
  for (const b of box.querySelectorAll('button,[role="button"],a')) {
    const t = (b.textContent || '').trim();
    if (t !== 'ถัดไป' && t !== 'Next') continue;
    if (b.disabled || b.getAttribute('aria-disabled') === 'true') return 'last';
    if (!vis(b)) continue;
    b.click();
    return 'ok';
  }
  return 'nonext';
}"""

# --- การ์ดในรายการไอเท็ม: หาโดย "Item ID" ไม่ใช่ลำดับที่เท่าไหร่ ---
# หน้าเว็บเปิด/ปิด/ย่อ/ขยายได้ ลำดับเลยเชื่อไม่ได้
JS_BUNDLE_MARK_CARD = """
([itemId, what, tiers]) => {
  function vis(el){
    if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden') return false;
    const r = el.getBoundingClientRect();
    return r.width > 0 && r.height > 0;
  }
  function idAfter(t, key) {
    const i = t.indexOf(key);
    if (i < 0) return null;
    let j = i + key.length, out = '';
    while (j < t.length && t.charCodeAt(j) <= 32) j++;
    while (j < t.length && t.charAt(j) >= '0' && t.charAt(j) <= '9') { out += t.charAt(j); j++; }
    return out || null;
  }
  // การ์ด = กล่องเล็กที่สุดที่มีทั้ง "Item ID <เลขนี้>" และช่อง "จำนวน (Quantity)"
  let card = null, bl = 1e9;
  for (const el of document.querySelectorAll('div,section,li')) {
    const t = el.textContent || '';
    if (t.indexOf('จำนวน (Quantity)') < 0) continue;
    if (idAfter(t, 'Item ID') !== String(itemId)) continue;
    if (!vis(el)) continue;
    if (t.length < bl) { bl = t.length; card = el; }
  }
  if (!card) return 'ไม่เจอการ์ดของไอเทม ' + itemId;
  let ctl = null;
  if (what === 'qty') {
    // หาเฉพาะ "ข้างในการ์ดใบนี้" เท่านั้น ห้ามไล่ขึ้นไปเกินขอบการ์ด
    // (เคยพลาด: ไล่ขึ้นไปจนออกนอกการ์ด แล้วไปคว้าช่องของการ์ดใบแรกแทน)
    ctl = [...card.querySelectorAll('input')].filter(
      e => e.type !== 'checkbox' && e.type !== 'file' && vis(e))[0];
  } else {
    ctl = card.querySelector('select');
    if (!ctl) {
      const names = (tiers || []).concat(['เลือก Tier']);
      ctl = [...card.querySelectorAll('[role="combobox"],button')].filter(
        e => vis(e) && names.indexOf((e.textContent || '').trim()) >= 0)[0];
    }
  }
  if (!ctl) return 'ไม่เจอช่อง ' + what + ' ในการ์ดของไอเทม ' + itemId;
  document.querySelectorAll('[data-trpu]').forEach(e => e.removeAttribute('data-trpu'));
  ctl.setAttribute('data-trpu', what);
  ctl.scrollIntoView({block: 'center'});
  return 'ok|' + (ctl.value !== undefined && ctl.tagName === 'INPUT'
                  ? ctl.value : (ctl.textContent || '').trim());
}"""

# การ์ด famepoint / exp ไม่มี Item ID เลยต้องหาจาก "ชื่อการ์ด" แทน
# แต่ยังใช้หลักเดิม: หากล่องที่เล็กที่สุดที่ครอบทั้งชื่อและช่องที่ต้องการ
JS_BUNDLE_MARK_NAMED = """
([name, what, tiers]) => {
  function vis(el){
    if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden') return false;
    const r = el.getBoundingClientRect();
    return r.width > 0 && r.height > 0;
  }
  const need = (what === 'kind') ? 'ประเภท' : 'Tier';
  let card = null, bl = 1e9;
  for (const el of document.querySelectorAll('div,section,li')) {
    const t = el.textContent || '';
    if (t.indexOf(name) < 0) continue;
    if (t.indexOf(need) < 0) continue;
    if (!vis(el)) continue;
    if (t.length < bl) { bl = t.length; card = el; }
  }
  if (!card) return 'ไม่เจอการ์ด ' + name;
  let ctl = null;
  if (what === 'kind') {
    ctl = card.querySelector('select');
    if (!ctl) {
      ctl = [...card.querySelectorAll('[role="combobox"],button')].filter(
        e => vis(e) && (e.textContent || '').indexOf('WALLET') >= 0)[0];
    }
    if (!ctl) {
      ctl = [...card.querySelectorAll('[role="combobox"]')].filter(vis)[0];
    }
  } else if (what === 'tier') {
    ctl = card.querySelector('select');
    if (!ctl) {
      const names = (tiers || []).concat(['เลือก Tier']);
      ctl = [...card.querySelectorAll('[role="combobox"],button')].filter(
        e => vis(e) && names.indexOf((e.textContent || '').trim()) >= 0)[0];
    }
  } else {
    ctl = [...card.querySelectorAll('input')].filter(
      e => e.type !== 'checkbox' && e.type !== 'file' && vis(e))[0];
  }
  if (!ctl) return 'ไม่เจอช่อง ' + what + ' ในการ์ด ' + name;
  document.querySelectorAll('[data-trpu]').forEach(e => e.removeAttribute('data-trpu'));
  ctl.setAttribute('data-trpu', what);
  ctl.scrollIntoView({block: 'center'});
  return 'ok|' + (ctl.tagName === 'INPUT' ? ctl.value : (ctl.textContent || '').trim());
}"""


# ===========================================================================
#  ตัวกลางเดียวสำหรับ "หาช่องในการ์ด แล้วลงมือเลย" — ครั้งเดียวจบใน JS
#
#  ทำไมต้องครั้งเดียวจบ:
#    เดิมใช้วิธี "ติดป้าย data-trpu ไว้ก่อน แล้วค่อยให้ Playwright มากด"
#    แต่เว็บเป็น React พอมีอะไรเปลี่ยน มัน render ใหม่ ป้ายที่ติดไว้หายเกลี้ยง
#    -> กดไม่โดน ขึ้น Timeout (เจอกับการ์ด Fame Point ที่เพิ่งถูกเพิ่มเข้ามา)
#
#  how  : 'id' = หาจาก Item ID ในการ์ด · 'name' = หาจากชื่อการ์ด (Fame Point / exp)
#  what : 'qty' | 'tier' | 'kind'
#  act  : 'read' | 'set' | 'click'
# ===========================================================================
# เช็กว่าดรอปดาวน์ตัวก่อนหน้าปิดสนิทหรือยัง
# เว็บพวกนี้ตอนเปิดดรอปดาวน์จะล็อกหน้าจอไว้ (pointer-events: none)
# ถ้ารีบเปิดตัวถัดไปทันที มันจะไม่เปิดให้ — เจอกับการ์ด Fame Point ที่มีดรอปดาวน์ 2 อันติดกัน
JS_BUNDLE_POPUP_OPEN = """
() => {
  const n = document.querySelectorAll('[role="option"],[role="listbox"]').length;
  const locked = getComputedStyle(document.body).pointerEvents === 'none';
  return {open: n > 0, locked: locked, options: n};
}"""

JS_BUNDLE_ACT = """
([how, key, what, act, value, tiers, nth]) => {
  function vis(el){
    if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden' || s.opacity === '0') return false;
    const r = el.getBoundingClientRect();
    return r.width > 0 && r.height > 0;
  }
  function idAfter(t, prefix) {
    const i = t.indexOf(prefix);
    if (i < 0) return null;
    let j = i + prefix.length, out = '';
    while (j < t.length && t.charCodeAt(j) <= 32) j++;
    while (j < t.length && t.charAt(j) >= '0' && t.charAt(j) <= '9') { out += t.charAt(j); j++; }
    return out || null;
  }
  // ---- 1. หา "การ์ด" ----
  const need = (what === 'kind') ? 'ประเภท'
             : (what === 'tier') ? 'Tier' : 'จำนวน (Quantity)';
  //
  //  ไอเทมตัวเดียวกันใส่ซ้ำได้หลายใบในบันเดิลเดียว (เช่น โบนัสแคช 100 ใส่ 9 ใบ)
  //  ถ้าเจอ "ใบแรก" แล้วจบ จะไปตั้งค่าใบเดิมซ้ำๆ ใบที่ 6 เป็นต้นไปเลยไม่ได้ Tier
  //  -> ต้องบอกได้ว่าเอา "ใบที่เท่าไหร่" ของไอเทมนั้น
  //
  const cands = [];
  for (const el of document.querySelectorAll('div,section,li')) {
    const t = el.textContent || '';
    if (t.indexOf(need) < 0) continue;
    if (how === 'id') {
      if (t.indexOf('จำนวน (Quantity)') < 0) continue;
      if (idAfter(t, 'Item ID') !== String(key)) continue;
    } else {
      if (t.indexOf(String(key)) < 0) continue;
    }
    if (!vis(el)) continue;
    cands.push(el);
  }
  // เก็บเฉพาะใบในสุด (ตัดกล่องที่ครอบใบอื่นอยู่ทิ้ง) -> ได้ใบละ 1 ตัว เรียงตามหน้าจอ
  const cards = cands.filter(a => !cands.some(b => b !== a && a.contains(b)));
  const want = Math.max(1, parseInt(nth || 1, 10));
  const card = cards[want - 1];
  if (!card) return cards.length
    ? ('ไอเทม ' + key + ' มีแค่ ' + cards.length + ' ใบ แต่ขอใบที่ ' + want)
    : ('ไม่เจอการ์ด ' + key);

  // ---- 2. หา "ช่อง" ในการ์ดใบนั้น ----
  //
  //  ยึด "ป้ายชื่อช่อง" เป็นหลัก ไม่ใช่ "ตัวแรกที่เจอในการ์ด"
  //  เพราะการ์ด Fame Point มี 2 ช่องอยู่ในใบเดียวกัน (ประเภท กับ Tier)
  //  ถ้าเอาตัวแรก จะได้ช่อง "ประเภท" มาแทน "Tier" ทุกครั้ง
  //  -> ตั้งค่าไม่เข้า แล้วไปกดเปิดดรอปดาวน์ผิดตัว (ต้นเหตุที่ Fame Point ไม่ได้ Tier)
  function clean(s) {
    s = (s || '').split('*').join(' ');
    let out = '';
    for (let i = 0; i < s.length; i++) {
      if (s.charCodeAt(i) <= 32) {
        if (out && out.charAt(out.length - 1) !== ' ') out += ' ';
      } else out += s.charAt(i);
    }
    return out.trim();
  }
  function ownText(el) {
    let s = '';
    for (const n of el.childNodes) if (n.nodeType === 3) s += n.nodeValue;
    return clean(s);
  }
  function ctlsIn(el) {
    if (!el) return [];
    if (what === 'qty')
      return [...el.querySelectorAll('input')].filter(
        e => e.type !== 'checkbox' && e.type !== 'file' && vis(e));
    // select ที่ซ่อนอยู่ไม่เอา — Radix/shadcn แอบใส่ <select> ซ่อนไว้
    // ซึ่ง textContent ของมันมีชื่อ "ทุกตัวเลือก" อยู่ครบ ทำให้อ่านค่าปัจจุบันผิด
    const ss = [...el.querySelectorAll('select')].filter(vis);
    if (ss.length) return ss;
    return [...el.querySelectorAll('[role="combobox"],button')].filter(vis);
  }
  // เอาตัวที่ "อยู่ถัดจากป้าย" ก่อนเสมอ — ช่องของป้ายไหนก็อยู่ใต้ป้ายนั้น
  function pickNear(list, lab) {
    for (const c of list) if (lab.compareDocumentPosition(c) & 4) return c;
    return list.length ? list[list.length - 1] : null;
  }
  const wantLab = (what === 'kind') ? 'ประเภท'
                : (what === 'tier') ? 'Tier' : 'จำนวน (Quantity)';
  let ctl = null;
  const labs = [];
  for (const el of card.querySelectorAll('label,div,span,p,td,th'))
    if (ownText(el) === wantLab) labs.push(el);
  for (const lab of labs) {
    let up = lab;
    for (let d = 0; d < 5 && up; d++) {
      const c = pickNear(ctlsIn(up), lab);
      if (c) { ctl = c; break; }
      if (up === card) break;
      up = up.parentElement;
    }
    if (ctl) break;
  }
  if (!ctl) {                       // ไม่เจอป้าย — ใช้วิธีเดิมเป็นตัวสำรอง
    if (what === 'qty') {
      ctl = ctlsIn(card)[0];
    } else {
      ctl = [...card.querySelectorAll('select')].filter(vis)[0];
      if (!ctl) {
        const names = (what === 'tier')
          ? (tiers || []).concat(['เลือก Tier'])
          : null;
        ctl = [...card.querySelectorAll('[role="combobox"],button')].filter(e => {
          if (!vis(e)) return false;
          const t = (e.textContent || '').trim();
          if (names) return names.indexOf(t) >= 0;
          return t.indexOf('WALLET') >= 0 || t.indexOf('เลือก') === 0;
        })[0];
      }
    }
  }
  if (!ctl) return 'ไม่เจอช่อง ' + what + ' ในการ์ด ' + key;

  // อ่านค่าปัจจุบัน — select ต้องอ่านจาก .value ไม่ใช่ textContent
  function cur() {
    if (ctl.tagName === 'SELECT') {
      const o = ctl.options[ctl.selectedIndex];
      return o ? o.text.trim() : (ctl.value || '');
    }
    if (ctl.tagName === 'INPUT') return ctl.value;
    return (ctl.dataset.value || ctl.textContent || '').trim();
  }

  if (act === 'read') return 'ok|' + cur();
  // บอกว่าช่องที่เจอเป็นชนิดไหน — ช่องแบบ <select> ห้ามกดเปิด
  // เพราะหน้าต่างตัวเลือกของมันเป็นของเบราว์เซอร์ ไม่ได้อยู่ในหน้าเว็บ
  // กดแล้วจะค้างและมองไม่เห็นตัวเลือกใดๆ เลย
  if (act === 'tag') return 'ok|' + ctl.tagName;

  if (act === 'set') {
    if (ctl.tagName === 'SELECT') {
      const opts = [...ctl.options];
      let j = opts.findIndex(o => o.text.trim() === String(value));
      if (j < 0) j = opts.findIndex(o => (o.text || '').indexOf(String(value)) >= 0
                                      || (o.value || '').indexOf(String(value)) >= 0);
      if (j < 0) return 'ไม่มีตัวเลือก ' + value + ' (ช่องนี้มี: '
                        + opts.map(o => (o.text || '').trim()).slice(0, 8).join(', ') + ')';
      Object.getOwnPropertyDescriptor(HTMLSelectElement.prototype, 'value')
            .set.call(ctl, opts[j].value);
    } else if (ctl.tagName === 'INPUT') {
      ctl.scrollIntoView({block: 'center'});
      ctl.focus();
      Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value')
            .set.call(ctl, String(value));
    } else {
      return 'combobox';
    }
    ctl.dispatchEvent(new Event('input', {bubbles: true}));
    ctl.dispatchEvent(new Event('change', {bubbles: true}));
    return 'ok';
  }

  if (act === 'click') {
    ctl.scrollIntoView({block: 'center'});
    ctl.click();
    return 'ok|' + cur();
  }
  if (act === 'press') {
    // ดรอปดาวน์แบบ shadcn/Radix เปิดตอน "กดเมาส์ลง" ไม่ใช่ตอน click
    // สั่ง .click() เฉยๆ เลยไม่เปิดให้
    ctl.scrollIntoView({block: 'center'});
    const r = ctl.getBoundingClientRect();
    const o = {bubbles: true, cancelable: true, composed: true,
               clientX: r.left + r.width / 2, clientY: r.top + r.height / 2,
               button: 0, buttons: 1, pointerId: 1, pointerType: 'mouse', isPrimary: true};
    try { ctl.dispatchEvent(new PointerEvent('pointerdown', o)); } catch (e) {}
    ctl.dispatchEvent(new MouseEvent('mousedown', o));
    try { ctl.dispatchEvent(new PointerEvent('pointerup', o)); } catch (e) {}
    ctl.dispatchEvent(new MouseEvent('mouseup', o));
    ctl.dispatchEvent(new MouseEvent('click', o));
    return 'ok|' + cur();
  }
  if (act === 'point') {
    // คืนพิกัดไว้กดด้วย "เมาส์จริง" — ดรอปดาวน์บางตัวเปิดเฉพาะตอนกดจริงเท่านั้น
    ctl.scrollIntoView({block: 'center'});
    const r = ctl.getBoundingClientRect();
    return 'ok|' + Math.round(r.left + r.width / 2) + ',' + Math.round(r.top + r.height / 2);
  }
  return 'ไม่รู้จักคำสั่ง ' + act;
}"""

JS_BUNDLE_SET_MARKED = """
([what, val]) => {
  const el = document.querySelector('[data-trpu="' + what + '"]');
  if (!el) return 'ไม่เจอช่องที่ทำเครื่องหมายไว้';
  el.scrollIntoView({block: 'center'});
  el.focus();
  if (el.tagName === 'SELECT') {
    const opts = [...el.options];
    const j = opts.findIndex(o => o.text.trim() === String(val));
    if (j < 0) return 'ไม่มีตัวเลือก ' + val;
    Object.getOwnPropertyDescriptor(HTMLSelectElement.prototype, 'value').set.call(el, opts[j].value);
  } else if (el.tagName === 'INPUT') {
    Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value').set.call(el, String(val));
  } else {
    return 'combobox';
  }
  el.dispatchEvent(new Event('input', {bubbles: true}));
  el.dispatchEvent(new Event('change', {bubbles: true}));
  return 'ok';
}"""

JS_BUNDLE_GET_MARKED = """
([what]) => {
  const el = document.querySelector('[data-trpu="' + what + '"]');
  if (!el) return null;
  if (el.tagName === 'SELECT' || el.tagName === 'INPUT') return el.value;
  return (el.dataset.value || el.textContent || '').trim();
}"""

# --- ดรอปดาวน์ในกล่องเพิ่มของ (Credit / Player Exp) — ต้องอยู่ในกล่องเท่านั้น ---
# ห้ามใช้ .last ทั้งหน้า เพราะจะไปโดน "ประเภท Bundle" ที่แผงขวา แล้วเปลี่ยนเป็น CHOICE
JS_BUNDLE_MARK_ADDBOX = """
([what]) => {
""" + _JS_BOX + """
  const box = addBox();
  if (!box) return 'ไม่เจอกล่องเพิ่มของเข้า Bundle';
  let ctl = null;
  if (what === 'pick') {
    ctl = [...box.querySelectorAll('[role="combobox"],select,button')].filter(e => {
      if (!vis(e)) return false;
      const t = (e.textContent || '').trim();
      return t.indexOf('เลือก') === 0 || e.getAttribute('role') === 'combobox'
             || e.tagName === 'SELECT';
    })[0];
  } else if (what === 'wqty') {
    ctl = [...box.querySelectorAll('input')].filter(
      e => vis(e) && e.type !== 'checkbox' && e.type !== 'file' &&
           e.getBoundingClientRect().width < 200).pop();
  } else if (what === 'wadd') {
    ctl = [...box.querySelectorAll('button')].filter(
      e => vis(e) && (e.textContent || '').indexOf('เพิ่ม') >= 0).pop();
  }
  if (!ctl) return 'ไม่เจอ ' + what + ' ในกล่องเพิ่มของ';
  document.querySelectorAll('[data-trpu]').forEach(e => e.removeAttribute('data-trpu'));
  ctl.setAttribute('data-trpu', what);
  ctl.scrollIntoView({block: 'center'});
  return 'ok|' + (ctl.textContent || ctl.value || '').trim().slice(0, 40);
}"""

# อ่าน "ประเภท Bundle" ไว้ตรวจว่าไม่โดนเปลี่ยนโดยไม่ตั้งใจ
JS_BUNDLE_TYPE_GET = """
() => {
  const nodes = [...document.querySelectorAll('*')].filter(
    e => e.children.length === 0 && (e.textContent || '').trim().indexOf('ประเภท Bundle') === 0);
  for (const lb of nodes) {
    let n = lb.parentElement;
    for (let i = 0; i < 4 && n; i++) {
      const c = n.querySelector('select,[role="combobox"],button');
      if (c) return (c.value || c.textContent || '').trim();
      n = n.parentElement;
    }
  }
  return null;
}"""

# นับการ์ดในรายการไอเท็ม
JS_BUNDLE_COUNT = """
() => {
  const t = document.body.innerText || '';
  const i = t.indexOf('รายการไอเท็ม');
  if (i < 0) return -1;
  const j = t.indexOf('(', i);
  if (j < 0) return -1;
  let k = j + 1, out = '';
  while (k < t.length && t.charAt(k) >= '0' && t.charAt(k) <= '9') { out += t.charAt(k); k++; }
  return out ? parseInt(out, 10) : -1;
}"""

JS_BUNDLE_ROW_QTY = """
([idx, val]) => {
  const boxes = [];
  for (const el of document.querySelectorAll('*')) {
    if (el.children.length) continue;
    const t = (el.textContent || '').trim();
    if (!t.startsWith('จำนวน (Quantity)')) continue;
    let n = el.parentElement;
    for (let i = 0; i < 5 && n; i++) {
      const inp = n.querySelector('input:not([type="checkbox"]):not([type="file"])');
      if (inp) {
        const st = getComputedStyle(inp);
        const rc = inp.getBoundingClientRect();
        if (st.display !== 'none' && st.visibility !== 'hidden' && rc.height > 0)
          boxes.push(inp);
        break;
      }
      n = n.parentElement;
    }
  }
  const el = boxes[idx - 1];
  if (!el) return 'ไม่เจอช่องจำนวน (' + boxes.length + ' ใบ)';
  el.scrollIntoView({block: 'center'});
  const d = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value');
  d.set.call(el, String(val));
  el.dispatchEvent(new Event('input', {bubbles: true}));
  el.dispatchEvent(new Event('change', {bubbles: true}));
  return 'ok';
}"""

# หา "การ์ดใบที่ N" แล้วติดป้าย data-trpu ไว้ที่ตัวเลือก Tier ของใบนั้น
# ห้ามนับ [role=combobox] ทั้งหน้า เพราะจะรวม "ประเภท Bundle" ที่แผงขวาเข้าไปด้วย
# แล้วเลื่อนผิดใบทั้งแถว (เคยพลาดมาแล้ว)
JS_BUNDLE_MARK_TIER = """
([idx, tiers]) => {
  function vis(el){
    if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden') return false;
    const r = el.getBoundingClientRect();
    return r.width > 0 && r.height > 0;
  }
  // การ์ด = กล่องที่มีทั้งช่อง "จำนวน (Quantity)" และคำว่า "Tier"
  const cards = [];
  for (const el of document.querySelectorAll('*')) {
    if (el.children.length) continue;
    if ((el.textContent || '').trim().indexOf('จำนวน (Quantity)') !== 0) continue;
    let n = el.parentElement;
    for (let i = 0; i < 6 && n; i++) {
      if ((n.textContent || '').indexOf('Tier') >= 0 && n.querySelector('input') && vis(n)) {
        if (cards.indexOf(n) < 0) cards.push(n);
        break;
      }
      n = n.parentElement;
    }
  }
  const card = cards[idx - 1];
  if (!card) return 'ไม่เจอการ์ดใบที่ ' + idx + ' (เจอ ' + cards.length + ' ใบ)';

  // ในการ์ดใบนั้น หาตัวเลือก Tier
  let ctl = card.querySelector('select');
  if (!ctl) {
    const names = (tiers || []).concat(['เลือก Tier']);
    const cands = [...card.querySelectorAll('[role="combobox"],button')].filter(e => {
      if (!vis(e)) return false;
      const t = (e.textContent || '').trim();
      return names.indexOf(t) >= 0;
    });
    ctl = cands[0];
  }
  if (!ctl) return 'ไม่เจอตัวเลือก Tier ในการ์ดใบที่ ' + idx;
  document.querySelectorAll('[data-trpu]').forEach(e => e.removeAttribute('data-trpu'));
  ctl.setAttribute('data-trpu', 'tier');
  ctl.scrollIntoView({block: 'center'});
  return 'ok|' + (ctl.textContent || '').trim();
}"""

# อ่านค่า Tier ปัจจุบันของการ์ดที่ติดป้ายไว้ — ใช้ตรวจว่าตั้งสำเร็จจริงไหม
JS_BUNDLE_READ_MARK = """
() => {
  const el = document.querySelector('[data-trpu="tier"]');
  if (!el) return null;
  if (el.tagName === 'SELECT') return el.value;
  return (el.dataset.value || el.textContent || '').trim();
}"""

# หาการ์ดของ Fame Point แล้วติดป้ายไว้ที่ดรอปดาวน์ "ประเภท" ของการ์ดนั้น
JS_BUNDLE_MARK_KIND = """
([name]) => {
  function vis(el){
    if (!el) return false;
    const s = getComputedStyle(el);
    if (s.display === 'none' || s.visibility === 'hidden') return false;
    const r = el.getBoundingClientRect();
    return r.width > 0 && r.height > 0;
  }
  let card = null, bl = 1e9;
  for (const el of document.querySelectorAll('div,section')) {
    const t = el.textContent || '';
    if (t.indexOf(name) < 0) continue;
    if (t.indexOf('ประเภท') < 0) continue;
    if (!vis(el)) continue;
    if (t.length < bl) { bl = t.length; card = el; }
  }
  if (!card) return 'ไม่เจอการ์ด ' + name;
  let ctl = card.querySelector('select');
  if (!ctl) {
    ctl = [...card.querySelectorAll('[role="combobox"],button')].filter(
      e => vis(e) && (e.textContent || '').indexOf('WALLET') >= 0)[0];
  }
  if (!ctl) {
    ctl = [...card.querySelectorAll('[role="combobox"]')].filter(vis)[0];
  }
  if (!ctl) return 'ไม่เจอดรอปดาวน์ประเภทในการ์ด ' + name;
  document.querySelectorAll('[data-trpu]').forEach(e => e.removeAttribute('data-trpu'));
  ctl.setAttribute('data-trpu', 'kind');
  ctl.scrollIntoView({block: 'center'});
  return 'ok|' + (ctl.textContent || '').trim();
}"""

JS_BUNDLE_ROW_TIER = """
([idx, tier]) => {
  // เฉพาะกรณีที่เว็บใช้ <select> ธรรมดา (เผื่อวันหลังเปลี่ยน) — ไม่ใช่ก็ให้ฝั่ง python คลิกเอง
  const el = document.querySelector('[data-trpu="tier"]');
  if (!el || el.tagName !== 'SELECT') return 'combobox';
  const opts = [...el.options];
  const j = opts.findIndex(o => o.text.trim() === String(tier));
  if (j < 0) return 'ไม่มีตัวเลือก ' + tier;
  const d = Object.getOwnPropertyDescriptor(HTMLSelectElement.prototype, 'value');
  d.set.call(el, opts[j].value);
  el.dispatchEvent(new Event('input', {bubbles: true}));
  el.dispatchEvent(new Event('change', {bubbles: true}));
  return 'ok';
}"""

JS_BUNDLE_FAME_TYPE = """
([idx, want]) => {
  // การ์ด Fame Point มีดรอปดาวน์ "ประเภท" — ถ้าเป็น <select> ตั้งได้เลย
  for (const s of document.querySelectorAll('select')) {
    const opts = [...s.options];
    const j = opts.findIndex(o => (o.text || '').includes(want) ||
                                  (o.value || '').includes(want));
    if (j >= 0) {
      const d = Object.getOwnPropertyDescriptor(HTMLSelectElement.prototype, 'value');
      d.set.call(s, opts[j].value);
      s.dispatchEvent(new Event('input', {bubbles: true}));
      s.dispatchEvent(new Event('change', {bubbles: true}));
      return 'ok';
    }
  }
  return 'combobox';
}"""



# ============================================================================
#  [5.4] อ่านชีทต้นฉบับแบบเดียวกับเครื่องมือเดิม (parse_master_rows)
#        ไม่ต้องกรอก template เอง — ยกกฎมาจาก tr_studio.py ทั้งดุ้น
#
#        anchor  = คอลัมน์ที่หัวตารางเขียนว่า fdItemNum (หรือ Item Kind)
#        name    = Name / Display Name / Item Name  ที่อยู่ในช่วงของ anchor นั้น
#        qty     = เลขที่อยู่หน้าคำว่า "ชิ้น" ในชื่อ   เช่น "กล่องอิลลิเฟียร์ 5 ชิ้น" -> 5
#        ชื่อโชว์ = ตัด " X ชิ้น" ท้ายออก                 -> "กล่องอิลลิเฟียร์"
#        dur     = เลขในคอลัมน์ระยะเวลา เฉพาะตอนมีคำว่า "วัน" ; "ถาวร"/ว่าง -> '' (ถาวร)
#        trade   = คอลัมน์ Itemmove (Yes/No, ว่าง = No) ; ไม่มีคอลัมน์เลย -> 'any'
#        หนึ่งแถวมีได้หลายตาราง (fdItemNum หลายคอลัมน์วางเรียงกัน) และตัด ID ซ้ำให้
# ============================================================================
MASTER_TRADE_YES = ('yes', 'y', 'true', 'ได้', 'แลกเปลี่ยนได้')


def read_sheet_rows(wb, sheet):
    return [list(r) if r else []
            for r in wb[sheet].iter_rows(min_row=1, max_row=SCAN_ROWS, values_only=True)]


MAX_SANE_NAME = 70          # ชื่อไอเทมยาวเกินนี้ ปกติจะเป็นข้อความโน้ตในชีท ไม่ใช่ชื่อไอเทม
#   วัดจากไฟล์จริงแล้ว ชื่อที่ยาวที่สุด (ประกอบคำต่อท้าย+วัน แล้ว) = 52 ตัวอักษร
#   ตั้ง 70 เผื่อไว้ ของจริงจะไม่โดนเตือน แต่ประโยคโน้ตในชีทจะโดน


def kind_mode_len(rows):
    """ความยาวเลข ItemKind ที่พบบ่อยที่สุดในชุดนี้ — ใช้เป็นไม้บรรทัดของชุดนั้นเอง
    ไม่ได้ตั้งกฎตายตัวว่าต้องกี่หลัก เพราะแต่ละไฟล์อาจไม่เหมือนกัน"""
    counts = {}
    for r in rows:
        k = str(r.get('kind') or '')
        if k:
            counts[len(k)] = counts.get(len(k), 0) + 1
    if not counts:
        return 0
    return max(counts.items(), key=lambda kv: (kv[1], kv[0]))[0]


def suspect_reasons(r, mode_len=0):
    """บอกว่าแถวนี้ดู "แปลกปลอม" ตรงไหน — บอกเฉยๆ ไม่ตัดทิ้ง
    ชีทบางใบเขียนโน้ตปนมาในคอลัมน์ตาราง เลยโผล่มาเป็นไอเทมได้
    แต่บางทีก็เป็นของจริง เลยให้คนตัดสินเองเป็นเคสๆ ไป"""
    out = []
    k = str(r.get('kind') or '')
    if mode_len and k and len(k) <= mode_len - 2:
        out.append('เลข ItemKind สั้นกว่าตัวอื่นในชุดนี้มาก (%d หลัก · ส่วนใหญ่ %d หลัก)'
                   % (len(k), mode_len))
    nm = str(r.get('raw') or r.get('name') or r.get('disp') or '')
    if len(nm) > MAX_SANE_NAME:
        out.append('ชื่อยาวผิดปกติ (%d ตัวอักษร) อาจเป็นข้อความโน้ตในชีท' % len(nm))
    return out


def mark_suspects(rows):
    """ติดป้าย warn ให้ทุกแถว แล้วคืนจำนวนแถวที่น่าสงสัย (ไม่ลบแถวไหนทิ้งเลย)"""
    m = kind_mode_len(rows)
    n = 0
    for r in rows:
        w = suspect_reasons(r, m)
        r['warn'] = w
        if w:
            n += 1
    return n


def split_name_qty(raw_name):
    """แยก "ชื่อ" กับ "จำนวนชิ้น" ออกจากกัน — กฎเดียวกับเครื่องมือเดิมเป๊ะ

    ตัด "N ชิ้น" ตัวสุดท้ายออกจากชื่อ แล้วต่อส่วนหน้า+ส่วนหลังเข้าด้วยกัน
    -> จำนวนชิ้น "ไม่" ติดไปในชื่อ แต่ ชม./อย่างอื่น ยังอยู่ครบ
       'กล่อง 2 ชั่วโมง 3 ชิ้น' -> ชื่อ 'กล่อง 2 ชั่วโมง' · จำนวน 3
    ชื่อที่ไม่มีคำว่า 'ชิ้น' -> คงชื่อเดิม จำนวน 1
    คืนค่า (ชื่อ, จำนวน, มีคำว่าชิ้นไหม)
    """
    s = (raw_name or '').strip()
    mqs = list(re.finditer(r'(\d+)\s*ชิ้น', s))
    if not mqs:
        return s, '1', False
    m = mqs[-1]
    base = (s[:m.start()] + s[m.end():]).strip()
    base = re.sub(r'\s{2,}', ' ', base)
    return base, m.group(1), True


def check_item_image(path):
    """ตรวจไฟล์รูปก่อนจะเอาไปอัปโหลด — คืน (ผ่านไหม, เหตุผล)

    ยึดตามที่เว็บเขียนไว้เอง: .png/.jpg/.webp ขนาดไม่เกิน 5 MB
    """
    p = str(path or '').strip()
    if not p:
        return False, 'ยังไม่ได้เลือกรูป'
    if not os.path.isfile(p):
        return False, 'ไม่เจอไฟล์รูปในเครื่อง: ' + p
    ext = os.path.splitext(p)[1].lower()
    if ext not in IMG_EXTS:
        return False, ('นามสกุล %s ใช้ไม่ได้ — เว็บรับแค่ %s'
                       % (ext or '(ไม่มี)', ' / '.join(IMG_EXTS)))
    try:
        mb = os.path.getsize(p) / (1024.0 * 1024.0)
    except OSError as ex:
        return False, 'อ่านไฟล์รูปไม่ได้: ' + str(ex)
    if mb > IMG_MAX_MB:
        return False, 'ไฟล์ใหญ่ %.1f MB — เว็บรับไม่เกิน %d MB' % (mb, IMG_MAX_MB)
    return True, '%s (%.0f KB)' % (os.path.basename(p), mb * 1024)


def make_item_name(cname, suffix='', dur=''):
    """ประกอบชื่อที่จะกรอกลงเว็บ: <ชื่อ> <คำต่อท้าย> (<X> วัน)

    กฎเดิม: คำต่อท้ายกับ '(X วัน)' เติมเฉพาะชื่อที่เคยมี 'ชิ้น' เท่านั้น
    (ตัวเรียกเป็นคนตัดสินใจ — ฟังก์ชันนี้ต่อให้ตามที่ส่งมา)
    """
    s = (cname or '').strip()
    suffix = (suffix or '').strip()
    dur = str(dur or '').strip()
    if suffix:
        s = (s + ' ' + suffix).strip()
    if dur:
        s = (s + ' (%s วัน)' % dur).strip()
    return s


def row_to_item(r, type_name=DEFAULT_TYPE, suffix=DEFAULT_SUFFIX, price='', mail=''):
    """แปลงแถวที่อ่านจากชีท -> ข้อมูลไอเทมสำหรับหน้า create"""
    cname = r.get('cname') or r.get('disp') or ''
    if r.get('has_qty'):
        name = make_item_name(cname, suffix, r.get('dur', ''))
    else:
        name = cname          # ชื่อไม่มี 'ชิ้น' -> ไม่เติมอะไรเลย
    return {
        'kind': r.get('kind', ''),
        'name': name,
        'cname': cname,
        'desc': r.get('desc', ''),
        'type': type_name,
        'price': str(price or ''),
        'dur': r.get('dur', ''),
        'mail': str(mail or ''),
        'qty': r.get('amount', '1'),
        'trade': r.get('trade', 'any') == 'yes',
        'web': True,
        'img': '',            # รูปไอเทม — ไฟล์ต้นฉบับไม่มีรูป ต้องเลือกเองในหน้าแก้ไข
    }


def miss_to_item(r, type_name=DEFAULT_TYPE, suffix=DEFAULT_SUFFIX, price='', mail=''):
    """แถวผลค้นหาที่ "ไม่เจอ" -> ข้อมูลไอเทมสำหรับคิวสร้าง

    ใช้กฎเดิมของการนำเข้าไฟล์ทุกอย่าง (row_to_item) — ไม่คิดกฎใหม่เอง
    เพราะแถวที่ไม่เจอก็มาจากแถวในชีทใบเดียวกันนั่นแหละ
    ถ้าเป็นการค้นทีละตัว (ไม่มีต้นทางจากชีท) ก็ใช้ชื่อกับเลขที่เห็นบนตารางแทน
    """
    src = dict(r.get('src') or {})
    if not (src.get('cname') or src.get('disp')):
        src['cname'] = r.get('name') or ''
        src.setdefault('has_qty', False)   # ไม่รู้ว่าชื่อเดิมมี "ชิ้น" ไหม -> ไม่เติมอะไร
    if not src.get('kind'):
        src['kind'] = r.get('kind') or r.get('req') or ''
    return row_to_item(src, type_name, suffix, price, mail)


def parse_master_rows(rows, seen=None):
    out = []
    if seen is None:
        seen = set()
    tables = []
    prev = []

    for row in rows:
        cells = [num_str(c) for c in row]
        low = [c.lower() for c in cells]

        anchors = [j for j, c in enumerate(low) if 'fditemnum' in c]
        if not anchors:
            anchors = [j for j, c in enumerate(low) if c in ('item kind', 'itemkind')]

        if anchors:
            top = [str(c or '').strip().lower() for c in prev]
            tables = []
            for idx, a in enumerate(anchors):
                end_c = anchors[idx + 1] if idx + 1 < len(anchors) else max(len(low), a + 12)
                span = list(range(a, min(end_c, len(low))))

                def findin(preds, _span=span, _low=low):
                    for j in _span:
                        h = _low[j]
                        if any(pr(h) for pr in preds):
                            return j
                    return None

                name_col = findin([lambda h: h == 'name', lambda h: 'display name' in h,
                                   lambda h: h == 'item name', lambda h: h == 'itemname'])
                desc_col = findin([lambda h: 'item des' in h, lambda h: h == 'des',
                                   lambda h: 'description' in h, lambda h: 'รายละเอียด' in h,
                                   lambda h: 'item dec' in h])
                dur_col = findin([lambda h: 'ระยะเวลา' in h, lambda h: 'ของขวัญ' in h,
                                  lambda h: 'duration' in h])
                move_col = findin([lambda h: 'itemmove' in h])
                if move_col is None:      # บางไฟล์เขียน Itemmove ไว้แถวบนของหัวตาราง
                    for j in span:
                        if j < len(top) and 'itemmove' in top[j]:
                            move_col = j
                            break
                tables.append({'kind': a, 'name': name_col, 'dur': dur_col,
                               'move': move_col, 'desc': desc_col})
            prev = cells
            continue

        prev = cells
        if not tables:
            continue

        for t in tables:
            a = t['kind']
            if a >= len(cells):
                continue
            kind = num_str(cells[a])
            if not re.fullmatch(r'\d+', kind):
                continue
            nc = t['name'] if t['name'] is not None else (a + 4)
            # ตัดอีโมจิออกตั้งแต่ตรงนี้ ทุกอย่างที่แตกออกไปจะสะอาดตามหมด
            name = clean_text(cells[nc] if nc < len(cells) else '')
            # ต้องมีตัวอักษรจริงในชื่อ ไม่ใช่ตัวเลขล้วน (กันแถวยอดรวม)
            if not name.strip() or not re.search(r'[^\d.,\s]', name):
                continue
            if kind in seen:
                continue
            seen.add(kind)

            m = re.search(r'(\d+)\s*ชิ้น', name)
            qty = m.group(1) if m else ''
            disp = re.sub(r'\s*\d+\s*ชิ้น.*$', '', name).strip() or name.strip()

            dc = t['dur'] if t['dur'] is not None else (a + 5)
            durc = cells[dc] if dc < len(cells) else ''
            md = re.search(r'(\d+)', durc) if 'วัน' in durc else None
            dur = md.group(1) if md else ''

            if t['move'] is not None:
                mc = cells[t['move']] if t['move'] < len(cells) else ''
                trade = 'yes' if mc.strip().lower() in MASTER_TRADE_YES else 'no'
            else:
                trade = 'any'

            dsc = ''
            if t.get('desc') is not None and t['desc'] < len(cells):
                dsc = clean_text(cells[t['desc']])

            cname, amount, has_qty = split_name_qty(name)

            out.append({'kind': kind, 'name': '', 'disp': disp,
                        'dur': dur, 'trade': trade, 'qty': qty,
                        'has_move': t['move'] is not None,
                        # ---- ใช้ตอน "สร้างไอเทม" ----
                        'raw': name.strip(), 'cname': cname,
                        'amount': amount, 'has_qty': has_qty, 'desc': dsc})
    return out



# ============================================================================
#  [5.8] อ่านไฟล์ต้นฉบับ -> บันเดิล
#        ยกกฎมาจากเครื่องมือเดิม (parse_source_sheet ใน tr_studio.py) ทั้งดุ้น
#
#        หัวตาราง = แถวที่มีคำว่า fdItemNum
#        Aztek Item Id = คอลัมน์ "ตัวเลขที่ไม่มีหัวตาราง" ที่อยู่ในบล็อกเดียวกัน
#                        (เลขซ้ายมือของตาราง — ตัวที่เอาไปค้นบนเว็บ)
#        เลข Bundle    = เลขในคอลัมน์เดียวกันนั้น ที่อยู่เหนือหัวตาราง 1-2 แถว
#        ชื่อบันเดิล    = Product Name > แถวหัวข้อเดี่ยวๆ > หัวคอลัมน์ Id > ชื่อชีท#n
#        จำนวน         = Amt / Amount / จำนวน
#        Tier          = Rank / ยศ / Tier
#        famepoint/exp = กล่อง "ได้รับ famepoint กับ exp" เหนือหัวตาราง
# ============================================================================
_SRC_TIER = {t.lower(): t for t in TIERS}
_SRC_STOP = {'start', 'end', 'limit', 'reset', 'category', 'product name', 'bundle',
             'date', 'announce', 'channel', 'detail', 'reward', 'fditemnum'}
_SRC_KNOWN = ('fditemnum', 'fdposition', 'fditemkind', 'rank', 'display name', 'item des',
              'status', 'ระยะเวลา', 'ของขวัญ', 'amount', 'amt', 'จำนวน', 'ราคา', 'price',
              'thb', 'รวม', 'duration')


def src_int(v):
    """อ่านจำนวนเต็มจากเซลล์ ('2074' / 2074.0 -> '2074') ไม่ใช่ตัวเลข -> None"""
    s = ('' if v is None else str(v)).strip()
    m = re.fullmatch(r'(\d+)(?:\.0+)?', s)
    return m.group(1) if m else None


def bundle_rewards(rows, hr):
    """จับกล่อง 'ได้รับ famepoint กับ exp' ที่อยู่เหนือหัวตาราง
    famepoint -> Credit (ต้องตั้งเป็นเครดิตเรียลไทม์)  ·  exp -> Player Experience

    ตัวเลขต้องเป็น "เซลล์ที่อยู่ใต้ป้ายนั้นตรงๆ" เท่านั้น
    ห้ามกวาดหาเลขไปทางขวาเรื่อยๆ เพราะข้างๆ กล่องมักมีคนพิมพ์เลขอื่นไว้
    (เช่น "โบนัสแคช 10%  900") ซึ่งไม่เกี่ยวกับ fame/exp เลย — เคยหยิบผิดมาแล้ว
    """
    region = rows[max(0, hr - 8):hr]
    out = []
    seen = set()

    def num_of(v):
        s = re.sub(r'[, ]', '', ('' if v is None else str(v)).strip())
        return re.sub(r'\.0+$', '', s) if re.fullmatch(r'\d+(\.\d+)?', s) else None

    for ri, row in enumerate(region):
        cs = [('' if c is None else str(c)).strip() for c in (row or [])]
        for j, c in enumerate(cs):
            cl = c.lower()
            if 'fame' not in cl:
                continue
            # ไม่ใช่ป้าย แต่เป็นโน้ตบอกวิธีตั้งค่า เช่น "Famepoint เซ็ตเป็น Wallet_realtime_credit"
            if 'wallet' in cl or 'เซ็ต' in c:
                continue
            # ใต้ป้ายเท่านั้น (เผื่อป้ายสูง 2 แถว ก็ไล่ลงไปได้ 2 แถว)
            val = None
            for d in (1, 2):
                if ri + d >= len(region):
                    break
                below = region[ri + d] or []
                val = num_of(below[j] if j < len(below) else None)
                if val:
                    break
            # ไม่มีข้างล่างจริงๆ ค่อยดูช่องที่ "ติดกันทางขวา" และต้องติดกันเท่านั้น
            if val is None:
                for k in (j + 1, j + 2):
                    if k < len(cs) and cs[k]:
                        val = num_of(cs[k])
                        break
            if not val:
                continue
            if 'credit' not in seen:
                out.append({'type': 'CREDIT', 'value': FAME_NAME, 'qty': val})
                seen.add('credit')
            if 'exp' in cl and 'exp' not in seen:
                out.append({'type': 'PLAYER_EXP', 'value': 'Player Experience', 'qty': val})
                seen.add('exp')
    return out


def parse_bundle_sheet(rows, sheet_name):
    """คืน (bundles, warnings) — แต่ละ bundle พร้อมเอาไปกรอกหน้าเว็บได้เลย"""
    # ชื่อชีทเองก็มีอีโมจิ (✅ 💛) และติดไปกับชื่อบันเดิลสำรอง/คำเตือน — ล้างตั้งแต่ต้นทาง
    sheet_name = clean_text(sheet_name) or str(sheet_name or '')
    bundles = []
    warns = []
    # หัวเรื่องบนสุดของชีท เช่น "Master Code 01/10/2026"
    # ใช้เฉพาะกับชีทแพทเทิร์น CODE เท่านั้น (ดูที่ _bundle_code ด้านล่าง)
    top_title = ''
    for r in rows[:4]:
        stop_row = False
        for c in (r or []):
            t = clean_text(c)
            if t.lower() == 'code':        # ถึงแถวโค้ดแล้ว = หมดเขตหัวเรื่อง เลิกหา
                stop_row = True
                break
            if not t or src_int(t) is not None or len(t) > 60:
                continue
            if re.match(r'^\d{4}-\d{2}-\d{2}', t) or re.match(r'^\d{1,2}[.:]\d{2}$', t):
                continue                   # วันที่/เวลา ไม่ใช่หัวเรื่อง
            if t.startswith('*'):          # บรรทัดหมายเหตุ ไม่ใช่หัวเรื่อง
                continue
            if t.lower() in _SRC_STOP:
                continue
            top_title = t
            break
        if stop_row or top_title:
            break
    if not top_title:                      # ชีทไหนไม่มีหัวเรื่อง ใช้ชื่อชีทแทน
        top_title = sheet_name
    hdr = [i for i, row in enumerate(rows)
           if any('fditemnum' in str(c).lower() for c in (row or []))]
    for bi, hr in enumerate(hdr):
        header = [('' if c is None else str(c)).strip() for c in rows[hr]]
        low = [h.lower() for h in header]

        def col(preds, _low=low):
            for j, h in enumerate(_low):
                if any(pr(h) for pr in preds):
                    return j
            return None
        kindc = col([lambda h: 'fditemnum' in h])
        rankc = col([lambda h: h == 'rank', lambda h: h == 'ยศ', lambda h: h == 'tier'])
        amtc = col([lambda h: h == 'amount', lambda h: h == 'amt', lambda h: h == 'จำนวน'])
        posc = col([lambda h: 'fdposition' in h])
        ikc = col([lambda h: 'fditemkind' in h])
        namec = col([lambda h: 'display name' in h, lambda h: h == 'name'])

        end = hdr[bi + 1] if bi + 1 < len(hdr) else len(rows)
        drows = []
        for r in range(hr + 1, end):
            if kindc is None:
                break
            row = rows[r] or []
            kv = src_int(row[kindc]) if kindc < len(row) else None
            if kv is None:
                if drows:
                    break
                continue
            drows.append(r)
        if not drows:
            continue

        # คอลัมน์ Aztek Item Id
        # ชีทส่วนใหญ่ตั้งหัวคอลัมน์นี้ว่า "Bundle" (บางใบเขียน Aztek Item Id / Item Id)
        # ถ้ามีหัวตารางบอกไว้ ให้เชื่อหัวตารางก่อนเสมอ — เดาเองเมื่อจำเป็นเท่านั้น
        idc_named = col([lambda h: h == 'bundle', lambda h: 'aztek' in h,
                         lambda h: h in ('item id', 'itemid', 'bundle id', 'bundleid')])

        # เดาเอง: คอลัมน์ตัวเลขที่ "ไม่มีหัวตารางที่รู้จัก"
        labeled = set(j for j, h in enumerate(low) if any(k in h for k in _SRC_KNOWN))
        for c in (kindc, posc, ikc, amtc, rankc, namec):
            if c is not None:
                labeled.add(c)
        ncols = max(len(rows[r] or []) for r in drows)
        cand = [c for c in range(ncols) if c not in labeled and
                sum(1 for r in drows
                    if c < len(rows[r] or []) and src_int((rows[r] or [])[c]) is not None)
                >= max(1, int(len(drows) * 0.6))]

        idcol = bid = None
        if idc_named is not None and any(
                idc_named < len(rows[r] or []) and
                src_int((rows[r] or [])[idc_named]) is not None for r in drows):
            idcol = idc_named
        else:
            # เลข Bundle: ดูแค่ 1-2 แถวเหนือหัวตาราง (กันไปหยิบเลขของบล็อกก่อนหน้า)
            above = hdr[bi - 1] + 1 if bi > 0 else 0
            bid_top = max(hr - 2, above)
            for c in cand:
                for r in range(hr - 1, bid_top - 1, -1):
                    row = rows[r] or []
                    b = src_int(row[c]) if c < len(row) else None
                    if b is not None:
                        idcol, bid = c, b
                        break
                if idcol is not None:
                    break
            if idcol is None and cand:
                # เอาคอลัมน์ที่ "อยู่ติดกับตารางไอเทม" ที่สุด ไม่ใช่ตัวซ้ายสุดของทั้งแผ่น
                # (ชีทที่มีตารางรายชื่อคนอยู่ข้างๆ จะมีคอลัมน์ลำดับ 1,2,3 หลอกอยู่)
                anchor = kindc if kindc is not None else 0
                idcol = min(cand, key=lambda c: (abs(c - anchor), c))
        if idcol is None:
            warns.append('%s บล็อก %d: หาคอลัมน์ Aztek Item Id ไม่เจอ -> ข้าม'
                         % (sheet_name, bi + 1))
            continue

        # ช่วงคอลัมน์ของตารางนี้ — ใช้กันไปหยิบของบล็อกที่วางคู่กันอยู่คนละฝั่ง
        band = [c for c in (kindc, posc, ikc, amtc, rankc, namec, idcol)
                if c is not None]
        bl, br = min(band), max(band)

        # ชื่อบันเดิล
        name = None

        # (0) ชีทแพทเทิร์น "Master Code": เหนือหัวตารางมีแถวที่เขียนว่า CODE เดี่ยวๆ
        #     แล้วช่องถัดไปทางขวาคือโค้ด เช่น  CODE | MARS8X3P9V2K
        #     -> ตั้งชื่อเป็น "<โค้ด> <หัวเรื่องบนสุดของชีท>"
        #     ยิงเฉพาะชีทที่มีคำว่า CODE เดี่ยวๆ จริงๆ เท่านั้น
        #     ("CODE START DATE" / "CODE: TALES RUNNER" / "Item Code" ไม่เข้าเงื่อนไข)
        #     และต้องอยู่ในช่วงคอลัมน์ของตารางนี้ ไม่ใช่ของบล็อกที่วางคู่กันอยู่
        for r in range(hr - 1, max(hr - 8, 0) - 1, -1):
            rl = [clean_text(c) for c in (rows[r] or [])]
            for j in range(bl, min(br + 1, len(rl))):
                if rl[j].lower() != 'code':
                    continue
                for k in range(j + 1, min(br + 2, len(rl))):
                    if rl[k] and src_int(rl[k]) is None:
                        name = (rl[k] + ' ' + top_title).strip()
                        break
                break
            if name:
                break

        for r in range(hr - 1, max(hr - 14, 0) - 1, -1):
            if name:
                break
            rl = [('' if c is None else str(c)).strip() for c in (rows[r] or [])]
            for j, cell in enumerate(rl):
                if 'product name' in cell.lower():
                    for k in range(j + 1, len(rl)):
                        if rl[k]:
                            name = clean_text(rl[k])
                            break
                    break
            if name:
                break
        if not name:
            # หัวข้อของบล็อกอยู่เหนือหัวตาราง "ในช่วงคอลัมน์เดียวกับตาราง"
            # ต้องดูเฉพาะช่วงคอลัมน์นั้น เพราะบางชีทมีตารางรายชื่อคนอยู่ข้างๆ
            # เรียงกันแบบ  ชื่อเรื่อง -> คำอธิบาย -> หัวตาราง  เลยเอา "บรรทัดบนสุด"
            picks = []
            for r in range(hr - 1, max(hr - 7, 0) - 1, -1):
                rl = [('' if c is None else str(c)).strip() for c in (rows[r] or [])]
                seg = [rl[j] for j in range(bl, min(br + 1, len(rl))) if rl[j]]
                if not seg:
                    break
                nz = [t for t in seg
                      if src_int(t) is None and t.lower() not in _SRC_STOP]
                if len(nz) != 1:
                    break
                picks.append(clean_text(nz[0]))
            if picks:
                name = picks[-1]
        if not name:
            for r in range(hr - 1, max(hr - 4, 0) - 1, -1):
                texts = [('' if c is None else str(c)).strip() for c in (rows[r] or [])]
                nz = [t for t in texts if t and src_int(t) is None]
                if len(nz) == 1 and nz[0].lower() not in _SRC_STOP:
                    name = clean_text(nz[0])
                    break
        if not name:
            h0 = header[idcol] if idcol < len(header) else ''
            if h0 and src_int(h0) is None and h0.lower() not in _SRC_STOP:
                name = clean_text(h0)
        if not name:
            name = '%s #%d' % (sheet_name, bi + 1)

        items = []
        for r in drows:
            row = rows[r] or []
            iid = src_int(row[idcol]) if idcol < len(row) else None
            disp = ''
            if namec is not None and namec < len(row):
                disp = clean_text(row[namec])
            if iid is None:
                warns.append('%s "%s" แถว %d: ไม่มี Aztek Item Id -> ข้าม'
                             % (sheet_name, name, r + 1))
                continue
            qty = (src_int(row[amtc]) if (amtc is not None and amtc < len(row)) else None) or '1'
            tier = DEFAULT_TIER
            if rankc is not None and rankc < len(row):
                rv = ('' if row[rankc] is None else str(row[rankc])).strip()
                t = _SRC_TIER.get(rv.lower())
                if t:
                    tier = t
                elif rv and rv != '-':
                    warns.append('%s "%s" แถว %d: Rank "%s" ไม่รู้จัก -> %s'
                                 % (sheet_name, name, r + 1, rv, DEFAULT_TIER))
            ikind = ''
            if kindc is not None and kindc < len(row):
                ikind = src_int(row[kindc]) or ''
            items.append({'id': iid, 'qty': qty, 'tier': tier, 'disp': disp,
                          'kind': ikind})
        # กันพลาดซ้ำรอย: อ่านไอเทมได้ไม่ครบทุกแถวในตาราง = มีอะไรผิดแน่ๆ ต้องดังไว้ก่อน
        if len(items) < len(drows):
            warns.append('%s "%s": ตารางมี %d แถว แต่อ่านมาได้ %d — '
                         'ไปดูคอลัมน์ Aztek Item Id ในชีทด้วย'
                         % (sheet_name, name, len(drows), len(items)))
        # เลขเรียง 1,2,3... ทั้งชุด = น่าจะไปหยิบคอลัมน์ "ลำดับ" มาแทน Aztek Item Id
        if idc_named is None and len(items) >= 3:
            nums = [int(i['id']) for i in items]
            if nums == list(range(nums[0], nums[0] + len(nums))) and nums[0] <= len(items):
                warns.append('%s "%s": Aztek Item Id ที่ได้เป็นเลขเรียง %s — '
                             'อาจหยิบผิดคอลัมน์ ตรวจก่อนสร้าง'
                             % (sheet_name, name, ', '.join(str(n) for n in nums[:4])))
        if items:
            bundles.append({'name': name, 'type': DEFAULT_BUNDLE_TYPE, 'deliver': True,
                            'items': items, 'rewards': bundle_rewards(rows, hr),
                            'sheet': sheet_name, 'src_bundle_id': bid})
    return bundles, warns


def scan_bundle_sheets_wb(wb, progress=None):
    """คืน list ของ (ชื่อชีท, จำนวนบันเดิลที่เจอ)"""
    out = []
    names = wb.sheetnames
    for i, s in enumerate(names, 1):
        if progress:
            progress(i, len(names), s)
        try:
            rows = [list(r) if r else [] for r in
                    wb[s].iter_rows(min_row=1, max_row=SCAN_ROWS, values_only=True)]
            b, _ = parse_bundle_sheet(rows, s)
        except Exception:
            b = []
        out.append((s, len(b)))
    return out


# ============================================================================
#  [5.9] อ่านไฟล์ต้นฉบับ -> Item Code (แถบ WR Master)
#        ชีทแพทเทิร์น "Master Code" — แต่ละบล็อกมีหน้าตาแบบนี้
#            CODE | <โค้ด>            ... Start | <วันที่> | <เวลา>
#                                     ... End   | <วันที่> | <เวลา>
#            Bundle                   ... Limit การใช้งาน | <จำนวน หรือ ไม่จำกัด>
#            <เลข Bundle>
#        อ่านโดยอ้าง "ป้ายข้อความ" ทั้งหมด ไม่ยึดตำแหน่งคอลัมน์
# ============================================================================
_WR_MONTH = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun',
             'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')


def wr_date(v):
    """อ่านวันที่จากเซลล์ -> (ปี, เดือน, วัน) ไม่ใช่วันที่ -> None"""
    if isinstance(v, datetime):
        return (v.year, v.month, v.day)
    s = clean_text(v)
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m:
        return (int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$', s)
    if m:
        return (int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def wr_time(v, end=False):
    """อ่านเวลาจากเซลล์ '13.00' / '23.59' -> 'HH:MM:SS'

    ปลายทางเวลาสิ้นสุดของทีมคือ 23.59 = 23:59:59 (เอาให้ถึงวินาทีสุดท้าย)
    เวลาเริ่มใช้วินาทีเป็น 00
    """
    s = clean_text(v).replace('น.', '').strip()
    m = re.match(r'^(\d{1,2})[.:](\d{2})(?:[.:](\d{2}))?$', s)
    if not m:
        return None
    hh, mm = int(m.group(1)), int(m.group(2))
    ss = int(m.group(3)) if m.group(3) else (59 if end else 0)
    if hh > 23 or mm > 59:
        return None
    return '%02d:%02d:%02d' % (hh, mm, ss)


def wr_name_date(ymd):
    """(2026, 10, 1) -> '1 Oct 2026' — รูปแบบชื่อที่ทีมใช้"""
    if not ymd:
        return ''
    return '%d %s %d' % (ymd[2], _WR_MONTH[ymd[1] - 1], ymd[0])


def _wr_after(cells, j, want_num=False):
    """ค่าถัดไปทางขวาของช่อง j ที่ไม่ว่าง"""
    for k in range(j + 1, len(cells)):
        t = clean_text(cells[k])
        if not t:
            continue
        if want_num and src_int(t) is None:
            continue
        return cells[k]
    return None


def parse_code_sheet(rows, sheet_name):
    """คืน (codes, warnings) — แต่ละตัวพร้อมเอาไปกรอกหน้าสร้าง Item Code

    ต้องเป็นชีทที่มีช่องเขียนว่า CODE เดี่ยวๆ เท่านั้น ชีทแบบอื่นคืนลิสต์ว่าง
    """
    sheet_name = clean_text(sheet_name) or str(sheet_name or '')
    out, warns = [], []
    starts = []
    for i, row in enumerate(rows):
        for j, c in enumerate(row or []):
            if clean_text(c).lower() == 'code':
                starts.append((i, j))
                break
    for n, (r0, jc) in enumerate(starts):
        end_r = starts[n + 1][0] if n + 1 < len(starts) else len(rows)
        code = clean_text(_wr_after(rows[r0], jc))
        if not code:
            warns.append('%s แถว %d: มีป้าย CODE แต่ไม่มีโค้ด -> ข้าม'
                         % (sheet_name, r0 + 1))
            continue

        d = {'code': code, 'sheet': sheet_name, 'row': r0 + 1,
             'start': '', 'end': '', 'limit': '', 'bundle': '', 'per_user': '1'}
        # ---- Start / End / Limit / Bundle : หาโดยอ้างป้าย ในช่วงของบล็อกนี้ ----
        for r in range(r0, min(r0 + 6, end_r)):
            cells = list(rows[r] or [])
            for j, c in enumerate(cells):
                lab = clean_text(c).lower().rstrip(':')
                if lab in ('start', 'end'):
                    ymd = wr_date(_wr_after(cells, j))
                    tm = None
                    for k in range(j + 1, len(cells)):
                        tm = wr_time(cells[k], end=(lab == 'end'))
                        if tm:
                            break
                    if ymd:
                        d[lab] = '%04d-%02d-%02d %s' % (
                            ymd[0], ymd[1], ymd[2], tm or
                            ('23:59:59' if lab == 'end' else '00:00:00'))
                        if lab == 'start':
                            d['name'] = code + ' ' + wr_name_date(ymd)
                elif lab.startswith('limit'):
                    v = clean_text(_wr_after(cells, j))
                    d['limit'] = v
                elif lab == 'bundle':
                    # เลข Bundle อยู่ "ใต้ป้าย" ในคอลัมน์เดียวกัน
                    for r2 in range(r + 1, min(r + 3, end_r)):
                        rw = list(rows[r2] or [])
                        b = src_int(rw[j]) if j < len(rw) else None
                        if b:
                            d['bundle'] = b
                            break

        if not d.get('name'):
            d['name'] = (code + ' ' + sheet_name).strip()
        cap = src_int(d['limit'])
        d['cap'] = str(int(cap) + WR_SPARE) if cap else ''      # '' = ไม่จำกัด
        d['unlimited'] = not cap
        if not d['start'] or not d['end']:
            warns.append('%s "%s": ไม่เจอเวลาเริ่ม/สิ้นสุดในชีท'
                         % (sheet_name, code))
        if not d['bundle']:
            warns.append('%s "%s": ไม่เจอเลข Bundle ในชีท -> ต้องใส่เอง'
                         % (sheet_name, code))
        out.append(d)
    return out, warns


def read_codes_wb(wb, sheet):
    rows = [list(r) if r else [] for r in
            wb[sheet].iter_rows(min_row=1, max_row=SCAN_ROWS, values_only=True)]
    return parse_code_sheet(rows, sheet)


def scan_code_sheets_wb(wb, progress=None):
    """คืน list ของ (ชื่อชีท, จำนวนโค้ดที่เจอ) — ชีทที่ไม่ใช่แพทเทิร์น CODE ได้ 0"""
    out = []
    names = wb.sheetnames
    for i, s in enumerate(names, 1):
        if progress:
            progress(i, len(names), s)
        try:
            c, _ = read_codes_wb(wb, s)
        except Exception:
            c = []
        out.append((s, len(c)))
    return out


def read_bundles_wb(wb, sheet):
    rows = [list(r) if r else [] for r in
            wb[sheet].iter_rows(min_row=1, max_row=SCAN_ROWS, values_only=True)]
    return parse_bundle_sheet(rows, sheet)

# ============================================================================
#  [5.5] หน้าต่างนำเข้า Excel — เลือกชีท แล้วอ่านให้อัตโนมัติ
# ============================================================================
class ImportDialog:
    def __init__(self, parent, path):
        self.path = path
        self.result = None
        self.sheet_name = ''
        self.sheet_names = []
        self.sheets = []
        self.rows = []
        self._pending = None

        self.top = tk.Toplevel(parent)
        self.top.title('นำเข้าจาก Excel')
        self.top.configure(bg=C['bg'])
        self.top.geometry('1010x640')
        self.top.transient(parent)
        self.top.grab_set()

        head = tk.Frame(self.top, bg=C['card'], height=56)
        head.pack(fill='x')
        head.pack_propagate(False)
        tk.Label(head, text='นำเข้าจาก Excel', bg=C['card'], fg=C['fg'],
                 font=('Segoe UI', 13, 'bold')).pack(side='left', padx=18)
        tk.Label(head, text=os.path.basename(path), bg=C['card'], fg=C['dim'],
                 font=('Segoe UI', 9)).pack(side='left')

        foot = tk.Frame(self.top, bg=C['card'], height=62)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)
        self.count_lbl = tk.Label(foot, text='', bg=C['card'], fg=C['dim'], font=('Segoe UI', 10))
        self.count_lbl.pack(side='left', padx=18)
        self.ok_btn = tk.Button(foot, text='ใช้ข้อมูลนี้', bg=C['accent'], fg='white', bd=0,
                                font=('Segoe UI', 10, 'bold'), cursor='hand2',
                                activebackground='#3d55cf', activeforeground='white',
                                state='disabled', command=self._ok)
        self.ok_btn.pack(side='right', padx=18, pady=12, ipadx=22, ipady=5)
        tk.Button(foot, text='ยกเลิก', bg=C['input'], fg=C['fg'], bd=0,
                  font=('Segoe UI', 10), cursor='hand2', activebackground=C['line'],
                  command=self._cancel).pack(side='right', pady=12, ipadx=16, ipady=5)

        body = tk.Frame(self.top, bg=C['bg'])
        body.pack(fill='both', expand=True, padx=14, pady=12)

        # ---- ซ้าย : รายชื่อชีท ----
        left = tk.Frame(body, bg=C['bg'], width=280)
        left.pack(side='left', fill='y')
        left.pack_propagate(False)
        tk.Label(left, text='เลือกชีท', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(anchor='w')
        lbwrap = tk.Frame(left, bg=C['bg'])
        lbwrap.pack(fill='both', expand=True, pady=(5, 0))
        # selectmode='multiple' = คลิกทีละชีทเพื่อเลือก/ยกเลิก เลือกหลายชีทได้โดยไม่ต้องกด Ctrl
        self.lb = tk.Listbox(lbwrap, bg=C['input'], fg=C['fg'], bd=0,
                             highlightthickness=1, highlightbackground=C['line'],
                             selectbackground=C['accent'], selectforeground='white',
                             font=('Segoe UI', 9), activestyle='none',
                             selectmode='multiple')
        sb = ttk.Scrollbar(lbwrap, orient='vertical', command=self.lb.yview)
        self.lb.configure(yscrollcommand=sb.set)
        self.lb.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')
        self.lb.bind('<<ListboxSelect>>', lambda e: self._on_sheet())

        self.sel_lbl = tk.Label(left, text='', bg=C['bg'], fg=C['fg'],
                                font=('Segoe UI', 9, 'bold'), anchor='w')
        self.sel_lbl.pack(fill='x', pady=(8, 2))
        btnf = tk.Frame(left, bg=C['bg'])
        btnf.pack(fill='x')
        for txt, cmd in (('เลือกชีทที่มีตารางทั้งหมด', self._pick_all),
                         ('ล้างที่เลือก', self._pick_none)):
            tk.Button(btnf, text=txt, bg=C['input'], fg=C['fg'], bd=0,
                      font=('Segoe UI', 9), cursor='hand2', activebackground=C['line'],
                      command=cmd).pack(side='left', padx=(0, 6), ipadx=6, ipady=3)
        tk.Label(left, text='คลิกชีทเพื่อเลือก คลิกซ้ำเพื่อเอาออก — เลือกได้หลายชีทพร้อมกัน '
                            'ID ที่ซ้ำกันข้ามชีทจะถูกตัดให้เหลือตัวเดียว',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), anchor='w',
                 justify='left', wraplength=265).pack(fill='x', pady=(8, 0))

        # ---- ขวา ----
        right = tk.Frame(body, bg=C['bg'])
        right.pack(side='left', fill='both', expand=True, padx=(14, 0))

        self.info = tk.Label(right, text='กำลังสแกนไฟล์…', bg=C['bg'], fg=C['dim'],
                             font=('Segoe UI', 9), anchor='w', justify='left')
        self.info.pack(fill='x')

        tk.Label(right, text='อ่านให้อัตโนมัติแบบเดียวกับเครื่องมือเดิม — '
                             'qty ดึงจากเลขหน้าคำว่า “ชิ้น” ในชื่อ · ระยะเวลาอ่านจากคอลัมน์ที่มีคำว่า “วัน” · '
                             'แลกเปลี่ยนอ่านจากคอลัมน์ Itemmove (ไม่มีคอลัมน์ = ไม่กรอง)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), wraplength=690,
                 justify='left').pack(anchor='w', pady=(6, 0))

        self.warn_lbl = tk.Label(right, text='', bg=C['bg'], fg=C['warn'],
                                 font=('Segoe UI', 9), anchor='w', justify='left',
                                 wraplength=690)
        self.warn_lbl.pack(fill='x', pady=(6, 0))

        tk.Label(right, text='ข้อมูลที่จะนำเข้า', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(anchor='w', pady=(12, 4))
        pv = tk.Frame(right, bg=C['bg'])
        pv.pack(fill='both', expand=True)
        cols = ('w', 'kind', 'name', 'dur', 'trade', 'qty')
        self.tree = ttk.Treeview(pv, columns=cols, show='headings', style='TR.Treeview', height=13)
        for c, t, w in (('w', '', 26), ('kind', 'ItemKind', 90), ('name', 'ชื่อไอเทม', 300),
                        ('dur', 'ระยะเวลา', 90), ('trade', 'แลกเปลี่ยน', 90),
                        ('qty', 'จำนวน', 70)):
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor='w')
        self.tree.tag_configure('warn', background='#3a3018', foreground='#e3b341')
        self.tree.bind('<<TreeviewSelect>>', lambda e: self._why())
        tsb = ttk.Scrollbar(pv, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=tsb.set)
        self.tree.pack(side='left', fill='both', expand=True)
        tsb.pack(side='right', fill='y')

        self.top.protocol('WM_DELETE_WINDOW', self._cancel)
        self.wb = None
        self.cache = {}
        self.lock = threading.Lock()
        self.q = queue.Queue()
        self.inflight = set()
        threading.Thread(target=self._scan_worker, daemon=True).start()
        self.top.after(80, self._pump)
        parent.wait_window(self.top)

    # ---------- รับผลจากเธรดเบื้องหลัง (บนเธรดหลักเท่านั้น) ----------
    def _pump(self):
        try:
            while True:
                msg = self.q.get_nowait()
                kind = msg[0]
                if kind == 'info':
                    self.info.config(text=msg[1])
                elif kind == 'sheets':
                    self.sheets = msg[1]
                    self._fill_sheets()
                elif kind == 'rows':
                    key, rows, note = msg[1], msg[2], msg[3]
                    self.cache[key] = (rows, note)
                    self.inflight.discard(key)
                    if self._key() == key:
                        self._show(rows, note)
        except queue.Empty:
            pass
        except Exception:
            pass
        try:
            self.top.after(80, self._pump)
        except Exception:
            pass

    # ---------- สแกนรายชื่อชีท ----------
    def _scan_worker(self):
        try:
            with self.lock:
                self.q.put(('info', 'กำลังเปิดไฟล์…'))
                self.wb = open_workbook(self.path)
                sheets = scan_sheets_wb(
                    self.wb,
                    progress=lambda k, n, name: self.q.put(
                        ('info', f'กำลังสแกน… {k}/{n}   {name}')))
        except Exception as ex:
            self.q.put(('info', 'อ่านไฟล์ไม่ได้: ' + str(ex)))
            return
        self.q.put(('sheets', sheets))

    def _fill_sheets(self):
        self.lb.delete(0, tk.END)
        first_hit = None
        for i, (name, n) in enumerate(self.sheets):
            self.lb.insert(tk.END, (f'★ ({n})  ' if n else '     ')
                           + (clean_text(name) or name))
            if n and first_hit is None:
                first_hit = i
        hits = sum(1 for _, n in self.sheets if n)
        self.info.config(text=f'ไฟล์นี้มี {len(self.sheets)} ชีท · พบตารางไอเทมใน {hits} ชีท'
                              ' · เลือกชีทที่ต้องการทางซ้าย')
        # ไม่เลือกชีทให้อัตโนมัติ — เพราะเลือกได้หลายชีท ถ้าเลือกไว้ให้ก่อน
        # พอผู้ใช้เลื่อนไปคลิกชีทที่ต้องการ ชีทที่เลือกไว้ให้จะติดมาด้วยโดยไม่รู้ตัว
        # (มันอยู่บนสุดของลิสต์ เลื่อนลงไปแล้วมองไม่เห็น)
        if first_hit is not None:
            self.lb.see(first_hit)
        self._on_sheet()

    # ---------- เลือกชีท (เลือกได้หลายชีทพร้อมกัน) ----------
    def _selected(self):
        """ชื่อชีทที่ถูกเลือกอยู่ เรียงตามลำดับในไฟล์"""
        return [self.sheets[i][0] for i in self.lb.curselection()
                if 0 <= i < len(self.sheets)]

    def _key(self):
        return '\x00'.join(self._selected())

    def _pick_all(self):
        self.lb.selection_clear(0, tk.END)
        for i, (_, n) in enumerate(self.sheets):
            if n:
                self.lb.selection_set(i)
        self._on_sheet()

    def _pick_none(self):
        self.lb.selection_clear(0, tk.END)
        self._on_sheet()

    def _on_sheet(self):
        """คลิกรัวๆ หลายชีทติดกัน ไม่ต้องอ่านไฟล์ทุกคลิก — รอให้หยุดคลิกก่อน"""
        names = self._selected()
        self.sheet_names = names
        self.sheet_name = names[0] if names else ''
        # โชว์ "ชื่อชีท" ที่เลือกไว้ด้วย ไม่ใช่แค่จำนวน — จะได้เห็นทันทีถ้ามีชีทติดมาเกิน
        if not names:
            self.sel_lbl.config(text='ยังไม่ได้เลือกชีท', fg=C['dim'])
        else:
            show = ', '.join(names[:4]) + (' …อีก %d' % (len(names) - 4) if len(names) > 4 else '')
            self.sel_lbl.config(text='เลือกไว้ %d ชีท: %s' % (len(names), show), fg=C['fg'])
        if self._pending is not None:
            try:
                self.top.after_cancel(self._pending)
            except Exception:
                pass
        self._pending = self.top.after(300, self._request)

    def _request(self):
        self._pending = None
        key = self._key()
        if not key:
            self._clear()
            self.info.config(text='เลือกชีทอย่างน้อยหนึ่งชีท')
            return
        if key in self.cache:
            rows, note = self.cache[key]
            self._show(rows, note)
            return
        self._clear()
        names = self._selected()
        self.info.config(text=(f'กำลังอ่าน {len(names)} ชีท…' if len(names) > 1
                               else f'กำลังอ่านชีท “{names[0]}” …'))
        if key in self.inflight:
            return
        self.inflight.add(key)
        threading.Thread(target=self._rows_worker, args=(key, names), daemon=True).start()

    def _rows_worker(self, key, names):
        rows, note = [], ''
        try:
            with self.lock:
                seen = set()          # ตัด ItemKind ซ้ำข้ามชีทให้เหลือตัวเดียว
                hit = []
                for k, name in enumerate(names, 1):
                    if len(names) > 1:
                        self.q.put(('info', 'กำลังอ่าน %d/%d ชีท…  %s  (ได้แล้ว %d รายการ)'
                                    % (k, len(names), name, len(rows))))
                    got = parse_master_rows(read_sheet_rows(self.wb, name), seen)
                    if got:
                        hit.append('%s (%d)' % (name, len(got)))
                    rows.extend(got)
                if len(names) == 1:
                    note = 'ชีท “%s”' % names[0]
                else:
                    note = 'รวม %d ชีท: %s' % (len(names), ', '.join(hit) or '-')
        except Exception as ex:
            self.q.put(('info', 'อ่านไม่ได้: ' + str(ex)))
        self.q.put(('rows', key, rows, note))

    # ---------- แสดงผล ----------
    def _clear(self):
        self.rows = []
        self.tree.delete(*self.tree.get_children())
        self.count_lbl.config(text='')
        self.ok_btn.config(state='disabled', bg=C['input'], fg=C['dim'])

    def _why(self):
        """คลิกแถวเหลือง แล้วบอกว่าเตือนเพราะอะไร"""
        s = self.tree.selection()
        if not s:
            return
        i = self.tree.index(s[0])
        if 0 <= i < len(self.rows):
            w = self.rows[i].get('warn') or []
            if w:
                self.warn_lbl.config(text='⚠  ' + ' · '.join(w), fg=C['warn'])
            else:
                self.warn_lbl.config(text='', fg=C['dim'])

    def _show(self, rows, note):
        self.rows = rows
        nwarn = mark_suspects(rows)
        self.tree.delete(*self.tree.get_children())
        for r in rows[:400]:
            dur = 'ถาวร' if r['dur'] == '' else (r['dur'] + ' วัน')
            trade = {'yes': 'ได้', 'no': 'ไม่ได้', 'any': '— ไม่กรอง —'}.get(r['trade'], r['trade'])
            self.tree.insert('', 'end', tags=('warn',) if r.get('warn') else (),
                             values=('⚠' if r.get('warn') else '', r['kind'], r['disp'],
                                     dur, trade, r['qty'] or '—'))
        has_move = any(r.get('has_move') for r in rows)
        extra = '' if has_move else '   (ชีทนี้ไม่มีคอลัมน์ Itemmove → ไม่กรองแลกเปลี่ยน)'
        self.info.config(text=f'{note} · พบ {len(rows)} รายการ{extra}')
        if nwarn:
            self.warn_lbl.config(
                text=f'⚠  มี {nwarn} แถวที่หน้าตาไม่เหมือนไอเทมทั่วไป (แถวสีเหลือง) — '
                     'ยังนำเข้าให้ครบ คลิกที่แถวเพื่อดูว่าเตือนเพราะอะไร',
                fg=C['warn'])
        else:
            self.warn_lbl.config(text='', fg=C['dim'])
        self.count_lbl.config(text=f'จะนำเข้า {len(rows)} รายการ'
                              + (f'  (น่าสงสัย {nwarn})' if nwarn else '')
                              + ('   (แสดงตัวอย่าง 400 แถวแรก)' if len(rows) > 400 else ''))
        self.ok_btn.config(state='normal' if rows else 'disabled',
                           bg=C['accent'] if rows else C['input'],
                           fg='white' if rows else C['dim'])

    def _close_wb(self):
        try:
            if self.wb:
                self.wb.close()
        except Exception:
            pass
        self.wb = None

    def _ok(self):
        self.result = list(self.rows)
        self._close_wb()
        self.top.destroy()

    def _cancel(self):
        self.result = None
        self._close_wb()
        self.top.destroy()


# ============================================================================
#  [5.9] หน้าต่างนำเข้าบันเดิลจากไฟล์ต้นฉบับ — เลือกชีท แล้วดูว่าได้บันเดิลอะไรบ้าง
# ============================================================================
class BundleImportDialog:
    # ข้อความบนหน้าต่าง — คลาสลูก (CodeImportDialog) เปลี่ยนได้โดยไม่ต้องก๊อปโค้ดทั้งก้อน
    TITLE = 'นำเข้าบันเดิลจาก Excel'
    PICK_ALL = 'เลือกชีทที่มีบันเดิลทั้งหมด'
    HINT = ('อ่านให้อัตโนมัติ — Aztek Item Id เอาจากคอลัมน์ Bundle ของตาราง · '
            'จำนวนเอาจาก Amt/Amount · Tier เอาจาก Rank · '
            'famepoint/exp เอาจากกล่อง “ได้รับ famepoint กับ exp”')
    TREE_TITLE = 'บันเดิลที่เจอ'
    TREE_HEAD = 'บันเดิล / ไอเทม'
    scan_fn = staticmethod(lambda wb, progress=None: scan_bundle_sheets_wb(wb, progress))
    read_fn = staticmethod(lambda wb, nm: read_bundles_wb(wb, nm))

    def __init__(self, parent, path):
        self.path = path
        self.result = None
        self.warns = []
        self.sheets = []
        self.bundles = []
        self._pending = None

        self.top = tk.Toplevel(parent)
        self.top.title(self.TITLE)
        self.top.configure(bg=C['bg'])
        self.top.geometry('1010x660')
        self.top.transient(parent)
        self.top.grab_set()

        head = tk.Frame(self.top, bg=C['card'], height=56)
        head.pack(fill='x')
        head.pack_propagate(False)
        tk.Label(head, text=self.TITLE, bg=C['card'], fg=C['fg'],
                 font=('Segoe UI', 13, 'bold')).pack(side='left', padx=18)
        tk.Label(head, text=os.path.basename(path), bg=C['card'], fg=C['dim'],
                 font=('Segoe UI', 9)).pack(side='left')

        foot = tk.Frame(self.top, bg=C['card'], height=62)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)
        self.count_lbl = tk.Label(foot, text='', bg=C['card'], fg=C['dim'],
                                  font=('Segoe UI', 10))
        self.count_lbl.pack(side='left', padx=18)
        self.ok_btn = tk.Button(foot, text='ใช้ข้อมูลนี้', bg=C['input'], fg=C['dim'], bd=0,
                                font=('Segoe UI', 10, 'bold'), cursor='hand2',
                                activebackground='#3d55cf', activeforeground='white',
                                state='disabled', command=self._ok)
        self.ok_btn.pack(side='right', padx=18, pady=12, ipadx=22, ipady=5)
        tk.Button(foot, text='ยกเลิก', bg=C['input'], fg=C['fg'], bd=0,
                  font=('Segoe UI', 10), cursor='hand2', activebackground=C['line'],
                  command=self._cancel).pack(side='right', pady=12, ipadx=16, ipady=5)

        body = tk.Frame(self.top, bg=C['bg'])
        body.pack(fill='both', expand=True, padx=14, pady=12)

        left = tk.Frame(body, bg=C['bg'], width=280)
        left.pack(side='left', fill='y')
        left.pack_propagate(False)
        tk.Label(left, text='เลือกชีท', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(anchor='w')
        lw = tk.Frame(left, bg=C['bg'])
        lw.pack(fill='both', expand=True, pady=(5, 0))
        self.lb = tk.Listbox(lw, bg=C['input'], fg=C['fg'], bd=0, highlightthickness=1,
                             highlightbackground=C['line'], selectbackground=C['accent'],
                             selectforeground='white', font=('Segoe UI', 9),
                             activestyle='none', selectmode='multiple')
        sb = ttk.Scrollbar(lw, orient='vertical', command=self.lb.yview)
        self.lb.configure(yscrollcommand=sb.set)
        self.lb.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')
        self.lb.bind('<<ListboxSelect>>', lambda e: self._on_sheet())
        self.sel_lbl = tk.Label(left, text='ยังไม่ได้เลือกชีท', bg=C['bg'], fg=C['dim'],
                                font=('Segoe UI', 9, 'bold'), anchor='w')
        self.sel_lbl.pack(fill='x', pady=(8, 2))
        bf = tk.Frame(left, bg=C['bg'])
        bf.pack(fill='x')
        for txt, cmd in ((self.PICK_ALL, self._pick_all),
                         ('ล้างที่เลือก', self._pick_none)):
            tk.Button(bf, text=txt, bg=C['input'], fg=C['fg'], bd=0, font=('Segoe UI', 9),
                      cursor='hand2', activebackground=C['line'], command=cmd).pack(
                side='left', padx=(0, 6), ipadx=6, ipady=3)
        tk.Label(left, text='คลิกชีทเพื่อเลือก คลิกซ้ำเพื่อเอาออก — เลือกได้หลายชีท',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), anchor='w',
                 justify='left', wraplength=265).pack(fill='x', pady=(8, 0))

        right = tk.Frame(body, bg=C['bg'])
        right.pack(side='left', fill='both', expand=True, padx=(14, 0))
        self.info = tk.Label(right, text='กำลังสแกนไฟล์…', bg=C['bg'], fg=C['dim'],
                             font=('Segoe UI', 9), anchor='w', justify='left')
        self.info.pack(fill='x')
        tk.Label(right, text=self.HINT,
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), wraplength=690,
                 justify='left', anchor='w').pack(fill='x', pady=(6, 0))
        self.warn_lbl = tk.Label(right, text='', bg=C['bg'], fg=C['warn'],
                                 font=('Segoe UI', 9), anchor='w', justify='left',
                                 wraplength=690)
        self.warn_lbl.pack(fill='x', pady=(6, 0))

        tk.Label(right, text=self.TREE_TITLE, bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(anchor='w', pady=(10, 4))
        pv = tk.Frame(right, bg=C['bg'])
        pv.pack(fill='both', expand=True)
        self.tree = ttk.Treeview(pv, columns=('qty', 'tier', 'info'), show='tree headings',
                                 style='TR.Treeview', height=13)
        self.tree.heading('#0', text=self.TREE_HEAD)
        self.tree.column('#0', width=330, anchor='w')
        for c, t, w in (('qty', 'จำนวน', 70), ('tier', 'Tier', 70), ('info', 'รายละเอียด', 220)):
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor='w')
        self.tree.tag_configure('rw', foreground=C['ok'])
        tsb = ttk.Scrollbar(pv, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=tsb.set)
        self.tree.pack(side='left', fill='both', expand=True)
        tsb.pack(side='right', fill='y')

        self.top.protocol('WM_DELETE_WINDOW', self._cancel)
        self.wb = None
        self.cache = {}
        self.lock = threading.Lock()
        self.q = queue.Queue()
        self.inflight = set()
        threading.Thread(target=self._scan, daemon=True).start()
        self.top.after(80, self._pump)
        parent.wait_window(self.top)

    def _pump(self):
        try:
            while True:
                msg = self.q.get_nowait()
                if msg[0] == 'info':
                    self.info.config(text=msg[1])
                elif msg[0] == 'sheets':
                    self.sheets = msg[1]
                    self._fill()
                elif msg[0] == 'rows':
                    key, bs, ws, note = msg[1], msg[2], msg[3], msg[4]
                    self.cache[key] = (bs, ws, note)
                    self.inflight.discard(key)
                    if self._key() == key:
                        self._show(bs, ws, note)
        except queue.Empty:
            pass
        except Exception:
            pass
        try:
            self.top.after(80, self._pump)
        except Exception:
            pass

    def _scan(self):
        try:
            with self.lock:
                self.q.put(('info', 'กำลังเปิดไฟล์…'))
                self.wb = open_workbook(self.path)
                sheets = self.scan_fn(
                    self.wb, progress=lambda k, n, nm: self.q.put(
                        ('info', 'กำลังสแกน… %d/%d   %s' % (k, n, nm))))
        except Exception as ex:
            self.q.put(('info', 'อ่านไฟล์ไม่ได้: ' + str(ex)))
            return
        self.q.put(('sheets', sheets))

    def _fill(self):
        self.lb.delete(0, tk.END)
        first = None
        for i, (name, n) in enumerate(self.sheets):
            self.lb.insert(tk.END, ('★ (%d)  ' % n if n else '     ')
                           + (clean_text(name) or name))
            if n and first is None:
                first = i
        hits = sum(1 for _, n in self.sheets if n)
        self.info.config(text='ไฟล์นี้มี %d ชีท · พบบันเดิลใน %d ชีท · เลือกชีททางซ้าย'
                              % (len(self.sheets), hits))
        if first is not None:
            self.lb.see(first)
        self._on_sheet()

    def _selected(self):
        return [self.sheets[i][0] for i in self.lb.curselection()
                if 0 <= i < len(self.sheets)]

    def _key(self):
        return '\x00'.join(self._selected())

    def _pick_all(self):
        self.lb.selection_clear(0, tk.END)
        for i, (_, n) in enumerate(self.sheets):
            if n:
                self.lb.selection_set(i)
        self._on_sheet()

    def _pick_none(self):
        self.lb.selection_clear(0, tk.END)
        self._on_sheet()

    def _on_sheet(self):
        names = self._selected()
        if not names:
            self.sel_lbl.config(text='ยังไม่ได้เลือกชีท', fg=C['dim'])
        else:
            show = ', '.join(names[:4]) + (' …อีก %d' % (len(names) - 4)
                                           if len(names) > 4 else '')
            self.sel_lbl.config(text='เลือกไว้ %d ชีท: %s' % (len(names), show), fg=C['fg'])
        if self._pending is not None:
            try:
                self.top.after_cancel(self._pending)
            except Exception:
                pass
        self._pending = self.top.after(300, self._request)

    def _request(self):
        self._pending = None
        key = self._key()
        if not key:
            self._clear()
            self.info.config(text='เลือกชีทอย่างน้อยหนึ่งชีท')
            return
        if key in self.cache:
            self._show(*self.cache[key])
            return
        self._clear()
        names = self._selected()
        self.info.config(text='กำลังอ่าน %d ชีท…' % len(names))
        if key in self.inflight:
            return
        self.inflight.add(key)
        threading.Thread(target=self._work, args=(key, names), daemon=True).start()

    def _work(self, key, names):
        bs, ws = [], []
        try:
            with self.lock:
                for nm in names:
                    b, w = self.read_fn(self.wb, nm)
                    bs.extend(b)
                    ws.extend(w)
        except Exception as ex:
            self.q.put(('info', 'อ่านไม่ได้: ' + str(ex)))
        note = ('ชีท “%s”' % names[0]) if len(names) == 1 else ('รวม %d ชีท' % len(names))
        self.q.put(('rows', key, bs, ws, note))

    def _clear(self):
        self.bundles = []
        self.warns = []
        self.tree.delete(*self.tree.get_children())
        self.count_lbl.config(text='')
        self.warn_lbl.config(text='')
        self.ok_btn.config(state='disabled', bg=C['input'], fg=C['dim'])

    def _show(self, bs, ws, note):
        self.bundles = bs
        self.warns = ws
        self.tree.delete(*self.tree.get_children())
        for b in bs:
            rw = ' · '.join(('famepoint' if r['type'] == 'CREDIT' else 'exp') + ' ' + r['qty']
                            for r in b.get('rewards', []))
            pid = self.tree.insert('', 'end', text='📦  ' + b['name'], open=True,
                                   values=('', '', '%d ไอเทม%s' % (len(b['items']),
                                                                   ('  ·  ' + rw) if rw else '')))
            for it in b['items']:
                self.tree.insert(pid, 'end', text='      ' + (it.get('disp') or ''),
                                 values=(it['qty'], it['tier'], 'Id ' + it['id']))
            for r in b.get('rewards', []):
                nm = FAME_NAME if r['type'] == 'CREDIT' else 'Player Experience'
                self.tree.insert(pid, 'end', text='      ⭐ ' + nm, tags=('rw',),
                                 values=(r['qty'], DEFAULT_TIER,
                                         'เครดิตเรียลไทม์' if r['type'] == 'CREDIT' else 'exp'))
        self.info.config(text='%s · พบ %d บันเดิล' % (note, len(bs)))
        self.warn_lbl.config(text=('⚠  %d จุดที่ต้องดู เช่น %s' % (len(ws), ws[0][:90]))
                             if ws else '')
        self.count_lbl.config(text='จะนำเข้า %d บันเดิล' % len(bs) if bs else '')
        self.ok_btn.config(state='normal' if bs else 'disabled',
                           bg=C['accent'] if bs else C['input'],
                           fg='white' if bs else C['dim'])

    def _close(self):
        try:
            if self.wb:
                self.wb.close()
        except Exception:
            pass
        self.wb = None

    def _ok(self):
        self.result = list(self.bundles)
        self._close()
        self.top.destroy()

    def _cancel(self):
        self.result = None
        self._close()
        self.top.destroy()


class CodeImportDialog(BundleImportDialog):
    """หน้าต่างนำเข้า Item Code — โครงเดียวกับนำเข้าบันเดิล เปลี่ยนแค่สิ่งที่อ่านกับที่โชว์"""
    TITLE = 'นำเข้า Item Code จาก Excel (WR Master)'
    PICK_ALL = 'เลือกชีทที่มีโค้ดทั้งหมด'
    HINT = ('อ่านเฉพาะชีทแบบ Master Code (มีแถว CODE) — '
            'ชื่อ Item Code = โค้ด + วันที่เริ่ม · เวลาเอาจากช่อง Start/End ในชีท · '
            'Limit ที่ชีทบอกจะเผื่อให้อีก %d · Bundle เอาจากเลขใต้ป้าย Bundle' % WR_SPARE)
    TREE_TITLE = 'Item Code ที่เจอ'
    TREE_HEAD = 'CODE / ชื่อ Item Code'
    scan_fn = staticmethod(lambda wb, progress=None: scan_code_sheets_wb(wb, progress))
    read_fn = staticmethod(lambda wb, nm: read_codes_wb(wb, nm))

    def _show(self, bs, ws, note):
        self.bundles = bs
        self.warns = ws
        self.tree.delete(*self.tree.get_children())
        for d in bs:
            pid = self.tree.insert(
                '', 'end', text='🎟  ' + d['code'], open=False,
                values=(d.get('cap') or 'ไม่จำกัด',
                        ('Bundle ' + d['bundle']) if d.get('bundle') else '⚠ ไม่มี Bundle',
                        d['name']))
            self.tree.insert(pid, 'end', text='      เริ่มใช้งาน',
                             values=('', '', d.get('start') or '⚠ ไม่เจอในชีท'))
            self.tree.insert(pid, 'end', text='      สิ้นสุด',
                             values=('', '', d.get('end') or '⚠ ไม่เจอในชีท'))
        nobd = sum(1 for d in bs if not d.get('bundle'))
        self.info.config(text='%s · พบ %d โค้ด' % (note, len(bs)))
        msg = ''
        if nobd:
            msg = '⚠  %d โค้ดยังไม่มีเลข Bundle ในชีท' % nobd
        if ws:
            msg = (msg + '  ·  ' if msg else '') + '%d จุดที่ต้องดู เช่น %s' % (
                len(ws), ws[0][:80])
        self.warn_lbl.config(text=msg)
        self.count_lbl.config(text='จะนำเข้า %d โค้ด' % len(bs) if bs else '')
        self.ok_btn.config(state='normal' if bs else 'disabled',
                           bg=C['accent'] if bs else C['input'],
                           fg='white' if bs else C['dim'])


# ============================================================================
#  [5.6] ระบบตรวจสอบตัวเอง (Self Test)
#        แบ่งเป็น 2 ชุด
#          A. ตรวจตรรกะ  — ไม่ต้องเปิดเว็บ เช็กว่ากฎการอ่านชีทยังถูกต้อง
#          B. ตรวจเว็บ    — เปิดเว็บจริง เช็กว่าหน้าตาเว็บยังตรงกับ CFG ไหม
#        ถ้าเว็บเปลี่ยน จะได้รู้ทันทีว่า "พังตรงไหน" และ "ต้องแก้ค่าตัวไหน"
# ============================================================================
def _t(ok, name, detail='', fix=''):
    return {'ok': bool(ok), 'name': name, 'detail': str(detail), 'fix': fix}


def run_logic_tests():
    """ตรวจกฎการอ่านข้อมูลทั้งหมด โดยไม่ต้องต่อเน็ต"""
    r = []

    # ---- กติกาช่องระยะเวลา ----
    r.append(_t(dur_from_text('') == '', 'ระยะเวลา: เว้นว่าง = ถาวร', repr(dur_from_text(''))))
    r.append(_t(dur_from_text('ถาวร') == '', 'ระยะเวลา: "ถาวร" = ถาวร', repr(dur_from_text('ถาวร'))))
    r.append(_t(dur_from_text('any') == 'any', 'ระยะเวลา: "any" = ไม่กรอง', repr(dur_from_text('any'))))
    r.append(_t(dur_from_text('15 วัน') == '15', 'ระยะเวลา: "15 วัน" = 15', repr(dur_from_text('15 วัน'))))

    # ---- ตัวเลขจาก Excel ----
    r.append(_t(num_str(112247.0) == '112247', 'ตัดทศนิยม .0 ที่ Excel ติดมา', num_str(112247.0)))

    # ---- Yes/No ----
    r.append(_t(all(norm_bool(v) == 'yes' for v in ('Yes', 'y', '1', 'TRUE')),
                'อ่านค่า Yes ได้ทุกแบบ'))
    r.append(_t(all(norm_bool(v) == 'no' for v in ('No', 'n', '0', 'false')),
                'อ่านค่า No ได้ทุกแบบ'))
    r.append(_t(norm_bool('') == 'any' and norm_bool('มั่ว') == 'any',
                'ค่าว่าง/อ่านไม่ออก = ไม่กรอง'))

    # ---- อ่านชีทต้นฉบับ (กฎเดียวกับเครื่องมือเดิม) ----
    rows = [
        ['400', 'fdItemNum', 'fdPosition', 'fdItemKind', 'Rank',
         'Display Name', 'ระยะเวลาของขวัญ', 'Amt', 'ราคา'],
        ['', '112247', '9999', '2648', 'SS', 'New.1 ตั๋วเปลี่ยนตำนาน II 1 ชิ้น', 'ถาวร', '1', '150'],
        ['', '128172', '250', '5485', 'S', 'กล่องอิลลิเฟียร์ 5 ชิ้น', 'ถาวร', '1', '95'],
        ['', '999001', '250', '5486', 'A', 'หมวกทดสอบ 3 ชิ้น', '7 วัน', '1', '10'],
        ['', '', '', '', '', '', '', '', '460'],
        ['', 'ขอรางวัลเริ่ดๆ เหมือนเดิมค่า'],
    ]
    got = parse_master_rows(rows)
    r.append(_t(len(got) == 3, 'อ่านชีทตัวอย่างได้ 3 รายการ (ตัดแถวยอดรวม/หมายเหตุ)', len(got)))
    if len(got) == 3:
        a, b, c = got
        r.append(_t(a['kind'] == '112247', 'ใช้คอลัมน์ fdItemNum เป็น ItemKind', a['kind']))
        r.append(_t(a['qty'] == '1' and b['qty'] == '5' and c['qty'] == '3',
                    'ดึงจำนวนจากเลขหน้าคำว่า "ชิ้น"',
                    f"{a['qty']}/{b['qty']}/{c['qty']}"))
        r.append(_t(b['disp'] == 'กล่องอิลลิเฟียร์', 'ตัดคำว่า "X ชิ้น" ออกจากชื่อ', b['disp']))
        r.append(_t(a['dur'] == '' and c['dur'] == '7',
                    'อ่านระยะเวลา: ถาวร / 7 วัน', f"{a['dur']!r}/{c['dur']!r}"))
        r.append(_t(a['trade'] == 'any', 'ไม่มีคอลัมน์ Itemmove = ไม่กรองแลกเปลี่ยน', a['trade']))

    # ---- มีคอลัมน์ Itemmove ----
    rows2 = [
        ['fdItemNum', 'Display Name', 'ระยะเวลาของขวัญ', 'Itemmove'],
        ['111111', 'ของทดสอบ 2 ชิ้น', 'ถาวร', 'Yes'],
        ['222222', 'ของทดสอบสอง 1 ชิ้น', 'ถาวร', ''],
    ]
    g2 = parse_master_rows(rows2)
    r.append(_t(len(g2) == 2 and g2[0]['trade'] == 'yes' and g2[1]['trade'] == 'no',
                'มีคอลัมน์ Itemmove: Yes = ได้ / ว่าง = ไม่ได้',
                [x['trade'] for x in g2]))

    # ---- ตัด ID ซ้ำ ----
    seen = set()
    d1 = parse_master_rows(rows, seen)
    d2 = parse_master_rows(rows, seen)
    r.append(_t(len(d1) == 3 and len(d2) == 0, 'ตัด ID ซ้ำข้ามชีทได้', f'{len(d1)} / {len(d2)}'))

    # ---- เงื่อนไข deep check ----
    item7 = {'duration': '7', 'trade': False, 'qty': '10'}
    r.append(_t(match_deep(item7, {'dur': '7', 'trade': 'any', 'qty': ''})[0],
                'deep check: ระยะเวลาตรง = ผ่าน'))
    r.append(_t(not match_deep(item7, {'dur': '15', 'trade': 'any', 'qty': ''})[0],
                'deep check: ระยะเวลาไม่ตรง = ไม่ผ่าน'))
    r.append(_t(match_deep({'duration': '0', 'trade': None, 'qty': ''},
                           {'dur': '', 'trade': 'any', 'qty': ''})[0],
                'deep check: ระยะเวลา 0 บนเว็บ = ถาวร'))
    r.append(_t(not match_deep(item7, {'dur': 'any', 'trade': 'yes', 'qty': ''})[0],
                'deep check: แลกเปลี่ยนไม่ตรง = ไม่ผ่าน'))
    r.append(_t(not match_deep(item7, {'dur': 'any', 'trade': 'any', 'qty': '5'})[0],
                'deep check: จำนวนไม่ตรง = ไม่ผ่าน'))

    # ---- ทางลัดใช้ฟิลเตอร์หน้า list ----
    r.append(_t(bool(duration_only_numeric({'dur': '15', 'trade': 'any', 'qty': ''})),
                'ใช้ทางลัดฟิลเตอร์เมื่อกรองแค่ระยะเวลา'))
    r.append(_t(not duration_only_numeric({'dur': '15', 'trade': 'yes', 'qty': ''}),
                'ไม่ใช้ทางลัดเมื่อมีเงื่อนไขอื่นด้วย'))

    # ---- ระบบบันทึกการใช้งาน ----
    try:
        before = len(read_events())
        log_event('selftest_ping')
        after = len(read_events())
        ok_log = after == before + 1
        detail = f'{after} เหตุการณ์ในไฟล์'
    except Exception as ex:
        ok_log, detail = False, str(ex)[:80]
    r.append(_t(ok_log, 'บันทึกการใช้งานลงไฟล์ได้', f'{detail}  ({LOCAL_LOG_DIR})',
                'เช็กสิทธิ์เขียนโฟลเดอร์ที่ตั้งไว้'))
    st = read_stats()
    r.append(_t(st.get('total_events', 0) > 0, 'ตัวนับสถิติทำงาน',
                f"รวม {st.get('total_events', 0)} ครั้ง · เริ่มเก็บ {st.get('first_used', '-')}"))
    if CENTRAL_LOG_DIR:
        r.append(_t(os.path.isdir(CENTRAL_LOG_DIR), 'โฟลเดอร์กลางของทีมเขียนได้',
                    CENTRAL_LOG_DIR, 'โฟลเดอร์หายไป — ตั้งใหม่ที่ปุ่มข้างบน'))
    else:
        r.append(_t(True, 'โฟลเดอร์กลางของทีม',
                    'ยังไม่ได้ตั้ง — เก็บเฉพาะในเครื่องนี้ (ตั้งได้ที่ปุ่มข้างบน)'))

    # ---- กติกาหน้า "สร้างไอเทม" ----
    a = split_name_qty('กล่องแอนิมอล 1 ชิ้น')
    r.append(_t(a == ('กล่องแอนิมอล', '1', True),
                'สร้างไอเทม: ตัด "จำนวนชิ้น" ออกจากชื่อ', repr(a),
                'กฎเดิม: จำนวนชิ้นไม่ติดไปในชื่อ'))
    b = split_name_qty('แพ็กเบ็ดตกปลา 2 ชั่วโมง 3 ชิ้น')
    r.append(_t(b == ('แพ็กเบ็ดตกปลา 2 ชั่วโมง', '3', True),
                'สร้างไอเทม: เก็บ "ชม." ไว้ในชื่อ ตัดแต่จำนวน', repr(b),
                'กฎเดิม: ชม. อยู่ในชื่อได้ จำนวนอยู่คนละช่อง'))
    c = split_name_qty('ค้อนซ่อมแซม')
    r.append(_t(c == ('ค้อนซ่อมแซม', '1', False),
                'สร้างไอเทม: ชื่อที่ไม่มี "ชิ้น" คงชื่อเดิม จำนวน 1', repr(c)))
    d1 = make_item_name('กล่องแอนิมอล', '(Gift)', '3')
    r.append(_t(d1 == 'กล่องแอนิมอล (Gift) (3 วัน)',
                'สร้างไอเทม: ประกอบชื่อ + คำต่อท้าย + (X วัน)', d1))
    e1 = row_to_item({'kind': '1', 'cname': 'ค้อน', 'amount': '1',
                      'has_qty': False, 'dur': '5', 'trade': 'yes'}, 'GENERAL', '(Gift)')
    r.append(_t(e1['name'] == 'ค้อน',
                'สร้างไอเทม: ชื่อไม่มี "ชิ้น" -> ไม่เติมคำต่อท้าย/วัน', e1['name'],
                'กฎเดิมเติมเฉพาะชื่อที่มี "ชิ้น"'))

    # ---- ระบบเตือนของแปลกปลอม (เตือนอย่างเดียว ห้ามตัดแถวทิ้ง) ----
    junk = [{'kind': '128107', 'raw': 'ก'}, {'kind': '128108', 'raw': 'ข'},
            {'kind': '128109', 'raw': 'ค'}, {'kind': '2', 'raw': 'แก้ UID เรียบร้อย'}]
    nw = mark_suspects(junk)
    r.append(_t(nw == 1 and len(junk) == 4,
                'เตือนแถวแปลกปลอมได้ โดยไม่ตัดแถวทิ้ง',
                'เตือน %d แถว · เหลือ %d แถว' % (nw, len(junk))))
    longest = 'แพ็กเบ็ดตกปลาเทลส์รันเนอร์ 2 ชั่วโมง (Gift) (30 วัน)'
    r.append(_t(suspect_reasons({'kind': '128909', 'raw': longest}, 6) == [],
                'ชื่อไอเทมจริงที่ยาวที่สุด ไม่ถูกเตือนผิด',
                '%d ตัวอักษร (เพดาน %d)' % (len(longest), MAX_SANE_NAME)))
    junk[3]['kind'] = '128110'
    r.append(_t(mark_suspects(junk) == 0,
                'แก้ ItemKind แล้วคำเตือนหายเอง'))

    # ---- กติกาหน้า "สร้างบันเดิล" ----
    brows = [
        ['', '', 'ได้รับ famepoint กับ exp'],
        ['', '', '99'],
        ['', 'Product Name', 'แพ็กเกจทดสอบ (Gift)'],
        ['2800', 'fdItemNum', 'fdPosition', 'fdItemKind', 'Name', 'Rank', 'Amount'],
        ['2074', '89369', '2', '5', 'ฮันบก ซ่อมแซม', 'A', '1'],
        ['2077', '110256', '121', '52', 'เสียงรีเมค 30 วัน', 'General', '2'],
    ]
    bb, _bw = parse_bundle_sheet(brows, 'ทดสอบ')
    r.append(_t(len(bb) == 1, 'บันเดิล: อ่านบล็อกจากชีทได้', '%d บันเดิล' % len(bb)))
    if bb:
        b0 = bb[0]
        r.append(_t([i['id'] for i in b0['items']] == ['2074', '2077'],
                    'บันเดิล: Aztek Item Id เอาจากเลขซ้ายมือของตาราง',
                    str([i['id'] for i in b0['items']]),
                    'เลขนี้คือตัวที่เอาไปค้นบนเว็บ'))
        r.append(_t([i['qty'] for i in b0['items']] == ['1', '2'],
                    'บันเดิล: จำนวนเอาจากคอลัมน์ Amount/Amt',
                    str([i['qty'] for i in b0['items']])))
        r.append(_t([i['tier'] for i in b0['items']] == ['A', 'General'],
                    'บันเดิล: Tier เอาจากคอลัมน์ Rank',
                    str([i['tier'] for i in b0['items']])))
        r.append(_t(b0['name'] == 'แพ็กเกจทดสอบ (Gift)',
                    'บันเดิล: ชื่อเอาจาก Product Name', b0['name']))
        fame = [x for x in b0['rewards'] if x['type'] == 'CREDIT']
        exp = [x for x in b0['rewards'] if x['type'] == 'PLAYER_EXP']
        r.append(_t(bool(fame) and fame[0]['qty'] == '99' and bool(exp) and exp[0]['qty'] == '99',
                    'บันเดิล: อ่าน famepoint กับ exp จากกล่องด้านบน',
                    str(b0['rewards'])))
    r.append(_t(FAME_CREDIT_TYPE == 'WALLET_REALTIME_CREDIT',
                'บันเดิล: famepoint ตั้งเป็นเครดิตเรียลไทม์', FAME_CREDIT_LABEL,
                'ถ้าเป็นเครดิตธรรมดาจะผิด'))

    # ---- สภาพแวดล้อม ----
    r.append(_t(XLSX_OK, 'อ่านไฟล์ Excel ได้ (openpyxl)', 'ok' if XLSX_OK else 'ไม่มี openpyxl'))
    chrome = find_chrome_exe()
    r.append(_t(chrome, 'เจอ Google Chrome ในเครื่อง', chrome or 'ไม่พบ',
                'ติดตั้ง Chrome ก่อนใช้งาน'))
    return r


JS_HEADERS = """() => [...document.querySelectorAll('table thead th')]
                        .map(x => (x.innerText || '').trim())"""


async def run_web_tests(page, log=None):
    """เปิดเว็บจริงแล้วตรวจว่าหน้าตายังตรงกับ CFG ไหม"""
    r = []

    def say(msg):
        if log:
            log(msg)

    # --- W1 เปิดหน้า Item ---
    try:
        await page.goto(ITEM_LIST_URL, wait_until='domcontentloaded', timeout=45000)
        await page.wait_for_timeout(1800)
        r.append(_t(True, 'เปิดหน้า Shop > Item ได้', page.url.split('?')[0]))
    except Exception as ex:
        r.append(_t(False, 'เปิดหน้า Shop > Item ได้', str(ex)[:90], 'เช็กอินเทอร์เน็ต / ITEM_LIST_URL'))
        return r
    say('  เปิดหน้า Item แล้ว')

    # --- W2 ล็อกอินอยู่ไหม ---
    has_search = await page.locator(SEL['search_input']).count() > 0
    r.append(_t(has_search, 'ล็อกอินอยู่ และเจอช่องค้นหา',
                'เจอ' if has_search else 'ไม่เจอช่องค้นหา',
                'กดปุ่ม "เปิดหน้า Login" แล้วล็อกอินก่อน'))
    if not has_search:
        return r

    # --- W3 หัวตาราง ---
    try:
        heads = await page.evaluate(JS_HEADERS)
    except Exception:
        heads = []
    ok_head = len(heads) >= 4 and 'Aztek Item Id' in ' '.join(heads) and 'ItemKind' in ' '.join(heads)
    r.append(_t(ok_head, 'ตารางมีคอลัมน์ Aztek Item Id และ ItemKind',
                ' | '.join(heads) or 'อ่านหัวตารางไม่ได้', 'แก้ COL ใน CONFIG'))

    # --- W4 อ่านแถวได้ ---
    rows = await page.evaluate(JS_READ_ROWS, COL)
    r.append(_t(len(rows) > 0, 'อ่านแถวในตารางได้', f'{len(rows)} แถว', 'แก้ SEL["row"] หรือ COL'))
    if not rows:
        return r
    first = rows[0]
    r.append(_t(re.fullmatch(r'\d+', first['id'] or '') and first['name'],
                'คอลัมน์ ID เป็นตัวเลข และมีชื่อไอเทม',
                f"id={first['id']} name={first['name'][:24]}", 'ลำดับคอลัมน์ COL เปลี่ยน'))

    # --- W5 ปุ่มค้นหา + ฟิลเตอร์ ---
    r.append(_t(await page.locator(SEL['search_button']).count() > 0,
                'เจอปุ่มค้นหา', SEL['search_button'], 'แก้ SEL["search_button"]'))
    r.append(_t(await page.locator(SEL['filter_duration']).count() > 0,
                'เจอฟิลเตอร์ DurationIndex', SEL['filter_duration'],
                'แก้ SEL["filter_duration"]'))
    nxt = await page.locator(f'button:has-text("{SEL["next_text"]}")').count()
    r.append(_t(nxt > 0, 'เจอปุ่มเปลี่ยนหน้า (Next)', f'{nxt} ปุ่ม', 'แก้ SEL["next_text"]'))

    # --- W6 ค้นหาจริงแล้วต้องเจอ ---
    kind = first['kind'] or first['id']
    say(f'  ทดลองค้นหา {kind}')
    try:
        await page.fill(SEL['search_input'], str(kind))
        await page.locator(SEL['search_button']).first.click()
        await page.wait_for_timeout(1600)
        found = await page.evaluate(JS_READ_ROWS, COL)
        hit = any(x['kind'] == first['kind'] or x['id'] == first['id'] for x in found)
        r.append(_t(hit, f'ค้นหา {kind} แล้วเจอไอเทมที่ต้องการ', f'ได้ {len(found)} แถว',
                    'ช่องค้นหาหรือปุ่มค้นหาเปลี่ยนไป'))
    except Exception as ex:
        r.append(_t(False, f'ค้นหา {kind} แล้วเจอไอเทมที่ต้องการ', str(ex)[:90]))
        return r

    # --- W7 เข้าหน้ารายละเอียด ---
    target = first['id']
    try:
        clicked = await page.evaluate(JS_CLICK_ROW, str(target))
        if clicked:
            await page.wait_for_selector(SEL['d_name'], timeout=15000)
            await page.wait_for_timeout(500)
        r.append(_t(clicked, 'คลิกเข้าหน้ารายละเอียดไอเทมได้',
                    page.url.split('/')[-1], 'ลิงก์ในตารางเปลี่ยนรูปแบบ'))
        if not clicked:
            return r
    except Exception as ex:
        r.append(_t(False, 'คลิกเข้าหน้ารายละเอียดไอเทมได้', str(ex)[:90]))
        return r

    # --- W8 ฟิลด์ในหน้ารายละเอียด ---
    for key, label in (('d_name', 'ชื่อไอเท็ม'), ('d_kind', 'game_item_id (ItemKind)'),
                       ('d_price', 'Price'), ('d_duration', 'ระยะเวลาไอเท็ม (วัน)'),
                       ('d_qty', 'จำนวน')):
        n = await page.locator(SEL[key]).count()
        r.append(_t(n > 0, f'เจอช่อง {label}', SEL[key], f'แก้ SEL["{key}"]'))

    detail = await page.evaluate(JS_READ_DETAIL, {
        'name': SEL['d_name'], 'kind': SEL['d_kind'], 'price': SEL['d_price'],
        'duration': SEL['d_duration'], 'qty': SEL['d_qty'],
        'tradeLabel': SEL['d_trade_label'],
    })
    r.append(_t(detail.get('trade') is not None, 'อ่านค่า "แลกเปลี่ยนได้" ได้',
                detail.get('trade'), 'แก้ SEL["d_trade_label"]'))

    # --- W9 ข้อมูลตรงกันระหว่างตารางกับหน้ารายละเอียด (ตรวจ mapping ทั้งเส้น) ---
    same = str(detail.get('kind') or '').strip() == str(first['kind'] or '').strip()
    r.append(_t(same, 'ItemKind ในตาราง ตรงกับในหน้ารายละเอียด',
                f"ตาราง={first['kind']}  รายละเอียด={detail.get('kind')}",
                'คอลัมน์ ItemKind ในตารางอาจสลับตำแหน่ง — แก้ COL'))
    return r

# ============================================================================
#  [6] หน้าต่างโปรแกรม
# ============================================================================
class SideTabs:
    """แถบแท็บทำเอง — กลุ่ม "ใช้งาน" อยู่ซ้าย กลุ่ม "ดูข้อมูล" ไปชิดขวา

    ttk.Notebook ดันแท็บไปชิดขวาไม่ได้ เลยทำแถบเอง
    จะได้ไม่ปนกันระหว่างแท็บที่ใช้ทำงาน กับแท็บที่ไว้ดูข้อมูล
    ใช้แทน Notebook ได้เลย — มี add / select / tab / index เหมือนกัน
    """

    def __init__(self, parent):
        self.bar = tk.Frame(parent, bg=C['bg'])
        self.bar.pack(fill='x', padx=10, pady=(8, 0))
        self._l = tk.Frame(self.bar, bg=C['bg'])
        self._l.pack(side='left')
        self._r = tk.Frame(self.bar, bg=C['bg'])
        self._r.pack(side='right')
        self.body = tk.Frame(parent, bg=C['bg'], highlightthickness=1,
                             highlightbackground=C['line'])
        self.body.pack(fill='both', expand=True, padx=10, pady=(0, 10))
        self._tabs = []
        self._cur = None

    def add(self, frame, text, side='left'):
        b = tk.Button(self._l if side == 'left' else self._r, text=text, bd=0,
                      cursor='hand2', font=FM, bg=C['card'], fg=C['dim'],
                      activebackground=C['line'], activeforeground=C['fg'],
                      padx=16, pady=7,
                      command=lambda f=frame: self.select(f))
        b.pack(side='left', padx=(0, 2))
        self._tabs.append({'frame': frame, 'text': text, 'btn': b})
        if self._cur is None:
            self.select(frame)
        return b

    def select(self, frame=None):
        if frame is None:
            return self._cur
        for t in self._tabs:
            on = t['frame'] is frame
            t['btn'].config(bg=C['bg'] if on else C['card'],
                            fg=C['fg'] if on else C['dim'],
                            font=FB if on else FM)
            if on:
                t['frame'].pack(fill='both', expand=True)
            else:
                t['frame'].pack_forget()
        self._cur = frame
        return frame

    def tab(self, key, option='text'):
        if isinstance(key, int):
            return self._tabs[key]['text'] if 0 <= key < len(self._tabs) else ''
        for t in self._tabs:
            if t['frame'] is key:
                return t['text']
        return ''

    def index(self, _what='end'):
        return len(self._tabs)


class App:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title(f'TR Plus Ultra  —  v{APP_VERSION}')
        self.root.configure(bg=C['bg'])
        self.root.geometry('1080x720')
        set_app_icon(self.root)
        self.prefs = load_prefs()

        self.results = []
        self.not_found = []
        self.imported = []
        self.imp_src = ''
        self.running = False
        self.cancel = False

        load_local_log_dir()
        load_central_dir()
        self._build_ui()
        trim_history()
        log_event('app_open', launcher=globals().get('TRPU_LAUNCHER', ''))
        for t in pending_items():
            self.log('⚠  ' + t, 'WARN')

    # ---------- UI ----------
    def _build_ui(self):
        top = tk.Frame(self.root, bg=C['card'], height=54)
        top.pack(fill='x')
        top.pack_propagate(False)
        tk.Label(top, text='TR Plus Ultra', bg=C['card'], fg=C['fg'],
                 font=('Segoe UI', 14, 'bold')).pack(side='left', padx=16)
        tk.Label(top, text=f'v{APP_VERSION}', bg=C['card'], fg=C['dim'],
                 font=('Segoe UI', 9)).pack(side='left')
        tk.Button(top, text='🔓  เปิดหน้า Login', bg=C['input'], fg=C['fg'], bd=0,
                  font=FM, cursor='hand2', activebackground=C['line'],
                  command=self.open_login).pack(side='right', padx=16, ipadx=12, ipady=5)
        # อยู่แท็บไหนก็แจ้งบั๊กได้ ไม่ต้องเดินไปหาแท็บ
        tk.Button(top, text='🐞  แจ้งบั๊ก', bg=C['input'], fg=C['fg'], bd=0,
                  font=FM, cursor='hand2', activebackground=C['line'],
                  command=self.bug_open_sheet).pack(side='right', ipadx=10, ipady=5)

        style = ttk.Style()
        try:
            style.theme_use('clam')
        except Exception:
            pass
        style.configure('TCombobox', fieldbackground=C['input'], background=C['input'],
                        foreground=C['fg'], arrowcolor=C['dim'], bordercolor=C['line'],
                        lightcolor=C['line'], darkcolor=C['line'], borderwidth=1)
        style.map('TCombobox',
                  fieldbackground=[('readonly', C['input'])],
                  background=[('readonly', C['input'])],
                  foreground=[('readonly', C['fg'])],
                  selectbackground=[('readonly', C['input'])],
                  selectforeground=[('readonly', C['fg'])])
        for k, v in (('*TCombobox*Listbox.background', C['input']),
                     ('*TCombobox*Listbox.foreground', C['fg']),
                     ('*TCombobox*Listbox.selectBackground', C['accent']),
                     ('*TCombobox*Listbox.selectForeground', 'white')):
            try:
                self.root.option_add(k, v)
            except Exception:
                pass

        self.nb = SideTabs(self.root)
        self.tab_search = tk.Frame(self.nb.body, bg=C['bg'])
        self.tab_create = tk.Frame(self.nb.body, bg=C['bg'])
        self.tab_bundle = tk.Frame(self.nb.body, bg=C['bg'])
        self.tab_wr = tk.Frame(self.nb.body, bg=C['bg'])
        self.tab_log = tk.Frame(self.nb.body, bg=C['bg'])
        self.tab_check = tk.Frame(self.nb.body, bg=C['bg'])
        # ซ้าย = แท็บที่ใช้ทำงาน · ขวา = แท็บไว้ดูข้อมูล จะได้ไม่ปนกัน
        self.nb.add(self.tab_search, '🔍  ค้นหา')
        self.nb.add(self.tab_create, '➕  สร้าง Item')
        self.nb.add(self.tab_bundle, '📦  สร้าง Bundle')
        self.nb.add(self.tab_wr, '🎟  WR Master')
        self.nb.add(self.tab_check, '🩺  ตรวจระบบ', side='right')
        self.nb.add(self.tab_log, '📜  Log', side='right')

        self._build_search()
        self._build_create()
        self._build_bundle()
        self._build_wr()
        self._build_log()
        self._build_check()

    def _card(self, parent, title, grow=False):
        outer = tk.LabelFrame(parent, text='  ' + title + '  ', bg=C['bg'], fg=C['dim'],
                              font=('Segoe UI', 9, 'bold'), bd=1,
                              relief='solid', highlightbackground=C['line'])
        outer.pack(fill='both' if grow else 'x', expand=grow, padx=14, pady=(8, 0))
        inner = tk.Frame(outer, bg=C['bg'])
        inner.pack(fill='both' if grow else 'x', expand=grow, padx=12, pady=8)
        return inner

    def _entry(self, parent, width=22):
        e = tk.Entry(parent, bg=C['input'], fg=C['fg'], insertbackground=C['fg'],
                     bd=0, font=FM, width=width, relief='flat',
                     highlightthickness=1, highlightbackground=C['line'],
                     highlightcolor=C['accent'])
        return e

    def _btn(self, parent, text, cmd, primary=False, width=None):
        b = tk.Button(parent, text=text, command=cmd, bd=0, cursor='hand2', font=FB if primary else FM,
                      bg=C['accent'] if primary else C['input'],
                      fg='white' if primary else C['fg'],
                      activebackground='#3d55cf' if primary else C['line'],
                      activeforeground='white')
        if width:
            b.config(width=width)
        return b

    def _build_search(self):
        p = self.tab_search

        s1 = self._card(p, 'ค้นหาทีละตัว')
        tk.Label(s1, text='ItemKind / Aztek Item Id', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9)).grid(row=0, column=0, sticky='w')
        tk.Label(s1, text='ชื่อไอเทม (ไม่บังคับ)', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9)).grid(row=0, column=1, sticky='w', padx=(12, 0))
        self.v_kind = self._entry(s1, 24)
        self.v_kind.grid(row=1, column=0, sticky='w', ipady=4)
        self.v_name = self._entry(s1, 34)
        self.v_name.grid(row=1, column=1, sticky='w', padx=(12, 0), ipady=4)
        self.btn_single = self._btn(s1, '🔍  ค้นหา', self.run_single, primary=True)
        self.btn_single.grid(row=1, column=2, padx=(14, 0), ipadx=18, ipady=3)

        s2 = self._card(p, 'ค้นหาหลายตัวจาก Excel')
        self._btn(s2, '⬇  โหลด Template', self.download_template).grid(row=0, column=0, ipadx=10, ipady=4)
        self._btn(s2, '📂  เลือกไฟล์', self.pick_file).grid(row=0, column=1, padx=8, ipadx=10, ipady=4)
        self.lbl_file = tk.Label(s2, text='ยังไม่ได้เลือกไฟล์', bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9))
        self.lbl_file.grid(row=0, column=2, padx=12, sticky='w')
        self.btn_multi = self._btn(s2, '🔍  ค้นหาทั้งหมด', self.run_multi, primary=True)
        self.btn_multi.grid(row=0, column=3, padx=(8, 0), ipadx=14, ipady=3)
        self.btn_multi.config(state='disabled')

        # ---- ปุ่มกาง/พับ รายการที่จะค้นหา ----
        self.imp_open = False
        self.btn_imp = self._btn(s2, '▸  ดูรายการที่จะค้นหา', self.toggle_imp)
        self.btn_imp.grid(row=1, column=0, columnspan=2, sticky='w', pady=(10, 0),
                          ipadx=8, ipady=3)
        self.btn_imp.config(state='disabled')
        self.lbl_imp = tk.Label(s2, text='', bg=C['bg'], fg=C['warn'], font=('Segoe UI', 9),
                                anchor='w')
        self.lbl_imp.grid(row=1, column=2, columnspan=2, sticky='w', padx=12, pady=(10, 0))

        # ---- แผงรายการ (ซ่อนไว้ก่อน กางเมื่อกดปุ่ม) ----
        self.imp_box = tk.LabelFrame(p, text='  รายการที่จะค้นหา  ', bg=C['bg'], fg=C['dim'],
                                     font=('Segoe UI', 9, 'bold'), bd=1, relief='solid')
        ib = tk.Frame(self.imp_box, bg=C['bg'])
        ib.pack(fill='both', expand=True, padx=12, pady=10)
        ibar = tk.Frame(ib, bg=C['bg'])
        ibar.pack(fill='x')
        self._btn(ibar, '🗑  เอาออกจากรายการ', self.imp_del).pack(side='left', ipadx=8, ipady=3)
        self._btn(ibar, '📋  คัดลอกเลขทั้งหมด', self.imp_copy).pack(side='left', padx=8,
                                                                    ipadx=8, ipady=3)
        self.imp_cnt = tk.Label(ibar, text='', bg=C['bg'], fg=C['dim'], font=FM)
        self.imp_cnt.pack(side='right')

        itw = tk.Frame(ib, bg=C['bg'])
        itw.pack(fill='both', expand=True, pady=(8, 0))
        icols = ('n', 'w', 'kind', 'name', 'dur', 'trade', 'qty')
        self.imp_tree = ttk.Treeview(itw, columns=icols, show='headings',
                                     style='TR.Treeview', height=6, selectmode='extended')
        for c, t, w in (('n', '#', 40), ('w', '', 26), ('kind', 'ItemKind', 95),
                        ('name', 'ชื่อไอเทม', 360), ('dur', 'ระยะเวลา', 90),
                        ('trade', 'แลกเปลี่ยน', 95), ('qty', 'จำนวน', 70)):
            self.imp_tree.heading(c, text=t)
            self.imp_tree.column(c, width=w, anchor='w')
        self.imp_tree.tag_configure('warn', background='#3a3018', foreground='#e3b341')
        isb = ttk.Scrollbar(itw, orient='vertical', command=self.imp_tree.yview)
        self.imp_tree.configure(yscrollcommand=isb.set)
        self.imp_tree.pack(side='left', fill='both', expand=True)
        isb.pack(side='right', fill='y')
        self.imp_why = tk.Label(ib, text='', bg=C['bg'], fg=C['warn'], font=('Segoe UI', 9),
                                anchor='w', justify='left', wraplength=940)
        self.imp_why.pack(fill='x', pady=(6, 0))
        self.imp_tree.bind('<<TreeviewSelect>>', lambda e: self._imp_why())

        # ---- Deep Check ไม่มีช่องกรอกบนหน้าจอแล้ว (v0.8.12) ----
        # ไฟล์ต้นฉบับพกเงื่อนไข (ระยะเวลา/แลกเปลี่ยน/จำนวน) มาในตัวอยู่แล้ว
        # โปรแกรมเปิด deep ให้เองรายตัวตามที่ไฟล์บอก — ช่องกรอกพวกนี้เลยไม่ได้ใช้ เปลืองที่เปล่าๆ
        # ตัวแปรยังอยู่ (ค่า "ไม่กรอง") เพื่อให้ส่วนอื่นที่อ้างถึงทำงานได้เหมือนเดิม
        self.v_deep = tk.BooleanVar(value=False)
        self.v_dur = tk.StringVar(value='any')
        self.v_trade = tk.StringVar(value='any')
        self.v_qty = tk.StringVar(value='')

        s4 = tk.Frame(p, bg=C['bg'])
        self.s_tail = s4
        s4.pack(fill='x', padx=14, pady=(14, 0))
        self.v_exact = tk.BooleanVar(value=bool(self.prefs.get('exact', True)))
        tk.Checkbutton(s4, text='เอาเฉพาะ ItemKind ที่ตรงเป๊ะ', variable=self.v_exact,
                       bg=C['bg'], fg=C['fg'], selectcolor=C['input'], activebackground=C['bg'],
                       activeforeground=C['fg'], font=('Segoe UI', 9), bd=0,
                       highlightthickness=0).pack(side='left')
        self.v_headless = tk.BooleanVar(value=bool(self.prefs.get('headless', False)))
        tk.Checkbutton(s4, text='ซ่อนหน้าต่าง Chrome ตอนทำงาน', variable=self.v_headless,
                       bg=C['bg'], fg=C['dim'], selectcolor=C['input'], activebackground=C['bg'],
                       activeforeground=C['fg'], font=('Segoe UI', 9), bd=0,
                       highlightthickness=0).pack(side='left', padx=(16, 0))
        self.btn_cancel = self._btn(s4, '■  ยกเลิก', self.do_cancel)
        self.btn_cancel.config(state='disabled', fg=C['err'])
        self.btn_cancel.pack(side='right', ipadx=14, ipady=4)

        s5 = tk.Frame(p, bg=C['bg'])
        s5.pack(fill='x', padx=14, pady=(12, 0))
        self.progress = ttk.Progressbar(s5, mode='determinate', maximum=100)
        self.progress.pack(fill='x')
        self.lbl_stat = tk.Label(s5, text='', bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9), anchor='w')
        self.lbl_stat.pack(fill='x', pady=(5, 0))

        # ---- ผลการค้นหา อยู่ในแท็บเดียวกับที่สั่งค้น จะได้ไม่ต้องสลับไปมา ----
        self._build_found(p)

    def _build_found(self, p):
        bar = tk.Frame(p, bg=C['bg'])
        bar.pack(fill='x', padx=14, pady=(14, 6))
        tk.Label(bar, text='ผลการค้นหา', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(side='left')
        self.lbl_count = tk.Label(bar, text='พบ 0 รายการ', bg=C['bg'], fg=C['dim'], font=FM)
        self.lbl_count.pack(side='left', padx=10)
        self._btn(bar, '🗑  ล้าง', self.clear_results).pack(side='right', ipadx=8, ipady=2)
        self._btn(bar, '⬇  CSV', lambda: self.export('csv')).pack(
            side='right', padx=6, ipadx=8, ipady=2)
        self._btn(bar, '⬇  Excel', lambda: self.export('xlsx')).pack(
            side='right', ipadx=8, ipady=2)
        self._btn(bar, '📋  คัดลอก ID ทั้งหมด', self.copy_ids).pack(
            side='right', padx=6, ipadx=8, ipady=2)
        self._btn(bar, '➕  ตัวที่ไม่เจอ → คิวสร้าง Item', self.miss_to_queue,
                  primary=True).pack(side='right', padx=6, ipadx=8, ipady=2)

        wrap = tk.Frame(p, bg=C['bg'])
        wrap.pack(fill='both', expand=True, padx=14, pady=(0, 6))
        style = ttk.Style()
        style.configure('TR.Treeview', background=C['card'], fieldbackground=C['card'],
                        foreground=C['fg'], rowheight=26, borderwidth=0, font=('Segoe UI', 9))
        style.configure('TR.Treeview.Heading', background=C['input'], foreground=C['dim'],
                        font=('Segoe UI', 9, 'bold'), borderwidth=0)
        cols = ('seq', 'id', 'name', 'type', 'kind', 'notes')
        self.tree = ttk.Treeview(wrap, columns=cols, show='headings', height=9,
                                 style='TR.Treeview', selectmode='extended')
        for c, t, w in (('seq', '#', 42), ('id', 'Aztek Item Id', 96), ('name', 'ชื่อ', 330),
                        ('type', 'ประเภท', 80), ('kind', 'ItemKind', 90),
                        ('notes', 'หมายเหตุ', 210)):
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor='w')
        # ซ้ำ = เหลือง (ลำดับเดียวเจอหลายตัว ต้องเลือกเอง) · ไม่เจอ = แดง
        self.tree.tag_configure('dup', background='#3a3018', foreground='#e3b341')
        self.tree.tag_configure('miss', background='#3a1f1e', foreground='#f0736a')
        sb = ttk.Scrollbar(wrap, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')
        tk.Label(p, text='# = ลำดับในรายการที่สั่งค้น  ·  แถวเหลือง = เว็บมีหลายตัว ต้องเลือกเอง'
                         '  ·  แถวแดง = หาไม่เจอ — ส่งเข้าคิวสร้าง Item ได้เลย '
                         '(เลือกแถวไว้ = เอาแค่ที่เลือก · ไม่เลือก = เอาแถวแดงทั้งหมด)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), anchor='w').pack(
            fill='x', padx=14, pady=(0, 10))

    # ========================================================================
    #  แท็บ "สร้าง Item"
    #  แนวคิด: โยนไฟล์ต้นฉบับเข้าไป -> ได้คิวไอเทม -> แก้ได้ทุกช่อง -> ยิงเข้าเว็บ
    #  ไม่ต้องกรอก template เองเหมือนแต่ก่อน
    # ========================================================================
    def _build_create(self):
        p = self.tab_create
        self.cq = []                       # คิวไอเทมที่จะสร้าง
        self.c_running = False
        self.c_cancel = False

        # ---------- ค่าเริ่มต้นที่ใช้กับทุกแถว ----------
        s1 = self._card(p, 'ค่าเริ่มต้น (ใช้ตอนนำเข้าไฟล์ — แก้รายตัวทีหลังได้)')
        self.cv_type = tk.StringVar(value=self.prefs.get('c_type', DEFAULT_TYPE))
        self.cv_suffix = tk.StringVar(value=self.prefs.get('c_suffix', DEFAULT_SUFFIX))
        self.cv_price = tk.StringVar(value=self.prefs.get('c_price', ''))
        self.cv_mail = tk.StringVar(value=self.prefs.get('c_mail', ''))

        for col, (lbl, var, w) in enumerate((
                ('ประเภทไอเทม', self.cv_type, 14),
                ('คำต่อท้ายชื่อ', self.cv_suffix, 12),
                ('Price', self.cv_price, 10),
                ('หัวข้อจดหมาย', self.cv_mail, 20))):
            tk.Label(s1, text=lbl, bg=C['bg'], fg=C['dim'],
                     font=('Segoe UI', 9)).grid(row=0, column=col * 2, sticky='w', padx=(0 if col == 0 else 14, 6))
            e = self._entry(s1, width=w)
            e.config(textvariable=var)
            e.grid(row=0, column=col * 2 + 1, sticky='w', ipady=3)

        tk.Label(s1, text='คำต่อท้ายกับ “(X วัน)” จะเติมให้เฉพาะชื่อที่มีคำว่า “ชิ้น” เท่านั้น — '
                          'จำนวนชิ้นจะถูก “ตัดออกจากชื่อ” แต่ ชม. ยังอยู่ในชื่อเหมือนเดิม',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), wraplength=940,
                 justify='left').grid(row=1, column=0, columnspan=8, sticky='w', pady=(8, 0))

        # ---------- คิว ----------
        s2 = self._card(p, 'คิวไอเทมที่จะสร้าง', grow=True)
        bar = tk.Frame(s2, bg=C['bg'])
        bar.pack(fill='x')
        self._btn(bar, '📂  นำเข้าไฟล์ต้นฉบับ (.xlsx)', self.c_import,
                  primary=True).pack(side='left', ipadx=10, ipady=4)
        self._btn(bar, '➕  เพิ่มเอง', self.c_add).pack(side='left', padx=(8, 0), ipadx=8, ipady=4)
        self._btn(bar, '✏  แก้ไข', self.c_edit).pack(side='left', padx=(8, 0), ipadx=8, ipady=4)
        self._btn(bar, '🗑  ลบที่เลือก', self.c_del).pack(side='left', padx=(8, 0), ipadx=8, ipady=4)
        self._btn(bar, 'ล้างคิว', self.c_clear).pack(side='left', padx=(8, 0), ipadx=8, ipady=4)
        self.c_count = tk.Label(bar, text='คิวว่าง', bg=C['bg'], fg=C['dim'], font=FM)
        self.c_count.pack(side='right')

        tw = tk.Frame(s2, bg=C['bg'])
        tw.pack(fill='both', expand=True, pady=(10, 0))
        cols = ('n', 'w', 'kind', 'name', 'type', 'price', 'dur', 'qty', 'trade', 'img')
        self.c_tree = ttk.Treeview(tw, columns=cols, show='headings',
                                   style='TR.Treeview', height=6)
        for c, t, w in (('n', '#', 38), ('w', '', 26), ('kind', 'ItemKind', 88),
                        ('name', 'ชื่อที่จะกรอกลงเว็บ', 280), ('type', 'ประเภท', 85),
                        ('price', 'Price', 60), ('dur', 'ระยะเวลา', 78),
                        ('qty', 'จำนวน', 58), ('trade', 'แลกเปลี่ยน', 82),
                        ('img', 'รูป', 120)):
            self.c_tree.heading(c, text=t)
            self.c_tree.column(c, width=w, anchor='w')
        self.c_tree.tag_configure('warn', background='#3a3018', foreground='#e3b341')
        csb = ttk.Scrollbar(tw, orient='vertical', command=self.c_tree.yview)
        self.c_tree.configure(yscrollcommand=csb.set)
        self.c_tree.pack(side='left', fill='both', expand=True)
        csb.pack(side='right', fill='y')
        self.c_tree.bind('<Double-1>', lambda e: self.c_edit())
        self.c_tree.bind('<<TreeviewSelect>>', lambda e: self._c_why())

        self.c_warn = tk.Label(s2, text='', bg=C['bg'], fg=C['warn'], font=('Segoe UI', 9),
                               anchor='w', justify='left', wraplength=940)
        self.c_warn.pack(fill='x', pady=(8, 0))

        # ---------- ลงมือ ----------
        s3 = self._card(p, 'ลงมือสร้าง')
        self.cv_do = tk.BooleanVar(value=False)
        cb = tk.Checkbutton(s3, variable=self.cv_do, command=self._c_do_changed,
                            text='กดปุ่ม “สร้าง Item” จริง',
                            bg=C['bg'], fg=C['fg'], selectcolor=C['input'],
                            activebackground=C['bg'], activeforeground=C['fg'],
                            font=FB, bd=0, highlightthickness=0)
        cb.grid(row=0, column=0, sticky='w')
        self.c_mode = tk.Label(s3, text='', bg=C['bg'], fg=C['warn'], font=FM)
        self.c_mode.grid(row=0, column=1, sticky='w', padx=(10, 0))

        self.cv_hold = tk.StringVar(value=self.prefs.get('c_hold', '3'))
        tk.Label(s3, text='โหมดทดสอบ: กรอกเสร็จแล้วค้างหน้าไว้ให้ดู (วินาที)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9)
                 ).grid(row=1, column=0, sticky='w', pady=(8, 0))
        eh = self._entry(s3, width=6)
        eh.config(textvariable=self.cv_hold)
        eh.grid(row=1, column=1, sticky='w', padx=(10, 0), pady=(8, 0), ipady=3)

        run = tk.Frame(s3, bg=C['bg'])
        run.grid(row=2, column=0, columnspan=3, sticky='w', pady=(12, 0))
        self.c_btn_run = self._btn(run, '▶  เริ่มทำงาน', self.c_start, primary=True)
        self.c_btn_run.pack(side='left', ipadx=18, ipady=5)
        self.c_btn_stop = self._btn(run, '■  ยกเลิก', self.c_stop)
        self.c_btn_stop.config(state='disabled')
        self.c_btn_stop.pack(side='left', padx=(8, 0), ipadx=12, ipady=5)

        # ---- ไอเทมที่สร้างสำเร็จ + เลข Aztek Item Id ที่เว็บออกให้ ----
        self.made_items = []
        ib = tk.Frame(p, bg=C['bg'])
        ib.pack(fill='x', padx=14, pady=(10, 4))
        iw = tk.Frame(p, bg=C['bg'])
        iw.pack(fill='both', expand=True, padx=14, pady=(0, 12))
        tk.Label(ib, text='Item ที่สร้างแล้ว', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(side='left')
        self.lbl_item = tk.Label(ib, text='ยังไม่ได้สร้าง', bg=C['bg'], fg=C['dim'],
                                 font=('Segoe UI', 9))
        self.lbl_item.pack(side='left', padx=10)
        self._btn(ib, '📋  คัดลอก Aztek Item Id', self.copy_made_items).pack(
            side='right', ipadx=8, ipady=2)
        tk.Label(ib, text='(เลือกแถวที่ต้องการก่อน · ไม่เลือก = เอาทั้งหมด)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8)).pack(
                     side='right', padx=8)
        ic = ('no', 'aid', 'kind', 'iname', 'pic', 'at', 'note')
        self.tree_item = ttk.Treeview(iw, columns=ic, show='headings', height=6,
                                      style='TR.Treeview', selectmode='extended')
        for c, t, w in (('no', '#', 42), ('aid', 'Aztek Item Id', 110),
                        ('kind', 'ItemKind', 90), ('iname', 'ชื่อไอเทม', 330),
                        ('pic', 'รูป', 44), ('at', 'เวลา', 78),
                        ('note', 'หมายเหตุ', 150)):
            self.tree_item.heading(c, text=t)
            self.tree_item.column(c, width=w, anchor='w')
        self.tree_item.tag_configure('bad', background='#3a1f1e', foreground='#f0736a')
        isb2 = ttk.Scrollbar(iw, orient='vertical', command=self.tree_item.yview)
        self.tree_item.configure(yscrollcommand=isb2.set)
        self.tree_item.pack(side='left', fill='both', expand=True)
        isb2.pack(side='right', fill='y')

        self._c_do_changed()
        self._c_refresh()

    def _c_do_changed(self):
        if self.cv_do.get():
            self.c_mode.config(text='⚠  จะสร้างไอเทมจริงบนเว็บ', fg=C['err'])
        else:
            self.c_mode.config(text='โหมดทดสอบ — กรอกฟอร์มให้ดูเฉยๆ ไม่กดสร้าง', fg=C['ok'])

    def _c_why(self):
        """คลิกแถวเหลืองในคิว แล้วบอกว่าเตือนเพราะอะไร"""
        i = self._c_sel()
        if i is None or not (0 <= i < len(self.cq)):
            return
        w = self.cq[i].get('warn') or []
        self.c_warn.config(text=('⚠  แถว #%d : %s' % (i + 1, ' · '.join(w))) if w else '',
                           fg=C['warn'])

    def _c_refresh(self):
        # คิดใหม่ทุกครั้ง — แก้ ItemKind หรือชื่อแล้วคำเตือนจะหายไปเอง
        nwarn = mark_suspects(self.cq)
        self.c_tree.delete(*self.c_tree.get_children())
        for i, d in enumerate(self.cq, 1):
            self.c_tree.insert('', 'end', tags=('warn',) if d.get('warn') else (), values=(
                i, '⚠' if d.get('warn') else '', d['kind'], d['name'], d['type'],
                d['price'] or '—', (d['dur'] + ' วัน') if d['dur'] else 'ถาวร',
                d['qty'] or '—', 'ได้' if d['trade'] else 'ไม่ได้',
                self._c_img_cell(d)))
        self.c_count.config(text=('คิวว่าง' if not self.cq else f'ในคิว {len(self.cq)} ไอเทม'))
        self.c_warn.config(
            text=(f'⚠  มี {nwarn} แถวหน้าตาไม่เหมือนไอเทมทั่วไป (สีเหลือง) — '
                  'คลิกที่แถวเพื่อดูเหตุผล · ดับเบิลคลิกเพื่อแก้ · ไม่เอาก็ลบออกจากคิวได้'
                  if nwarn else ''), fg=C['warn'])

    @staticmethod
    def _c_img_cell(d):
        """ช่อง "รูป" ในคิว — บอกได้ทันทีว่าตัวไหนยังไม่มีรูป / รูปเสีย"""
        p = str(d.get('img') or '').strip()
        if not p:
            return '—'
        good, _ = check_item_image(p)
        return ('🖼 ' if good else '⚠ ') + os.path.basename(p)

    def _c_sel(self):
        s = self.c_tree.selection()
        if not s:
            return None
        return self.c_tree.index(s[0])

    # ---------- นำเข้า / แก้ไขคิว ----------
    def c_import(self):
        if self.c_running:
            return
        path = filedialog.askopenfilename(title='เลือกไฟล์ต้นฉบับ',
                                          filetypes=[('Excel', '*.xlsx *.xlsm'), ('ทุกไฟล์', '*.*')])
        if not path:
            return
        dlg = ImportDialog(self.root, path)
        if not dlg.result:
            return
        self.save_now()
        t = self.cv_type.get().strip() or DEFAULT_TYPE
        sfx = self.cv_suffix.get().strip()
        pr = self.cv_price.get().strip()
        ml = self.cv_mail.get().strip()
        got = [row_to_item(r, t, sfx, pr, ml) for r in dlg.result]
        self.cq.extend(got)
        self._c_refresh()
        self.log(f'นำเข้าเข้าคิวสร้างไอเทม {len(got)} รายการ (รวม {len(self.cq)})', 'OK')
        log_event('create_import', count=len(got), total=len(self.cq),
                  file=os.path.basename(path))

    def c_add(self):
        if self.c_running:
            return
        d = {'kind': '', 'name': '', 'cname': '', 'desc': '',
             'type': self.cv_type.get().strip() or DEFAULT_TYPE,
             'price': self.cv_price.get().strip(), 'dur': '',
             'mail': self.cv_mail.get().strip(), 'qty': '1',
             'trade': False, 'web': True, 'img': ''}
        if self._c_form(d, 'เพิ่มไอเทมเข้าคิว'):
            self.cq.append(d)
            self._c_refresh()

    def c_edit(self):
        if self.c_running:
            return
        i = self._c_sel()
        if i is None:
            return messagebox.showinfo('แก้ไข', 'เลือกแถวในคิวก่อนนะ')
        d = dict(self.cq[i])
        if self._c_form(d, f'แก้ไขไอเทม #{i + 1}'):
            self.cq[i] = d
            self._c_refresh()

    def c_del(self):
        i = self._c_sel()
        if i is None:
            return
        del self.cq[i]
        self._c_refresh()

    def c_clear(self):
        if self.c_running or not self.cq:
            return
        if messagebox.askyesno('ล้างคิว', f'ลบทั้ง {len(self.cq)} รายการออกจากคิว?'):
            self.cq = []
            self._c_refresh()

    def _c_form(self, d, title):
        """หน้าต่างแก้ไขไอเทม — แก้ได้ทุกช่อง คืน True ถ้ากดบันทึก"""
        top = tk.Toplevel(self.root)
        top.title(title)
        top.configure(bg=C['bg'])
        top.transient(self.root)
        top.grab_set()
        try:    # วางกลางหน้าต่างหลัก ไม่ให้ไปโผล่มุมจอ
            self.root.update_idletasks()
            x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - 660) // 2)
            y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - 600) // 2)
            top.geometry('660x600+%d+%d' % (x, y))
        except Exception:
            top.geometry('660x600')
        ok = {'v': False}

        body = tk.Frame(top, bg=C['bg'])
        body.pack(fill='both', expand=True, padx=18, pady=14)

        fields = (('ชื่อไอเทม (ที่จะกรอกลงเว็บ)', 'name', 48),
                  ('ItemKind / game_item_id', 'kind', 20),
                  ('ประเภทไอเทม', 'type', 20),
                  ('Price', 'price', 12),
                  ('ระยะเวลา (วัน) — เว้นว่าง = ถาวร', 'dur', 12),
                  ('จำนวน (แสดงบนเว็บ)', 'qty', 12),
                  ('หัวข้อจดหมาย', 'mail', 32),
                  ('คำอธิบาย', 'desc', 48))
        vs = {}
        for r, (lbl, key, w) in enumerate(fields):
            tk.Label(body, text=lbl, bg=C['bg'], fg=C['dim'],
                     font=('Segoe UI', 9)).grid(row=r, column=0, sticky='w', pady=4)
            v = tk.StringVar(value=str(d.get(key, '') or ''))
            vs[key] = v
            e = self._entry(body, width=w)
            e.config(textvariable=v)
            e.grid(row=r, column=1, sticky='w', padx=(12, 0), ipady=3)

        vt = tk.BooleanVar(value=bool(d.get('trade')))
        vw = tk.BooleanVar(value=bool(d.get('web', True)))
        for r, (txt, var) in enumerate(((' แลกเปลี่ยนได้', vt),
                                        (' เปิดใช้งานการแสดงผลบนเว็บ', vw)),
                                       start=len(fields)):
            tk.Checkbutton(body, text=txt, variable=var, bg=C['bg'], fg=C['fg'],
                           selectcolor=C['input'], activebackground=C['bg'],
                           activeforeground=C['fg'], font=FM, bd=0,
                           highlightthickness=0).grid(row=r, column=1, sticky='w',
                                                      padx=(12, 0), pady=2)

        # ---- รูปภาพไอเทม (เว็บมีช่องอัปโหลด — เลือกไฟล์จากเครื่องได้เลย) ----
        pic = {'v': str(d.get('img') or '')}
        prow = tk.Frame(body, bg=C['bg'])
        prow.grid(row=len(fields) + 2, column=0, columnspan=2, sticky='w', pady=(10, 0))
        tk.Label(prow, text='รูปภาพไอเทม', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9)).pack(side='left')
        lb_pic = tk.Label(prow, text='', bg=C['bg'], fg=C['dim'],
                          font=('Segoe UI', 8), wraplength=300, justify='left')

        def _show_pic():
            if not pic['v']:
                lb_pic.config(text='ยังไม่ได้เลือกรูป — สร้างได้ แต่ไอเทมจะไม่มีรูป',
                              fg=C['dim'])
                return
            good, why = check_item_image(pic['v'])
            lb_pic.config(text=('🖼  ' if good else '⚠  ') + why,
                          fg=C['ok'] if good else C['err'])

        def _pick():
            p = filedialog.askopenfilename(
                parent=top, title='เลือกรูปไอเทม',
                filetypes=[('รูปภาพ (.png .jpg .webp)', '*.png *.jpg *.jpeg *.webp'),
                           ('ทุกไฟล์', '*.*')])
            if not p:
                return
            good, why = check_item_image(p)
            if not good:
                return messagebox.showwarning('รูปนี้ใช้ไม่ได้', why, parent=top)
            pic['v'] = p
            _show_pic()

        def _drop():
            pic['v'] = ''
            _show_pic()

        self._btn(prow, '📁  เลือกรูป…', _pick).pack(side='left', padx=(12, 0),
                                                     ipadx=8, ipady=2)
        self._btn(prow, '✕  เอารูปออก', _drop).pack(side='left', padx=(6, 0),
                                                     ipadx=6, ipady=2)
        lb_pic.pack(side='left', padx=(10, 0))
        _show_pic()
        tk.Label(body, text='เว็บรับ .png / .jpg / .webp ขนาดไม่เกิน %d MB  ·  '
                            'ไม่ใส่ก็สร้างได้ แค่ไอเทมจะไม่มีรูป' % IMG_MAX_MB,
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8)).grid(
                     row=len(fields) + 3, column=0, columnspan=2, sticky='w')

        # ตัวช่วย: ประกอบชื่อใหม่จาก "ชื่อฐาน" + คำต่อท้าย + ระยะเวลา
        hint = tk.Label(body, text='ชื่อฐาน (ตัดจำนวนชิ้นออกแล้ว): ' + (d.get('cname') or '—'),
                        bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8),
                        wraplength=560, justify='left')
        hint.grid(row=len(fields) + 4, column=0, columnspan=2, sticky='w', pady=(10, 0))

        def _rebuild():
            vs['name'].set(make_item_name(d.get('cname') or vs['name'].get(),
                                          self.cv_suffix.get().strip(),
                                          vs['dur'].get().strip()))
        self._btn(body, '↻ ประกอบชื่อใหม่จากชื่อฐาน + คำต่อท้าย + ระยะเวลา', _rebuild
                  ).grid(row=len(fields) + 5, column=0, columnspan=2,
                         sticky='w', pady=(6, 0), ipadx=6, ipady=3)

        foot = tk.Frame(top, bg=C['card'], height=58)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)

        def _save():
            for key, v in vs.items():
                d[key] = v.get().strip()
            d['trade'] = vt.get()
            d['web'] = vw.get()
            d['img'] = pic['v']
            if not d['kind'] or not d['name']:
                return messagebox.showwarning('กรอกไม่ครบ', 'ต้องมี ItemKind และชื่อไอเทม')
            ok['v'] = True
            top.destroy()

        self._btn(foot, 'บันทึก', _save, primary=True).pack(side='right', padx=18, pady=12,
                                                            ipadx=20, ipady=4)
        self._btn(foot, 'ยกเลิก', top.destroy).pack(side='right', pady=12, ipadx=14, ipady=4)
        self.root.wait_window(top)
        return ok['v']

    # ---------- รัน ----------
    def c_stop(self):
        self.c_cancel = True
        self.log('กำลังยกเลิกการสร้างไอเทม...', 'WARN')

    def c_start(self):
        if self.c_running or self.running:
            return messagebox.showinfo('กำลังทำงาน', 'รอให้งานปัจจุบันเสร็จก่อนนะ')
        if not self.cq:
            return messagebox.showwarning('คิวว่าง', 'ยังไม่มีไอเทมในคิว — นำเข้าไฟล์ต้นฉบับก่อน')
        do = self.cv_do.get()
        nwarn = sum(1 for d in self.cq if d.get('warn'))
        nopic = sum(1 for d in self.cq if not check_item_image(d.get('img'))[0])
        if do:
            msg = f'จะสร้างไอเทมจริงบนเว็บ {len(self.cq)} รายการ\n'
            if nopic:
                msg += (f'\n🖼  มี {nopic} รายการที่ยังไม่มีรูป (หรือรูปใช้ไม่ได้)\n'
                        '    จะสร้างต่อให้โดยไม่มีรูป — ใส่รูปทีหลังในเว็บได้\n')
            if nwarn:
                msg += (f'\n⚠  ในคิวมี {nwarn} แถวที่หน้าตาไม่เหมือนไอเทมทั่วไป (แถวสีเหลือง)\n'
                        '    ถ้ายังไม่ได้ดู แนะนำให้กดยกเลิกแล้วไล่ดูก่อน\n')
            msg += '\nตรวจชื่อ/ItemKind ในคิวเรียบร้อยแล้วใช่ไหม?'
            if not messagebox.askyesno('ยืนยัน', msg):
                return
        self.save_now()
        self.c_running = True
        self.c_cancel = False
        self.c_btn_run.config(state='disabled')
        self.c_btn_stop.config(state='normal')
        self.nb.select(self.tab_log)
        self.log('=' * 46, 'STEP')
        self.log(('เริ่มสร้างไอเทมจริง ' if do else 'เริ่มทดสอบกรอกฟอร์ม ')
                 + f'{len(self.cq)} รายการ', 'STEP')
        if nopic:
            self.log('ยังไม่มีรูป %d จาก %d รายการ — ตัวที่ไม่มีจะสร้างโดยไม่ใส่รูป'
                     % (nopic, len(self.cq)), 'WARN')
        log_event('create_start', count=len(self.cq), commit=bool(do),
                  no_image=nopic)
        threading.Thread(target=self._c_thread, args=(list(self.cq), do), daemon=True).start()

    def _c_thread(self, queue_rows, do):
        try:
            asyncio.run(self._c_work(queue_rows, do))
        except Exception as ex:
            log_event('error', where='create', message=str(ex)[:300])
            self.log('ผิดพลาด: ' + str(ex), 'ERR')
            self.log(traceback.format_exc(), 'ERR')
        finally:
            self.c_running = False

            def _rst():
                self.c_btn_run.config(state='normal')
                self.c_btn_stop.config(state='disabled')
                # ผลของการสร้างไอเทมอยู่ในแท็บของตัวเอง — เด้งกลับไปให้ดูเลย
                if self.made_items:
                    self.nb.select(self.tab_create)
            self.root.after(0, _rst)

    async def _c_work(self, rows, do):
        async with async_playwright() as pw:
            browser = await pw.chromium.launch_persistent_context(**launch_kwargs(False))
            page = browser.pages[0] if browser.pages else await browser.new_page()
            try:
                try:
                    hold = max(0, int(float(self.cv_hold.get() or 0)))
                except Exception:
                    hold = 3
                okc = errc = 0
                for i, d in enumerate(rows, 1):
                    if self.c_cancel:
                        self.log('ยกเลิกแล้ว', 'WARN')
                        break
                    self.set_progress(i - 1, len(rows), d['name'][:28])
                    self.log(f'[{i}/{len(rows)}] {d["name"]}  (ItemKind={d["kind"]})', 'STEP')
                    try:
                        if await self._c_one(page, d, do, hold):
                            okc += 1
                        else:
                            errc += 1
                    except Exception as ex:
                        errc += 1
                        self.log('   ✗ ' + str(ex)[:160], 'ERR')
                        log_event('error', where='create_one', kind_id=d.get('kind'),
                                  message=str(ex)[:200])
                self.set_progress(len(rows), len(rows), 'เสร็จ')
                self.log(f'จบ — สำเร็จ {okc} · ไม่ผ่าน {errc}', 'OK' if not errc else 'WARN')
                log_event('create_done', ok=okc, fail=errc, commit=bool(do))
            finally:
                try:
                    await browser.close()
                except Exception:
                    pass

    async def _c_fill(self, page, key, value, label=''):
        """กรอกช่องเดียว: ลอง id ก่อน ถ้าไม่เจอค่อยหาโดยอ้างข้อความข้างๆ"""
        if value is None or str(value) == '':
            return True
        sel = SEL_CREATE.get(key, '')
        r = await page.evaluate(JS_SET_VALUE, [sel, str(value)])
        if r != 'ok' and label:
            r = await page.evaluate(JS_SET_BY_LABEL, [label, str(value)])
        if r == 'ok':
            self.log(f'   · {label or key} = {str(value)[:40]}', 'INFO')
            return True
        self.log(f'   ! ไม่เจอช่อง {label or key} ({sel})', 'WARN')
        return False

    async def _click_confirm(self, page, hits=None, tries=12):
        """กดปุ่มยืนยันในป๊อปอัปที่เด้งขึ้นมาหลังกดปุ่มสร้าง

        ป๊อปอัปโผล่ช้ากว่าการคลิกนิดนึง เลยต้องวนดูหลายรอบ ไม่ใช่เช็กครั้งเดียวแล้วเลิก
        ถ้าเว็บตอบกลับมาแล้ว (hits) แปลว่ารอบนี้ไม่มีป๊อปอัป — เลิกรอทันที ไม่ถ่วงเวลา
        """
        last = 'ไม่มีป๊อปอัป'
        for _ in range(tries):
            if hits:
                return False
            r = str(await page.evaluate(JS_CONFIRM_POPUP, [CONFIRM_YES, CONFIRM_NO]))
            if r.startswith('ok'):
                self.log('   · กดยืนยันในป๊อปอัปแล้ว (%s)' % r.split('|', 1)[1], 'INFO')
                return True
            last = r
            await page.wait_for_timeout(250)
        if last != 'ไม่มีป๊อปอัป':
            self.log('   ! ' + last, 'WARN')
        return False

    async def _c_read_type(self, page):
        try:
            return str(await page.evaluate(
                JS_READ_TYPE, [SEL_CREATE['type_label'], 'game_item_type']) or '').strip()
        except Exception:
            return ''

    async def _c_type(self, page, want):
        """ตั้งประเภทไอเทม (game_item_type)

        หน้าสร้างไอเทมของเว็บ "ล็อก" ช่องนี้ไว้เป็น GENERAL กดเลือกไม่ได้เลย
        (คนกดเองก็ไม่ได้) ถ้าค่าบนหน้าตรงกับที่เราต้องการอยู่แล้ว = เรียบร้อย
        ไม่ใช่ข้อผิดพลาด เลยต้องเช็กก่อนแล้วค่อยลองเลือก
        """
        want_s = str(want).strip()
        cur = await self._c_read_type(page)
        if cur and cur.lower() == want_s.lower():
            self.log('   · ประเภทไอเทม = %s  (เว็บล็อกค่านี้ไว้ เลือกไม่ได้ '
                     'แต่ตรงกับที่ต้องการอยู่แล้ว)' % cur, 'INFO')
            return True
        r = await page.evaluate(JS_PICK_TYPE, [want_s, SEL_CREATE['type_ph']])
        if r == 'ok' or await self._c_pick_combo(page, want_s):
            self.log('   · ประเภทไอเทม = %s' % want_s, 'INFO')
            return True
        cur = await self._c_read_type(page)
        if cur and cur.lower() == want_s.lower():
            self.log('   · ประเภทไอเทม = %s  (เว็บล็อกค่านี้ไว้ เลือกไม่ได้ '
                     'แต่ตรงกับที่ต้องการอยู่แล้ว)' % cur, 'INFO')
            return True
        self.log('   ! ประเภทไอเทมไม่ตรง — อยากได้ “%s” แต่หน้าเว็บเป็น “%s” '
                 'ต้องแก้เองบนเว็บ' % (want_s, cur or 'อ่านค่าไม่ได้'), 'WARN')
        return False

    # ==================================================================
    #  WR Master — กรอกหน้าสร้าง Item Code
    # ==================================================================
    async def _w_text(self, page, label, side, value, what=''):
        r = str(await page.evaluate(JS_CODE_TEXT, [label, side, 'set', str(value)]))
        if r == 'ok':
            self.log('   · %s = %s' % (what or label, str(value)[:44]), 'INFO')
            return True
        self.log('   ! กรอก “%s” ไม่ได้ (%s)' % (what or label, r), 'WARN')
        return False

    async def _w_read(self, page, label, side):
        r = str(await page.evaluate(JS_CODE_TEXT, [label, side, 'read', '']))
        return r.split('|', 1)[1] if r.startswith('ok|') else None

    async def _w_select(self, page, label, side, value, what=''):
        cur = str(await page.evaluate(JS_CODE_SELECT, [label, side, 'read', '']))
        if cur.startswith('ok|') and value.lower() in cur.split('|', 1)[1].lower():
            self.log('   · %s = %s  (เป็นค่านี้อยู่แล้ว)'
                     % (what or label, cur.split('|', 1)[1]), 'INFO')
            return True
        r = str(await page.evaluate(JS_CODE_SELECT, [label, side, 'set', value]))
        if r == 'ok':
            self.log('   · %s = %s' % (what or label, value), 'INFO')
            return True
        self.log('   ! เลือก “%s” เป็น %s ไม่ได้ (%s)' % (what or label, value, r), 'WARN')
        return False

    async def _w_switch(self, page, label, side, want, what=''):
        """สวิตช์ของหน้านี้เป็น React เหมือนหน้าสร้าง Item — กดแล้วต้องรอค่อยอ่าน"""
        want = bool(want)

        async def read():
            r = str(await page.evaluate(JS_CODE_SWITCH, [label, side, 'read']))
            return (r.split('|', 1)[1] == '1') if r.startswith('ok|') else None

        cur = await read()
        if cur is None:
            self.log('   ! ไม่เจอสวิตช์ “%s”' % (what or label), 'WARN')
            return False
        if cur == want:
            self.log('   · %s = %s  (เป็นแบบนี้อยู่แล้ว)'
                     % (what or label, 'เปิด' if want else 'ปิด'), 'INFO')
            return True
        await page.evaluate(JS_CODE_SWITCH, [label, side, 'click'])
        for _ in range(6):
            await page.wait_for_timeout(200)
            if (await read()) == want:
                self.log('   · %s = %s  (เปลี่ยนให้แล้ว)'
                         % (what or label, 'เปิด' if want else 'ปิด'), 'INFO')
                await page.wait_for_timeout(200)     # รอช่องที่ซ่อนอยู่โผล่มา
                return True
        self.log('   ! ตั้งสวิตช์ “%s” ไม่ได้' % (what or label), 'WARN')
        return False

    async def _w_bundle(self, page, bid):
        """เลือก Bundle จากเลขในชีท — เปิดป๊อปอัป ค้นด้วยเลข แล้วคลิกแถวที่ตรงเป๊ะ"""
        bid = str(bid or '').strip()
        if not bid:
            self.log('   ! ไม่มีเลข Bundle ในชีท — ต้องเลือกเองบนเว็บ', 'WARN')
            return False
        r = str(await page.evaluate(JS_CODE_BUNDLE, ['open', '']))
        if r != 'ok':
            self.log('   ! ' + r, 'WARN')
            return False
        await page.wait_for_timeout(500)
        r = str(await page.evaluate(JS_CODE_BUNDLE, ['search', bid]))
        if r != 'ok':
            self.log('   ! ' + r, 'WARN')
            return False
        for _ in range(10):                       # รอรายการโหลด แล้วค่อยคลิก
            await page.wait_for_timeout(400)
            r = str(await page.evaluate(JS_CODE_BUNDLE, ['pick', bid]))
            if r.startswith('ok'):
                await page.wait_for_timeout(400)
                self.log('   · Bundle = %s' % (r.split('|', 1)[1] if '|' in r else bid),
                         'INFO')
                return True
        self.log('   ! %s' % r, 'WARN')
        return False

    async def _w_one(self, page, d, do, hold):
        """กรอก Item Code หนึ่งตัวให้ครบทุกช่องตามที่ทีมใช้จริง"""
        await page.goto(ITEMCODE_CREATE_URL, wait_until='domcontentloaded', timeout=45000)
        await page.wait_for_timeout(2000)
        if any(k in page.url.lower() for k in ('login', 'signin', 'auth')):
            self.log('   ✗ ยังไม่ได้ล็อกอิน — กด “เปิดหน้า Login” ด้านบนก่อน', 'ERR')
            return False

        cap = str(d.get('cap') or '').strip()
        limited = bool(cap)

        # ---------- ฝั่งซ้าย ----------
        await self._w_text(page, SEL_CODE['name_th'], 'left', d['name'], 'ชื่อ Item Code (ไทย)')
        await self._w_text(page, SEL_CODE['name_en'], 'left', d['name'], 'ชื่อ Item Code (อังกฤษ)')
        await self._w_select(page, SEL_CODE['kind'], 'left', SEL_CODE['kind_val'], 'ประเภท')
        await self._w_text(page, SEL_CODE['per_user'], 'left', d.get('per_user', '1'),
                           'จำนวนการใช้งานต่อ 1 User (ซ้าย)')
        if d.get('start'):
            await self._w_text(page, SEL_CODE['start'], 'left', d['start'], 'เวลาเริ่มใช้งาน')
        if d.get('end'):
            await self._w_text(page, SEL_CODE['end'], 'left', d['end'], 'เวลาสิ้นสุด')
        await self._w_switch(page, SEL_CODE['limit_sw'], 'left', limited, 'จำกัดจำนวน (ซ้าย)')
        if limited:
            await self._w_text(page, SEL_CODE['limit_max'], 'left', cap,
                               'จำนวนครั้งที่สามารถใช้งานได้')
            await self._w_text(page, SEL_CODE['limit_left'], 'left', cap,
                               'จำนวนคงเหลือ (ซ้าย)')

        # ---------- ฝั่งขวา (ของรางวัล) ----------
        await self._w_text(page, SEL_CODE['rw_th'], 'right', d['code'], 'ชื่อรางวัล (ไทย)')
        await self._w_text(page, SEL_CODE['rw_en'], 'right', d['code'], 'ชื่อรางวัล (อังกฤษ)')
        await self._w_text(page, SEL_CODE['per_user'], 'right', d.get('per_user', '1'),
                           'จำนวนการใช้งานต่อ 1 User (ขวา)')
        await self._w_switch(page, SEL_CODE['rw_sw'], 'right', limited, 'จำกัดจำนวน Code')
        if limited:
            await self._w_text(page, SEL_CODE['rw_max'], 'right', cap, 'จำนวนซอง')
            await self._w_text(page, SEL_CODE['rw_left'], 'right', cap, 'จำนวนคงเหลือ (ขวา)')
        await self._w_select(page, SEL_CODE['code_type'], 'right', SEL_CODE['code_fix'],
                             'ประเภทของ Code')
        await self._w_text(page, SEL_CODE['code_list'], 'right', d['code'], 'รายการ Code')
        await self._w_bundle(page, d.get('bundle'))

        # ---------- ตรวจว่าค่าเข้าไปจริง ----------
        bad = []
        for lbl, side, want, what in (
                (SEL_CODE['name_th'], 'left', d['name'], 'ชื่อ Item Code'),
                (SEL_CODE['code_list'], 'right', d['code'], 'รายการ Code')):
            got = await self._w_read(page, lbl, side)
            if (got or '').strip() != str(want).strip():
                bad.append('%s: อยากได้ “%s” แต่ในฟอร์มเป็น “%s”' % (what, want, got))
        for b in bad:
            self.log('   ! ' + b, 'WARN')

        if not do:
            self.log('   ✓ กรอกฟอร์มครบแล้ว (โหมดทดสอบ — ไม่กดสร้าง)', 'OK')
            if hold:
                await page.wait_for_timeout(hold * 1000)
            self.add_made_code(d, '', 'กรอกฟอร์มแล้ว (โหมดทดสอบ — ไม่ได้สร้าง)')
            return not bad

        # ---------- กดสร้างจริง + ป๊อปอัปยืนยัน ----------
        btn = page.locator('button:has-text("%s")' % SEL_CODE['submit']).last
        if await btn.count() == 0:
            self.log('   ✗ ไม่เจอปุ่ม “%s”' % SEL_CODE['submit'], 'ERR')
            return False
        hits = []

        def _grab(resp):
            try:
                if resp.request.method in ('POST', 'PUT', 'PATCH') \
                        and 'code' in resp.url.lower():
                    hits.append(resp)
            except Exception:
                pass
        page.on('response', _grab)
        try:
            await btn.click(timeout=10000)
            await page.wait_for_timeout(600)
            await self._click_confirm(page, hits)
            waited = 0
            while waited < 20000 and not hits:
                await page.wait_for_timeout(300)
                waited += 300
            await page.wait_for_timeout(1500)
        finally:
            try:
                page.remove_listener('response', _grab)
            except Exception:
                pass

        code_http = 0
        cid = ''
        for resp in hits:
            try:
                code_http = resp.status
                body = await resp.json()
            except Exception:
                body = None
            cid = cid or _dig_id(body)
        good = (code_http == 0) or (200 <= code_http < 300)
        if good:
            self.log('   ✓ สร้างแล้ว%s%s'
                     % ('  ·  Item Code Id = %s' % cid if cid else '',
                        '  (HTTP %d)' % code_http if code_http else ''), 'OK')
        else:
            self.log('   ✗ เว็บตอบ HTTP %d — ยังไม่ได้สร้าง' % code_http, 'ERR')
        self.add_made_code(d, cid if good else '',
                           'สร้างแล้ว' if good else 'ไม่สำเร็จ (HTTP %d)' % code_http,
                           ok=good)
        log_event('create_itemcode', code=d['code'], name=d['name'],
                  bundle=d.get('bundle'), ok=bool(good), http=code_http)
        return good

    async def _c_one(self, page, d, do, hold):
        await page.goto(ITEM_CREATE_URL, wait_until='domcontentloaded', timeout=45000)
        await page.wait_for_timeout(2000)
        if any(k in page.url.lower() for k in ('login', 'signin', 'auth')):
            self.log('   ✗ ยังไม่ได้ล็อกอิน — กด “เปิดหน้า Login” ด้านบนก่อน', 'ERR')
            return False
        try:
            await page.wait_for_selector(SEL_CREATE['name'], timeout=15000)
        except Exception:
            self.log('   ! ไม่เจอฟอร์มตาม id เดิม — จะลองหาโดยอ้างชื่อช่องแทน', 'WARN')

        # [1] รายละเอียดไอเทม
        await self._c_fill(page, 'name', d['name'], 'ชื่อไอเทม')
        await self._c_fill(page, 'desc', d.get('desc', ''), 'คำอธิบาย')
        if d.get('type'):
            await self._c_type(page, d['type'])

        # [2] พารามิเตอร์ส่งเข้าเกม
        await self._c_fill(page, 'kind', d['kind'], 'game_item_id (ItemKind)')
        await self._c_fill(page, 'price', d.get('price', ''), 'Price')
        await self._c_fill(page, 'duration', d.get('dur', ''), 'ระยะเวลาไอเทม (วัน)')
        await self._c_fill(page, 'mail', d.get('mail', ''), 'หัวข้อจดหมาย')
        await self._c_switch(page, SEL_CREATE['trade_label'], bool(d.get('trade')))

        # [3] พารามิเตอร์แสดงบนเว็บ
        await self._c_switch(page, SEL_CREATE['web_label'], bool(d.get('web', True)))
        await page.wait_for_timeout(300)
        await self._c_fill(page, 'qty', d.get('qty', ''), 'จำนวน')

        # [3.5] รูปภาพไอเทม — ไม่มีรูปก็สร้างต่อ แต่ต้องเตือนไว้ใน Log
        pic_ok = await self._c_image(page, d)

        # ตรวจว่าค่าเข้าไปจริง
        got = await page.evaluate(JS_READ_FORM, SEL_CREATE)
        bad = []
        for k, want in (('name', d['name']), ('kind', d['kind'])):
            if (got.get(k) or '').strip() != str(want).strip():
                bad.append(f'{k}: อยากได้ “{want}” แต่ในฟอร์มเป็น “{got.get(k)}”')
        if bad:
            for b in bad:
                self.log('   ! ' + b, 'WARN')

        if not do:
            self.log('   ✓ กรอกฟอร์มครบแล้ว (โหมดทดสอบ — ไม่กดสร้าง)', 'OK')
            if hold:
                await page.wait_for_timeout(hold * 1000)
            self.add_made_item(d, '', 'กรอกฟอร์มแล้ว (โหมดทดสอบ — ไม่ได้สร้าง)',
                               pic=pic_ok)
            return not bad

        # [4] กดสร้างจริง — ดักคำตอบของเว็บไว้ เอา Aztek Item Id ออกมาให้ได้
        btn = page.locator(f'button:has-text("{SEL_CREATE["submit"]}")').last
        if await btn.count() == 0:
            self.log(f'   ✗ ไม่เจอปุ่ม “{SEL_CREATE["submit"]}”', 'ERR')
            return False

        hits = []

        def _grab(resp):
            try:
                if resp.request.method in ('POST', 'PUT', 'PATCH') \
                        and 'item' in resp.url.lower():
                    hits.append(resp)
            except Exception:
                pass
        page.on('response', _grab)
        try:
            await btn.click(timeout=10000)
            await page.wait_for_timeout(600)
            # เว็บเด้งป๊อปอัป “ยืนยันการสร้าง” ต่ออีกที — ไม่กดยืนยัน = ไม่ได้สร้างจริง
            await self._click_confirm(page, hits)
            waited = 0
            while waited < 20000 and not hits:
                await page.wait_for_timeout(300)
                waited += 300
            await page.wait_for_timeout(1500)
        finally:
            try:
                page.remove_listener('response', _grab)
            except Exception:
                pass

        code = 0
        aid = ''
        for resp in hits:
            try:
                code = resp.status
                body = await resp.json()
            except Exception:
                body = None
            aid = aid or _dig_id(body)
        if not aid:
            try:
                aid = str(await page.evaluate(JS_ITEM_MADE_ID) or '').strip()
            except Exception:
                aid = ''

        good = (code == 0) or (200 <= code < 300)
        if good:
            self.log('   ✓ สร้างแล้ว%s%s'
                     % ('  ·  Aztek Item Id = %s' % aid if aid
                        else '  (เว็บไม่ได้ส่งเลข Aztek Item Id กลับมา)',
                        '  (HTTP %d)' % code if code else ''), 'OK')
        else:
            self.log('   ✗ เว็บตอบ HTTP %d — ยังไม่ได้สร้าง' % code, 'ERR')
        self.add_made_item(d, aid if good else '',
                           'สร้างแล้ว' if good else 'ไม่สำเร็จ (HTTP %d)' % code,
                           pic=pic_ok, ok=good)
        log_event('create_item', kind_id=d['kind'], name=d['name'],
                  ok=bool(good), http=code, aztek_id=aid)
        return good

    async def _c_image(self, page, d):
        """อัปโหลดรูปไอเทม — คืน True ถ้าใส่รูปเข้าไปได้จริง

        ไม่มีรูป หรือใส่ไม่ได้ ก็ไม่ล้มงาน แค่เตือนไว้ใน Log แล้วสร้างต่อ
        """
        path = str(d.get('img') or '').strip()
        if not path:
            self.log('   ! ไม่ได้เลือกรูปให้ไอเทมนี้ — จะสร้างโดยไม่มีรูป', 'WARN')
            return False
        ok, why = check_item_image(path)
        if not ok:
            self.log('   ! รูปใช้ไม่ได้ (' + why + ') — จะสร้างโดยไม่มีรูป', 'WARN')
            return False
        r = await page.evaluate(JS_MARK_FILE, [SEL_CREATE['img_label']])
        if r != 'ok':
            self.log('   ! ' + str(r) + ' — จะสร้างโดยไม่มีรูป', 'WARN')
            return False
        before = await page.evaluate(JS_PIC_STATE, [SEL_CREATE['img_label']]) or {}
        try:
            await page.locator('input[type="file"][data-trpu-pic]').first.set_input_files(
                path, timeout=15000)
        except Exception as ex:
            self.log('   ! ใส่รูปไม่สำเร็จ: ' + str(ex)[:120] + ' — จะสร้างโดยไม่มีรูป', 'WARN')
            return False
        # เว็บอัปโหลดทันทีแล้วล้างช่อง input ทิ้ง — ดู el.files อย่างเดียวไม่พอ
        # ต้องรอแล้วดูหลักฐานฝั่งหน้าเว็บด้วย (รูปตัวอย่างโผล่ / ข้อความอัปโหลดสำเร็จ)
        st = {}
        for _ in range(12):
            await page.wait_for_timeout(500)
            st = await page.evaluate(JS_PIC_STATE, [SEL_CREATE['img_label']]) or {}
            if self._pic_ok(before, st):
                break
        name = st.get('name') or os.path.basename(path)
        if self._pic_ok(before, st):
            self.log('   · รูปภาพไอเทม = %s  (เว็บรับรูปแล้ว)' % name, 'INFO')
            return True
        # ส่งไฟล์เข้าช่องไปแล้ว แต่ดูไม่ออกว่าเว็บรับรึยัง — ห้ามฟันธงว่าไม่รับ
        self.log('   ? ส่งรูป %s เข้าช่องแล้ว แต่หน้าเว็บยังไม่ขึ้นอะไรให้ยืนยัน '
                 '— ดูที่กล่อง “รูปภาพไอเทม” บนเว็บอีกที' % name, 'WARN')
        return False

    @staticmethod
    def _pic_ok(before, after):
        """เว็บรับรูปแล้วรึยัง — ดูหลายทางเพราะแต่ละเว็บโชว์ไม่เหมือนกัน"""
        if not after:
            return False
        if (after.get('n') or 0) > 0:
            return True                                  # ไฟล์ยังอยู่ในช่อง
        if (after.get('imgs') or 0) > (before.get('imgs') or 0):
            return True                                  # รูปตัวอย่างโผล่มาใหม่
        if after.get('src') and after.get('src') != before.get('src'):
            return True                                  # รูปตัวอย่างเปลี่ยนเป็นรูปใหม่
        return bool(after.get('toast')) and not before.get('toast')

    async def _switch_read(self, page, label):
        """อ่านว่าสวิตช์นี้เปิดอยู่ไหม — คืน True/False หรือ None ถ้าไม่เจอสวิตช์"""
        r = str(await page.evaluate(JS_SWITCH, [label, 'read']))
        if not r.startswith('ok'):
            return None
        return r.split('|', 1)[1] == '1'

    async def _c_switch(self, page, label, want):
        """ติ๊ก/ปลดติ๊กสวิตช์ให้เป็นค่าที่ต้องการ

        สำคัญ: เว็บเป็น React — กดแล้วค่าบนหน้าไม่ได้เปลี่ยนทันทีในจังหวะเดียวกัน
        ต้อง "กด แล้วรอ แล้วค่อยอ่านใหม่" ไม่ใช่อ่านทันทีหลังกด
        ไม่งั้นจะหาว่า "กดแล้วค่าไม่เปลี่ยน" ทั้งที่หน้าเว็บเปลี่ยนให้เรียบร้อยแล้ว
        """
        want = bool(want)
        cur = await self._switch_read(page, label)
        if cur is None:
            if not want:
                # ไม่เจอสวิตช์ แต่เราต้องการ "ไม่ติ๊ก" อยู่แล้ว = ตรงตามที่ต้องการ ไม่ใช่บั๊ก
                self.log('   · %s = ปิด  (หน้านี้ไม่มีให้ติ๊ก ค่าเริ่มต้นคือไม่ติ๊กอยู่แล้ว)'
                         % label, 'INFO')
                return True
            self.log('   ! เปิด “%s” ให้ไม่ได้ (ไม่เจอสวิตช์) — ต้องไปติ๊กเองบนเว็บ'
                     % label, 'WARN')
            return False
        if cur == want:
            self.log('   · %s = %s  (เป็นแบบนี้อยู่แล้ว)'
                     % (label, 'เปิด' if want else 'ปิด'), 'INFO')
            return True

        for act in ('click', 'label'):        # กดตัวควบคุมก่อน ไม่ขยับค่อยลองกดที่ป้าย
            r = str(await page.evaluate(JS_SWITCH, [label, act]))
            if not r.startswith('ok'):
                continue
            for _ in range(6):                # รอ React วาดใหม่ก่อนค่อยอ่านซ้ำ
                await page.wait_for_timeout(200)
                if (await self._switch_read(page, label)) == want:
                    self.log('   · %s = %s  (เปลี่ยนให้แล้ว)'
                             % (label, 'เปิด' if want else 'ปิด'), 'INFO')
                    return True
        self.log('   ! ตั้ง “%s” เป็น%s ไม่ได้ — ต้องไปกดเองบนเว็บ'
                 % (label, 'เปิด' if want else 'ปิด'), 'WARN')
        return False

    async def _c_pick_combo(self, page, value):
        """combobox แบบ React — คลิกเปิดแล้วคลิกตัวเลือก"""
        for trig in ('[role="combobox"]', 'text="%s"' % SEL_CREATE['type_ph']):
            try:
                t = page.locator(trig).first
                if await t.count() == 0:
                    continue
                await t.click(timeout=4000)
                await page.wait_for_timeout(400)
                opt = page.locator(f'[role="option"]:has-text("{value}")').first
                if await opt.count() == 0:
                    opt = page.locator(f'text="{value}"').last
                if await opt.count() > 0:
                    await opt.click(timeout=4000)
                    self.log(f'   · ประเภทไอเทม = {value}', 'INFO')
                    return True
            except Exception:
                continue
        self.log(f'   ! เลือกประเภทไอเทม “{value}” ไม่ได้', 'WARN')
        return False

    # ========================================================================
    #  แท็บ "สร้าง Bundle"
    #  โยนไฟล์ต้นฉบับเข้าไป -> ได้บันเดิลพร้อมไอเทมข้างใน -> ติ๊กเลือกอันที่จะสร้าง
    # ========================================================================
    def _build_bundle(self):
        p = self.tab_bundle
        self.bq = []                 # บันเดิลทั้งหมดที่นำเข้ามา
        self.b_running = False
        self.b_cancel = False
        self._b_rowmap = {}          # iid ของ treeview -> ('b', i) หรือ ('i', i, j)
        self.b_open_new = False      # นำเข้ามาใหม่ = หุบไว้ก่อน จะได้เห็นรายชื่อครบๆ

        s1 = self._card(p, 'นำเข้าจากไฟล์ต้นฉบับ')
        bar = tk.Frame(s1, bg=C['bg'])
        bar.pack(fill='x')
        self._btn(bar, '📂  นำเข้าไฟล์ต้นฉบับ (.xlsx)', self.b_import,
                  primary=True).pack(side='left', ipadx=10, ipady=4)
        self._btn(bar, '✔  เลือกทั้งหมด', lambda: self._b_all(True)).pack(
            side='left', padx=(8, 0), ipadx=8, ipady=4)
        self._btn(bar, '✗  ไม่เลือกเลย', lambda: self._b_all(False)).pack(
            side='left', padx=(8, 0), ipadx=8, ipady=4)
        self._btn(bar, '🗑  ล้าง', self.b_clear).pack(side='left', padx=(8, 0), ipadx=8, ipady=4)
        self.b_count = tk.Label(bar, text='ยังไม่ได้นำเข้า', bg=C['bg'], fg=C['dim'], font=FM)
        self.b_count.pack(side='right')

        bar2 = tk.Frame(s1, bg=C['bg'])
        bar2.pack(fill='x', pady=(8, 0))
        self._btn(bar2, '✏  แก้ชื่อบันเดิลทั้งหมด', self.b_rename_all,
                  primary=True).pack(side='left', ipadx=10, ipady=3)
        self._btn(bar2, '⌄  กางทั้งหมด', lambda: self._b_expand_all(True)).pack(
            side='left', padx=(8, 0), ipadx=8, ipady=3)
        self._btn(bar2, '⌃  หุบทั้งหมด', lambda: self._b_expand_all(False)).pack(
            side='left', padx=(8, 0), ipadx=8, ipady=3)
        tk.Label(s1, text='คลิกช่อง “ใช้” เพื่อเลือก/ไม่เลือกบันเดิล · ดับเบิลคลิกเพื่อแก้ '
                          '(ชื่อบันเดิล / จำนวน / Tier / famepoint / exp)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), anchor='w').pack(
            fill='x', pady=(8, 0))

        s2 = self._card(p, 'บันเดิลที่จะสร้าง', grow=True)
        tw = tk.Frame(s2, bg=C['bg'])
        tw.pack(fill='both', expand=True)
        cols = ('use', 'qty', 'tier', 'info')
        self.b_tree = ttk.Treeview(tw, columns=cols, show='tree headings',
                                   style='TR.Treeview', height=4)
        self.b_tree.heading('#0', text='บันเดิล / ไอเทมข้างใน')
        self.b_tree.column('#0', width=420, anchor='w')
        for c, t, w in (('use', 'ใช้', 44), ('qty', 'จำนวน', 70),
                        ('tier', 'Tier', 80), ('info', 'รายละเอียด', 300)):
            self.b_tree.heading(c, text=t)
            self.b_tree.column(c, width=w, anchor='w')
        self.b_tree.tag_configure('off', foreground=C['dim'])
        self.b_tree.tag_configure('rw', foreground=C['ok'])
        bsb = ttk.Scrollbar(tw, orient='vertical', command=self.b_tree.yview)
        self.b_tree.configure(yscrollcommand=bsb.set)
        self.b_tree.pack(side='left', fill='both', expand=True)
        bsb.pack(side='right', fill='y')
        self.b_tree.bind('<Button-1>', self._b_click)
        self.b_tree.bind('<Double-1>', lambda e: self.b_edit())

        s3 = self._card(p, 'ลงมือสร้าง')
        self.bv_do = tk.BooleanVar(value=False)
        tk.Checkbutton(s3, variable=self.bv_do, command=self._b_do_changed,
                       text='กดปุ่ม “สร้าง Bundle” จริง', bg=C['bg'], fg=C['fg'],
                       selectcolor=C['input'], activebackground=C['bg'],
                       activeforeground=C['fg'], font=FB, bd=0,
                       highlightthickness=0).grid(row=0, column=0, sticky='w')
        self.b_mode = tk.Label(s3, text='', bg=C['bg'], fg=C['warn'], font=FM)
        self.b_mode.grid(row=0, column=1, sticky='w', padx=(10, 0))

        self.bv_hold = tk.StringVar(value=self.prefs.get('b_hold', '4'))
        tk.Label(s3, text='โหมดทดสอบ: กรอกเสร็จแล้วค้างหน้าไว้ให้ดู (วินาที)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9)).grid(
            row=1, column=0, sticky='w', pady=(8, 0))
        eh = self._entry(s3, width=6)
        eh.config(textvariable=self.bv_hold)
        eh.grid(row=1, column=1, sticky='w', padx=(10, 0), pady=(8, 0), ipady=3)

        run = tk.Frame(s3, bg=C['bg'])
        run.grid(row=2, column=0, columnspan=3, sticky='w', pady=(12, 0))
        self.b_btn_run = self._btn(run, '▶  เริ่มทำงาน', self.b_start, primary=True)
        self.b_btn_run.pack(side='left', ipadx=18, ipady=5)
        self.b_btn_stop = self._btn(run, '■  ยกเลิก', self.b_stop)
        self.b_btn_stop.config(state='disabled')
        self.b_btn_stop.pack(side='left', padx=(8, 0), ipadx=12, ipady=5)

        # ---- เลข Bundle ที่สร้างสำเร็จ อยู่ในแท็บเดียวกับที่สั่งสร้าง ----
        self.made = []
        mb = tk.Frame(p, bg=C['bg'])
        mb.pack(fill='x', padx=14, pady=(10, 4))
        mw = tk.Frame(p, bg=C['bg'])
        mw.pack(fill='both', expand=True, padx=14, pady=(0, 12))
        tk.Label(mb, text='Bundle ที่สร้างแล้ว', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(side='left')
        self.lbl_made = tk.Label(mb, text='ยังไม่ได้สร้าง', bg=C['bg'], fg=C['dim'],
                                 font=('Segoe UI', 9))
        self.lbl_made.pack(side='left', padx=10)
        self._btn(mb, '📋  คัดลอกเลข Bundle', self.copy_made).pack(
            side='right', ipadx=8, ipady=2)
        tk.Label(mb, text='(เลือกแถวที่ต้องการก่อน · ไม่เลือก = เอาทั้งหมด)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8)).pack(
                     side='right', padx=8)
        mc = ('no', 'bid', 'bname', 'cnt', 'at')
        self.tree_made = ttk.Treeview(mw, columns=mc, show='headings', height=4,
                                      style='TR.Treeview', selectmode='extended')
        for c, t, w in (('no', '#', 42), ('bid', 'Bundle ID', 100),
                        ('bname', 'ชื่อ Bundle', 430), ('cnt', 'ไอเทม', 60),
                        ('at', 'เวลา', 80)):
            self.tree_made.heading(c, text=t)
            self.tree_made.column(c, width=w, anchor='w')
        msb = ttk.Scrollbar(mw, orient='vertical', command=self.tree_made.yview)
        self.tree_made.configure(yscrollcommand=msb.set)
        self.tree_made.pack(side='left', fill='both', expand=True)
        msb.pack(side='right', fill='y')

        self._b_do_changed()
        self._b_refresh()

    def _b_do_changed(self):
        if self.bv_do.get():
            self.b_mode.config(text='⚠  จะสร้างบันเดิลจริงบนเว็บ', fg=C['err'])
        else:
            self.b_mode.config(text='โหมดทดสอบ — กรอกให้ดูเฉยๆ ไม่กดสร้าง', fg=C['ok'])

    # ---------- ตาราง ----------
    def _b_refresh(self):
        # จำไว้ก่อนว่าใบไหนกางอยู่ · เลื่อนอยู่ตรงไหน · เลือกอะไรไว้
        # ไม่งั้นแก้ชื่อทีเดียวแล้วทุกใบกางออกหมด ต้องมาไล่หุบใหม่ทุกครั้ง
        was_open = {}
        for row, m in getattr(self, '_b_rowmap', {}).items():
            if m[0] == 'b':
                try:
                    was_open[m[1]] = bool(self.b_tree.item(row, 'open'))
                except Exception:
                    pass
        sel_key = None
        try:
            cur = self.b_tree.selection()
            if cur:
                sel_key = getattr(self, '_b_rowmap', {}).get(cur[0])
        except Exception:
            pass
        try:
            top_frac = self.b_tree.yview()[0]
        except Exception:
            top_frac = 0.0

        self.b_tree.delete(*self.b_tree.get_children())
        self._b_rowmap = {}
        self._b_open = was_open
        for i, b in enumerate(self.bq):
            rw = []
            for r in b.get('rewards', []):
                rw.append(('famepoint' if r['type'] == 'CREDIT' else 'exp') + ' ' + str(r['qty']))
            info = '%d ไอเทม' % len(b['items'])
            if rw:
                info += '  ·  ' + ' · '.join(rw)
            if b.get('sheet'):
                info += '  ·  ชีท ' + b['sheet']
            on = b.get('use', True)
            pid = self.b_tree.insert('', 'end', text='📦  ' + b['name'],
                                     values=('✔' if on else '✗', '', '', info),
                                     open=was_open.get(i, self.b_open_new),
                                     tags=() if on else ('off',))
            self._b_rowmap[pid] = ('b', i)
            for j, it in enumerate(b['items']):
                cid = self.b_tree.insert(
                    pid, 'end', text='      %d. %s' % (j + 1, it.get('disp') or ''),
                    values=('', it['qty'], it['tier'], 'Aztek Item Id = ' + it['id']),
                    tags=() if on else ('off',))
                self._b_rowmap[cid] = ('i', i, j)
            for r in b.get('rewards', []):
                nm = FAME_NAME if r['type'] == 'CREDIT' else 'Player Experience'
                extra = ('ประเภท: ' + FAME_CREDIT_LABEL) if r['type'] == 'CREDIT' else 'จากช่อง exp'
                rid = self.b_tree.insert(pid, 'end', text='      ⭐ ' + nm,
                                         values=('', r['qty'], DEFAULT_TIER, extra),
                                         tags=('rw',) if on else ('off',))
                self._b_rowmap[rid] = ('r', i, r['type'])
        # กลับไปอยู่ที่เดิม ทั้งแถวที่เลือกและจุดที่เลื่อนค้างไว้
        if sel_key:
            for row, m in self._b_rowmap.items():
                if m == sel_key:
                    try:
                        self.b_tree.selection_set(row)
                        self.b_tree.see(row)
                    except Exception:
                        pass
                    break
        try:
            self.b_tree.update_idletasks()
            self.b_tree.yview_moveto(top_frac)
        except Exception:
            pass

        n = len(self.bq)
        use = sum(1 for b in self.bq if b.get('use', True))
        self.b_count.config(text=('ยังไม่ได้นำเข้า' if not n
                                  else 'ทั้งหมด %d บันเดิล · จะสร้าง %d' % (n, use)))

    def _b_click(self, ev):
        """คลิกที่ช่อง "ใช้" ของแถวบันเดิล = สลับเลือก/ไม่เลือก"""
        if self.b_running:
            return
        if self.b_tree.identify_region(ev.x, ev.y) != 'cell':
            return
        if self.b_tree.identify_column(ev.x) != '#1':
            return
        row = self.b_tree.identify_row(ev.y)
        m = self._b_rowmap.get(row)
        if not m or m[0] != 'b':
            return
        b = self.bq[m[1]]
        b['use'] = not b.get('use', True)
        self._b_refresh()

    def _b_all(self, on):
        for b in self.bq:
            b['use'] = on
        self._b_refresh()

    def b_clear(self):
        if self.b_running or not self.bq:
            return
        if messagebox.askyesno('ล้าง', 'ลบบันเดิลทั้งหมด %d อันออกจากรายการ?' % len(self.bq)):
            self.bq = []
            self._b_refresh()

    # ---------- นำเข้า ----------
    def b_import(self):
        if self.b_running:
            return
        path = filedialog.askopenfilename(title='เลือกไฟล์ต้นฉบับ',
                                          filetypes=[('Excel', '*.xlsx *.xlsm'), ('ทุกไฟล์', '*.*')])
        if not path:
            return
        dlg = BundleImportDialog(self.root, path)
        if not dlg.result:
            return
        for b in dlg.result:
            b['use'] = True
            self.bq.append(b)
        self._b_refresh()
        self.log('นำเข้าบันเดิล %d อัน (รวม %d)' % (len(dlg.result), len(self.bq)), 'OK')
        for w in (dlg.warns or [])[:40]:
            self.log('  ⚠ ' + w, 'WARN')
        # เรื่องที่ "ของหาย/เลขอาจผิด" ต้องเด้งบอกเลย ไม่ใช่ซ่อนไว้ในแท็บ Log
        # (เคยมีเคสไอเทมหายไป 2 ใบต่อบันเดิลแล้วไม่มีใครเห็นคำเตือน)
        heavy = [w for w in (dlg.warns or [])
                 if 'ข้าม' in w or 'เลขเรียง' in w or 'อ่านมาได้' in w]
        if heavy:
            messagebox.showwarning(
                'อ่านไฟล์ได้ไม่ครบ',
                'นำเข้ามาแล้ว แต่มี %d เรื่องที่ควรดูก่อนสร้าง:\n\n%s%s\n\n'
                'ดูทั้งหมดได้ในแท็บ Log'
                % (len(heavy), '\n'.join('· ' + w for w in heavy[:6]),
                   '\n· ...' if len(heavy) > 6 else ''))
        log_event('bundle_import', count=len(dlg.result), total=len(self.bq),
                  file=os.path.basename(path))

    # ---------- แก้ไข ----------
    def b_edit(self):
        if self.b_running:
            return
        sel = self.b_tree.selection()
        if not sel:
            return
        m = self._b_rowmap.get(sel[0])
        if not m:
            return
        if m[0] == 'b':
            self._b_form_bundle(self.bq[m[1]])
        elif m[0] == 'i':
            self._b_form_item(self.bq[m[1]]['items'][m[2]])
        else:
            b = self.bq[m[1]]
            for r in b['rewards']:
                if r['type'] == m[2]:
                    self._b_form_reward(r)
                    break
        self._b_refresh()

    def _b_dialog(self, title, h=300):
        top = tk.Toplevel(self.root)
        top.title(title)
        top.configure(bg=C['bg'])
        top.transient(self.root)
        top.grab_set()
        try:
            self.root.update_idletasks()
            x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - 560) // 2)
            y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - h) // 2)
            top.geometry('560x%d+%d+%d' % (h, x, y))
        except Exception:
            top.geometry('560x%d' % h)
        return top

    def _b_expand_all(self, on):
        """กาง/หุบทุกใบทีเดียว แล้วจำไว้ว่าจะให้ใบใหม่เป็นแบบไหน"""
        self.b_open_new = bool(on)
        for row, m in self._b_rowmap.items():
            if m[0] == 'b':
                try:
                    self.b_tree.item(row, open=bool(on))
                except Exception:
                    pass

    def b_rename_all(self):
        """แก้ชื่อบันเดิลทุกอันในหน้าเดียว — ไม่ต้องดับเบิลคลิกทีละใบ

        ของเดิมต้องดับเบิลคลิก -> แก้ -> ปิด -> เลื่อนหาใบถัดไป วนไปเรื่อยๆ
        บันเดิลเยอะๆ (40-50 อัน) เสียเวลามาก
        หน้านี้ไล่ Tab ลงไปพิมพ์ทีเดียวจบ แล้วกดบันทึกครั้งเดียว
        """
        if not self.bq:
            messagebox.showinfo('ยังไม่มีบันเดิล', 'นำเข้าไฟล์ต้นฉบับก่อนนะ')
            return
        top = self._b_dialog('แก้ชื่อบันเดิลทั้งหมด (%d อัน)' % len(self.bq), 560)
        try:      # หน้านี้กว้างกว่าปกติ จัดให้อยู่กลางจออีกที
            self.root.update_idletasks()
            x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - 820) // 2)
            y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - 560) // 2)
            top.geometry('820x560+%d+%d' % (x, y))
        except Exception:
            top.geometry('820x560')

        head = tk.Frame(top, bg=C['bg'])
        head.pack(fill='x', padx=18, pady=(14, 6))
        tk.Label(head, text='พิมพ์ชื่อใหม่ได้เลย · กด Tab เพื่อไปช่องถัดไป · '
                            'เว้นว่าง = ใช้ชื่อเดิม',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9), anchor='w').pack(side='left')

        # แถบเลื่อนเอง เพราะรายการยาวกว่าหน้าจอ
        wrap = tk.Frame(top, bg=C['bg'])
        wrap.pack(fill='both', expand=True, padx=18)
        cv = tk.Canvas(wrap, bg=C['bg'], highlightthickness=0)
        sb = ttk.Scrollbar(wrap, orient='vertical', command=cv.yview)
        inner = tk.Frame(cv, bg=C['bg'])
        cv.configure(yscrollcommand=sb.set)
        cv.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')
        win = cv.create_window((0, 0), window=inner, anchor='nw')

        def _fit(_e=None):
            cv.configure(scrollregion=cv.bbox('all'))
            try:
                cv.itemconfigure(win, width=cv.winfo_width())
            except Exception:
                pass
        inner.bind('<Configure>', _fit)
        cv.bind('<Configure>', _fit)

        def _wheel(ev):
            cv.yview_scroll(-1 * (ev.delta // 120 or (1 if ev.delta < 0 else -1)), 'units')
        for w in (cv, inner, top):
            w.bind('<MouseWheel>', _wheel)

        vs = []
        for i, b in enumerate(self.bq):
            row = tk.Frame(inner, bg=C['bg'])
            row.pack(fill='x', pady=2)
            tk.Label(row, text='%2d.' % (i + 1), bg=C['bg'], fg=C['dim'],
                     font=('Consolas', 9), width=4, anchor='e').pack(side='left')
            v = tk.StringVar(value=b['name'])
            e = self._entry(row, width=52)
            e.config(textvariable=v)
            e.pack(side='left', ipady=3, padx=(6, 8))
            e.bind('<MouseWheel>', _wheel)
            info = '%d ไอเทม' % len(b['items'])
            if b.get('sheet'):
                info += '  ·  ' + b['sheet']
            tk.Label(row, text=info, bg=C['bg'], fg=C['dim'],
                     font=('Segoe UI', 8), anchor='w').pack(side='left')
            vs.append(v)
            if i == 0:
                e.focus_set()

        def _save():
            n = 0
            for b, v in zip(self.bq, vs):
                nm = v.get().strip()
                if nm and nm != b['name']:
                    b['name'] = nm
                    n += 1
            top.destroy()
            self._b_refresh()
            self.log('แก้ชื่อบันเดิล %d อัน' % n, 'OK' if n else 'INFO')

        foot = tk.Frame(top, bg=C['card'], height=58)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)
        self._btn(foot, 'บันทึกทั้งหมด', _save, primary=True).pack(
            side='right', padx=18, pady=12, ipadx=20, ipady=4)
        self._btn(foot, 'ยกเลิก', top.destroy).pack(side='right', pady=12, ipadx=14, ipady=4)
        self.root.wait_window(top)

    def _b_form_bundle(self, b):
        top = self._b_dialog('แก้ไขบันเดิล', 300)
        body = tk.Frame(top, bg=C['bg'])
        body.pack(fill='both', expand=True, padx=18, pady=14)
        tk.Label(body, text='ชื่อบันเดิล', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9)).grid(row=0, column=0, sticky='w', pady=4)
        v_name = tk.StringVar(value=b['name'])
        e = self._entry(body, width=46)
        e.config(textvariable=v_name)
        e.grid(row=0, column=1, sticky='w', padx=(12, 0), ipady=3)

        tk.Label(body, text='ประเภท Bundle', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9)).grid(row=1, column=0, sticky='w', pady=4)
        v_type = ttk.Combobox(body, values=BUNDLE_TYPES, width=20, state='readonly', font=FM)
        v_type.set(b.get('type', DEFAULT_BUNDLE_TYPE))
        v_type.grid(row=1, column=1, sticky='w', padx=(12, 0))

        v_del = tk.BooleanVar(value=bool(b.get('deliver', True)))
        tk.Checkbutton(body, text=' ส่งทันที (Immediately Send)', variable=v_del,
                       bg=C['bg'], fg=C['fg'], selectcolor=C['input'],
                       activebackground=C['bg'], activeforeground=C['fg'], font=FM,
                       bd=0, highlightthickness=0).grid(row=2, column=1, sticky='w',
                                                        padx=(12, 0), pady=6)
        ok = {'v': False}

        def _save():
            b['name'] = v_name.get().strip() or b['name']
            b['type'] = v_type.get()
            b['deliver'] = v_del.get()
            ok['v'] = True
            top.destroy()
        foot = tk.Frame(top, bg=C['card'], height=58)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)
        self._btn(foot, 'บันทึก', _save, primary=True).pack(side='right', padx=18, pady=12,
                                                            ipadx=20, ipady=4)
        self._btn(foot, 'ยกเลิก', top.destroy).pack(side='right', pady=12, ipadx=14, ipady=4)
        self.root.wait_window(top)

    def _b_form_item(self, it):
        top = self._b_dialog('แก้ไขไอเทมในบันเดิล', 290)
        body = tk.Frame(top, bg=C['bg'])
        body.pack(fill='both', expand=True, padx=18, pady=14)
        vs = {}
        for r, (lbl, key, w) in enumerate((('Aztek Item Id (เลขที่เอาไปค้น)', 'id', 18),
                                           ('จำนวน (Quantity)', 'qty', 12))):
            tk.Label(body, text=lbl, bg=C['bg'], fg=C['dim'],
                     font=('Segoe UI', 9)).grid(row=r, column=0, sticky='w', pady=5)
            v = tk.StringVar(value=str(it.get(key, '') or ''))
            vs[key] = v
            e = self._entry(body, width=w)
            e.config(textvariable=v)
            e.grid(row=r, column=1, sticky='w', padx=(12, 0), ipady=3)
        tk.Label(body, text='Tier', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9)).grid(row=2, column=0, sticky='w', pady=5)
        cb = ttk.Combobox(body, values=TIERS, width=12, state='readonly', font=FM)
        cb.set(it.get('tier', DEFAULT_TIER))
        cb.grid(row=2, column=1, sticky='w', padx=(12, 0))
        tk.Label(body, text='ชื่อที่ชีทเขียนไว้: ' + (it.get('disp') or '—'),
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), wraplength=500,
                 anchor='w', justify='left').grid(row=3, column=0, columnspan=2,
                                                  sticky='w', pady=(10, 0))

        def _save():
            nid = vs['id'].get().strip()
            if not nid.isdigit():
                return messagebox.showwarning('ไม่ถูกต้อง', 'Aztek Item Id ต้องเป็นตัวเลข')
            it['id'] = nid
            it['qty'] = vs['qty'].get().strip() or '1'
            it['tier'] = cb.get()
            top.destroy()
        foot = tk.Frame(top, bg=C['card'], height=58)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)
        self._btn(foot, 'บันทึก', _save, primary=True).pack(side='right', padx=18, pady=12,
                                                            ipadx=20, ipady=4)
        self._btn(foot, 'ยกเลิก', top.destroy).pack(side='right', pady=12, ipadx=14, ipady=4)
        self.root.wait_window(top)

    def _b_form_reward(self, r):
        nm = FAME_NAME if r['type'] == 'CREDIT' else 'Player Experience'
        top = self._b_dialog('แก้ไข ' + nm, 250)
        body = tk.Frame(top, bg=C['bg'])
        body.pack(fill='both', expand=True, padx=18, pady=14)
        tk.Label(body, text='จำนวน', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9)).grid(row=0, column=0, sticky='w', pady=5)
        v = tk.StringVar(value=str(r.get('qty', '')))
        e = self._entry(body, width=14)
        e.config(textvariable=v)
        e.grid(row=0, column=1, sticky='w', padx=(12, 0), ipady=3)
        note = ('บนเว็บจะตั้งประเภทเป็น “%s” ให้อัตโนมัติ' % FAME_CREDIT_LABEL
                if r['type'] == 'CREDIT' else 'เพิ่มจากแท็บ Player Exp.')
        tk.Label(body, text=note, bg=C['bg'], fg=C['warn'], font=('Segoe UI', 8),
                 wraplength=500, anchor='w', justify='left').grid(
            row=1, column=0, columnspan=2, sticky='w', pady=(10, 0))

        def _save():
            r['qty'] = v.get().strip() or r['qty']
            top.destroy()
        foot = tk.Frame(top, bg=C['card'], height=58)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)
        self._btn(foot, 'บันทึก', _save, primary=True).pack(side='right', padx=18, pady=12,
                                                            ipadx=20, ipady=4)
        self._btn(foot, 'ยกเลิก', top.destroy).pack(side='right', pady=12, ipadx=14, ipady=4)
        self.root.wait_window(top)

    # ---------- รัน ----------
    def b_stop(self):
        self.b_cancel = True
        self.log('กำลังยกเลิกการสร้างบันเดิล...', 'WARN')

    def b_start(self):
        if self.b_running or self.running or getattr(self, 'c_running', False):
            return messagebox.showinfo('กำลังทำงาน', 'รอให้งานปัจจุบันเสร็จก่อนนะ')
        use = [b for b in self.bq if b.get('use', True)]
        if not use:
            return messagebox.showwarning('ยังไม่ได้เลือก', 'ติ๊กเลือกบันเดิลที่จะสร้างก่อนนะ')
        do = self.bv_do.get()
        if do:
            msg = 'จะสร้างบันเดิลจริงบนเว็บ %d อัน\n\n' % len(use)
            msg += '\n'.join('  • %s (%d ไอเทม)' % (b['name'][:40], len(b['items']))
                             for b in use[:8])
            if len(use) > 8:
                msg += '\n  … อีก %d อัน' % (len(use) - 8)
            msg += '\n\nตรวจชื่อ/จำนวน/Tier เรียบร้อยแล้วใช่ไหม?'
            if not messagebox.askyesno('ยืนยัน', msg):
                return
        self.prefs['b_hold'] = self.bv_hold.get().strip()
        save_prefs(self.prefs)
        self.b_running = True
        self.b_cancel = False
        self.b_btn_run.config(state='disabled')
        self.b_btn_stop.config(state='normal')
        self.nb.select(self.tab_log)
        self.log('=' * 46, 'STEP')
        self.log(('เริ่มสร้างบันเดิลจริง ' if do else 'เริ่มทดสอบกรอกบันเดิล ')
                 + '%d อัน' % len(use), 'STEP')
        log_event('bundle_start', count=len(use), commit=bool(do))
        threading.Thread(target=self._b_thread, args=([dict(b) for b in use], do),
                         daemon=True).start()

    def _b_thread(self, rows, do):
        try:
            asyncio.run(self._b_work(rows, do))
        except Exception as ex:
            log_event('error', where='bundle', message=str(ex)[:300])
            self.log('ผิดพลาด: ' + str(ex), 'ERR')
            self.log(traceback.format_exc(), 'ERR')
        finally:
            self.b_running = False

            def _rst():
                self.b_btn_run.config(state='normal')
                self.b_btn_stop.config(state='disabled')
            self.root.after(0, _rst)

    async def _b_work(self, bundles, do):
        async with async_playwright() as pw:
            browser = await pw.chromium.launch_persistent_context(**launch_kwargs(False))
            page = browser.pages[0] if browser.pages else await browser.new_page()
            try:
                try:
                    hold = max(0, int(float(self.bv_hold.get() or 0)))
                except Exception:
                    hold = 4
                okc = errc = 0
                for i, b in enumerate(bundles, 1):
                    if self.b_cancel:
                        self.log('ยกเลิกแล้ว', 'WARN')
                        break
                    self.set_progress(i - 1, len(bundles), b['name'][:30])
                    self.log('[%d/%d] %s  (%d ไอเทม)'
                             % (i, len(bundles), b['name'], len(b['items'])), 'STEP')
                    try:
                        if await self._b_one(page, b, do, hold):
                            okc += 1
                        else:
                            errc += 1
                    except Exception as ex:
                        errc += 1
                        self.log('   ✗ ' + str(ex)[:160], 'ERR')
                        log_event('error', where='bundle_one', name=b['name'],
                                  message=str(ex)[:200])
                self.set_progress(len(bundles), len(bundles), 'เสร็จ')
                self.log('จบ — สำเร็จ %d · ไม่ผ่าน %d' % (okc, errc),
                         'OK' if not errc else 'WARN')
                log_event('bundle_done', ok=okc, fail=errc, commit=bool(do))
            finally:
                try:
                    await browser.close()
                except Exception:
                    pass

    async def _b_one(self, page, b, do, hold):
        await page.goto(BUNDLE_CREATE_URL, wait_until='domcontentloaded', timeout=45000)
        await page.wait_for_timeout(2500)
        if any(k in page.url.lower() for k in ('login', 'signin', 'auth')):
            self.log('   ✗ ยังไม่ได้ล็อกอิน — กด “เปิดหน้า Login” ด้านบนก่อน', 'ERR')
            return False

        # [1] ข้อมูลบันเดิล
        r = await page.evaluate(JS_BUNDLE_NAME, b['name'])
        self.log('   · ชื่อ Bundle = %s%s' % (b['name'], '' if r == 'ok' else '  (%s)' % r),
                 'INFO' if r == 'ok' else 'WARN')
        if b.get('type') and b['type'] != DEFAULT_BUNDLE_TYPE:
            await self._b_pick(page, SEL_BUNDLE['type'], b['type'])

        # [2] ไอเทม — เพิ่มให้ครบก่อน แล้วค่อยกางการ์ดทีเดียวเพื่อกรอกจำนวน/Tier
        added = 0
        done = []
        missed = []
        for it in b['items']:
            if self.b_cancel:
                break
            if await self._b_add_item(page, it):
                added += 1
                done.append(it)
            else:
                missed.append(it['id'])
        if done:
            await self._b_expand(page)
            nth = {}      # ไอเทมซ้ำ = คนละใบ ต้องตั้งค่าให้ครบทุกใบ
            for it in done:
                nth[it['id']] = nth.get(it['id'], 0) + 1
                await self._b_set_row(page, it['id'], it['qty'], it['tier'], nth[it['id']])
        if missed:
            self.log('   ⚠ ไอเทมที่เพิ่มไม่ได้: %s' % ', '.join(missed), 'ERR')

        # [3] famepoint / exp
        wallet_bad = []
        for rw in b.get('rewards', []):
            if self.b_cancel:
                break
            if rw['type'] == 'CREDIT':
                if await self._b_add_wallet(page, SEL_BUNDLE['tab_credit'],
                                            FAME_OPTION, rw['qty']):
                    added += 1
                    if not await self._b_finish_wallet(page, FAME_NAME, True):
                        wallet_bad.append(FAME_NAME)
            else:
                if await self._b_add_wallet(page, SEL_BUNDLE['tab_exp'], None, rw['qty']):
                    added += 1
                    if not await self._b_finish_wallet(page, EXP_NAME, False):
                        wallet_bad.append(EXP_NAME)

        self.log('   เพิ่มเข้าบันเดิลแล้ว %d รายการ' % added,
                 'INFO' if added else 'WARN')

        # ตรวจ "ประเภท Bundle" ว่ายังเป็นค่าที่ตั้งใจ (เคยโดนดรอปดาวน์อื่นเปลี่ยนเป็น CHOICE)
        want_type = b.get('type') or DEFAULT_BUNDLE_TYPE
        got_type = await page.evaluate(JS_BUNDLE_TYPE_GET)
        if got_type and want_type not in str(got_type):
            self.log('   ! ประเภท Bundle กลายเป็น "%s" (ควรเป็น %s) — แก้กลับให้'
                     % (got_type, want_type), 'WARN')
            await self._b_pick(page, SEL_BUNDLE['type'], want_type)
            got_type = await page.evaluate(JS_BUNDLE_TYPE_GET)
            if got_type and want_type not in str(got_type):
                self.log('   ✗ ประเภท Bundle ยังเป็น "%s" — แก้เองบนเว็บก่อนกดสร้าง'
                         % got_type, 'ERR')
                if do:
                    return False

        # ตรวจชื่อบันเดิลอีกรอบ — ถ้าโดนเขียนทับด้วยเลขไอเทม จะจับได้ตรงนี้
        got = await page.evaluate(JS_BUNDLE_NAME_GET)
        if (got or '').strip() != b['name'].strip():
            self.log('   ! ชื่อ Bundle ไม่ตรง (ในช่องเป็น "%s") — กรอกใหม่ให้' % (got or ''),
                     'WARN')
            await page.evaluate(JS_BUNDLE_NAME, b['name'])
            got = await page.evaluate(JS_BUNDLE_NAME_GET)
            if (got or '').strip() != b['name'].strip():
                self.log('   ✗ แก้ชื่อ Bundle ไม่สำเร็จ (ยังเป็น "%s")' % (got or ''), 'ERR')
                if do:
                    return False

        if wallet_bad:
            self.log('   ✗ ยังตั้งค่าไม่ครบ: %s — เว็บจะกดสร้างไม่ผ่าน'
                     % ', '.join(wallet_bad), 'ERR')
            if do:
                return False

        if not do:
            self.log('   ✓ กรอกครบแล้ว (โหมดทดสอบ — ไม่กดสร้าง)', 'OK')
            if hold:
                await page.wait_for_timeout(hold * 1000)
            return added > 0

        btn = page.locator('button:has-text("%s")' % SEL_BUNDLE['submit']).last
        if await btn.count() == 0:
            self.log('   ✗ ไม่เจอปุ่ม “%s”' % SEL_BUNDLE['submit'], 'ERR')
            return False

        # ดักคำตอบของเว็บไว้ตั้งแต่ก่อนกด — เลข Bundle อยู่ในคำตอบนั้น
        hits = []

        async def _grab(resp):
            try:
                if resp.request.method in ('POST', 'PUT', 'PATCH') \
                        and 'bundle' in resp.url.lower():
                    hits.append(resp)
            except Exception:
                pass
        page.on('response', _grab)
        try:
            await btn.click(timeout=10000)
            await page.wait_for_timeout(900)

            # เว็บเด้งป๊อปอัป "ยืนยันการสร้าง" ต่อ — ไม่กดยืนยัน = ไม่ได้สร้างจริง
            r = await page.evaluate(JS_BUNDLE_CONFIRM, [BUNDLE_YES, BUNDLE_NO])
            if str(r).startswith('ok'):
                self.log('   · กดยืนยันในป๊อปอัปแล้ว (%s)'
                         % str(r).split('|', 1)[1], 'INFO')
            elif r != 'ไม่มีป๊อปอัป':
                self.log('   ! %s' % r, 'WARN')

            waited = 0
            while waited < 25000 and not hits:
                await page.wait_for_timeout(300)
                waited += 300
            await page.wait_for_timeout(1200)
        finally:
            try:
                page.remove_listener('response', _grab)
            except Exception:
                pass

        code = 0
        bid = ''
        for resp in hits:
            try:
                code = resp.status
                body = await resp.json()
            except Exception:
                body = None
            bid = bid or _dig_id(body)
        if not bid:
            try:
                bid = str(await page.evaluate(JS_BUNDLE_MADE_ID) or '').strip()
            except Exception:
                bid = ''

        good = 200 <= code < 300
        if good:
            self.log('   ✓ สร้างบันเดิลแล้ว%s'
                     % ('  ·  Bundle ID = %s' % bid if bid else
                        '  (เว็บไม่ได้ส่งเลข Bundle กลับมา)'), 'OK')
            self.add_made_bundle(b['name'], bid, len(b['items']))
        elif code:
            self.log('   ✗ เว็บตอบ HTTP %d — ยังไม่ได้สร้าง' % code, 'ERR')
        else:
            self.log('   ✗ กดสร้างแล้วเว็บไม่ตอบอะไรกลับมาเลย — ถือว่ายังไม่ได้สร้าง '
                     '(เช็กบนเว็บอีกทีก่อนสร้างซ้ำ)', 'ERR')
        log_event('bundle_create', name=b['name'], items=len(b['items']),
                  ok=bool(good), http=code, bundle_id=bid)
        return good

    # ---------- ชิ้นส่วนของหน้า bundle ----------
    async def _b_tab(self, page, name):
        try:
            t = page.locator('button:has-text("%s"), [role="tab"]:has-text("%s")'
                             % (name, name)).first
            if await t.count() > 0:
                await t.click(timeout=5000)
                await page.wait_for_timeout(500)
                return True
        except Exception:
            pass
        return False

    async def _b_seek(self, page, term, want_id):
        """ค้นด้วยคำหนึ่งคำ แล้วไล่ดูทีละหน้าจนเจอแถวที่ ID ตรงเป๊ะ

        สำคัญ: ต้อง "ดูหน้าที่ยืนอยู่ก่อน แล้วค่อยกดหน้าถัดไป"
        ของเดิมนับหน้าผิด เลยกดไปหน้าสุดท้ายแล้วจบ โดยยังไม่ได้ดูหน้านั้นเลย
        (ไอเทมเลข 2 หลักอย่าง 51 อยู่หน้าสุดท้ายพอดี เลยหาไม่เจอ)
        """
        r = await page.evaluate(JS_BUNDLE_SEARCH, str(term))
        if r != 'ok':
            return None, {'err': r}
        checked, total = 0, 0
        while checked < BUNDLE_MAX_PAGES and not self.b_cancel:
            st = await self._b_wait_rows(page, want_id)
            checked += 1
            if not total:
                try:
                    total = int((await page.evaluate(JS_BUNDLE_PAGES)).get('total') or 0)
                except Exception:
                    total = 0
            if st.get('hit'):
                return True, {'page': checked, 'total': total}
            if st.get('rows', 0) == 0:
                return False, {'page': checked, 'total': total, 'empty': True}
            if total and checked >= total:
                return False, {'page': checked, 'total': total}
            if await page.evaluate(JS_BUNDLE_NEXT_PAGE) != 'ok':
                return False, {'page': checked, 'total': total}
            await page.wait_for_timeout(400)
        return False, {'page': checked, 'total': total, 'over': True}

    async def _b_add_item(self, page, it):
        """แท็บ Item -> ค้นหา -> กดปุ่มเพิ่มของแถวที่ ID ตรงเป๊ะเท่านั้น

        ค้นด้วย ItemKind ก่อนถ้ามี เพราะเจาะจงกว่า ผลน้อยกว่ามาก
        ไม่เจอค่อยค้นด้วย Aztek Item Id แล้วไล่หน้าไปเรื่อยๆ
        (เลขสั้นๆ อย่าง 51 / 134 ผลจะเยอะและกระจายหลายหน้า)
        """
        await self._b_tab(page, SEL_BUNDLE['tab_item'])
        terms = []
        if str(it.get('kind') or '').strip():
            terms.append((str(it['kind']).strip(), 'ItemKind'))
        terms.append((str(it['id']), 'Item ID'))

        last = {}
        for term, what in terms:
            found, info = await self._b_seek(page, term, it['id'])
            if found is None:
                self.log('   ✗ ไม่เจอช่องค้นหา Item (%s) — ข้ามไอเทม %s'
                         % (info.get('err'), it['id']), 'ERR')
                return False
            if found:
                if what != 'Item ID' or info.get('page', 1) > 1:
                    self.log('   · หาไอเทม %s เจอจาก %s (หน้า %d/%d)'
                             % (it['id'], what, info.get('page', 1), info.get('total') or 1),
                             'INFO')
                break
            last = info
        else:
            n, tot = last.get('page', 0), last.get('total') or 0
            why = ('ค้นแล้วไม่มีผลเลย' if last.get('empty')
                   else ('ผลเยอะเกิน ดูไป %d หน้า จาก %d หน้า' % (n, tot) if last.get('over')
                         else 'ดูครบ %d หน้าแล้วไม่มีแถวที่ ID ตรง' % n))
            self.log('   ✗ ค้นหา %s ไม่สำเร็จ — %s (ไม่กดเพิ่ม กันได้ของผิด)'
                     % (it['id'], why), 'ERR')
            return False

        r2 = await page.evaluate(JS_BUNDLE_PICK_ROW, [str(it['id'])])
        if not str(r2).startswith('ok'):
            self.log('   ✗ กดเพิ่ม %s ไม่สำเร็จ (%s)' % (it['id'], r2), 'ERR')
            return False
        await page.wait_for_timeout(700)
        seen = str(r2).split('|', 1)[1] if '|' in str(r2) else ''
        self.log('   · เพิ่มไอเทม %s (qty=%s tier=%s)  [%s]'
                 % (it['id'], it['qty'], it['tier'], seen[:60]), 'INFO')
        return True

    async def _b_wait_rows(self, page, item_id, timeout=15000):
        """รอจนผลค้นหาโหลดเสร็จ — คืนสถานะล่าสุด"""
        waited = 0
        st = {}
        while waited < timeout:
            st = await page.evaluate(JS_BUNDLE_ROW_STATE, [str(item_id)])
            if st.get('hit'):
                return st
            if st.get('state') == 'ready' and st.get('rows', 0) > 0:
                return st
            await page.wait_for_timeout(300)
            waited += 300
        return st

    async def _b_add_wallet(self, page, tab, option, qty):
        """แท็บ Credit / Player Exp. -> เลือกจากดรอปดาวน์ -> ใส่จำนวน -> กด + เพิ่ม

        ทุกอย่างจำกัดอยู่ใน "กล่องเพิ่มของเข้า Bundle" เท่านั้น
        ห้ามเล็งทั้งหน้า ไม่งั้นจะไปโดน "ประเภท Bundle" ที่แผงขวา แล้วเปลี่ยนค่าเองโดยไม่ตั้งใจ
        """
        if not await self._b_tab(page, tab):
            self.log('   ! ไม่เจอแท็บ %s' % tab, 'WARN')
            return False
        mk = await page.evaluate(JS_BUNDLE_MARK_ADDBOX, ['pick'])
        if not str(mk).startswith('ok'):
            self.log('   ✗ %s: %s' % (tab, mk), 'ERR')
            return False
        try:
            await page.locator('[data-trpu="pick"]').first.click(timeout=5000)
            await page.wait_for_timeout(600)
            if option:
                box = page.locator('input[placeholder*="ค้นหา"]').last
                if await box.count() > 0:
                    await box.fill(option, timeout=4000)
                    await page.wait_for_timeout(800)
            opts = page.locator('[role="option"]')
            if await opts.count() == 0:
                self.log('   ✗ %s: ดรอปดาวน์ไม่มีตัวเลือกให้เลือก' % tab, 'ERR')
                await page.keyboard.press('Escape')
                return False
            await opts.first.click(timeout=5000)
            await page.wait_for_timeout(500)
        except Exception as ex:
            self.log('   ✗ เลือกใน %s ไม่ได้: %s' % (tab, str(ex)[:60]), 'ERR')
            try:
                await page.keyboard.press('Escape')
            except Exception:
                pass
            return False

        mk = await page.evaluate(JS_BUNDLE_MARK_ADDBOX, ['wqty'])
        if str(mk).startswith('ok'):
            await page.evaluate(JS_BUNDLE_SET_MARKED, ['wqty', str(qty)])
        else:
            self.log('   ! %s: %s' % (tab, mk), 'WARN')

        mk = await page.evaluate(JS_BUNDLE_MARK_ADDBOX, ['wadd'])
        if not str(mk).startswith('ok'):
            self.log('   ✗ %s: ไม่เจอปุ่มเพิ่มในกล่อง' % tab, 'ERR')
            return False
        try:
            await page.locator('[data-trpu="wadd"]').first.click(timeout=6000)
            await page.wait_for_timeout(900)
        except Exception as ex:
            self.log('   ✗ กดเพิ่มใน %s ไม่ได้: %s' % (tab, str(ex)[:60]), 'ERR')
            return False
        self.log('   · เพิ่ม %s จำนวน %s' % (tab, qty), 'INFO')
        return True

    async def _b_expand(self, page):
        """กด "ขยายทั้งหมด" — การ์ดที่ย่ออยู่ไม่มีช่องจำนวน/Tier ให้กรอก"""
        try:
            r = await page.evaluate(JS_BUNDLE_EXPAND)
            if r != 'ok':
                btn = page.locator('button:has-text("ขยายทั้งหมด")').first
                if await btn.count() > 0:
                    await btn.click(timeout=4000)
                    r = 'ok'
            await page.wait_for_timeout(700)
            return r == 'ok'
        except Exception:
            return False

    async def _b_settle(self, page, timeout=4000):
        """รอให้ดรอปดาวน์ตัวก่อนหน้าปิดสนิทจริงๆ ก่อนเปิดตัวถัดไป

        เว็บล็อกหน้าจอไว้ตอนดรอปดาวน์เปิด (pointer-events: none)
        ถ้ารีบเปิดตัวถัดไป มันจะไม่เปิดให้ แล้วหาตัวเลือกไม่เจอสักอัน
        """
        waited = 0
        while waited < timeout:
            try:
                st = await page.evaluate(JS_BUNDLE_POPUP_OPEN)
            except Exception:
                return
            if not st.get('open') and not st.get('locked'):
                await page.wait_for_timeout(250)     # เผื่อแอนิเมชันปิด
                return
            try:
                await page.keyboard.press('Escape')
            except Exception:
                pass
            await page.wait_for_timeout(200)
            waited += 200

    async def _b_open_dd(self, page, how, key, what, nth=1):
        """เปิดดรอปดาวน์ให้ได้จริง — ลองกดด้วย JS ก่อน ไม่ติดค่อยกดด้วยเมาส์จริง"""
        for attempt in range(3):
            await self._b_settle(page)
            if attempt == 0:
                # กดด้วยเมาส์จริงก่อน — เหมือนคนกดที่สุด
                r = await page.evaluate(JS_BUNDLE_ACT,
                                        [how, str(key), what, 'point', '', TIERS, nth])
                if str(r).startswith('ok'):
                    try:
                        x, y = str(r).split('|', 1)[1].split(',')
                        await page.mouse.click(float(x), float(y))
                    except Exception:
                        pass
            elif attempt == 1:
                # ยิงเหตุการณ์กดเมาส์ครบชุด — ดรอปดาวน์พวกนี้เปิดตอน "กดลง"
                r = await page.evaluate(JS_BUNDLE_ACT,
                                        [how, str(key), what, 'press', '', TIERS, nth])
            else:
                r = await page.evaluate(JS_BUNDLE_ACT,
                                        [how, str(key), what, 'click', '', TIERS, nth])
            if not str(r).startswith('ok'):
                return r
            for _ in range(12):
                await page.wait_for_timeout(200)
                try:
                    if await page.locator('[role="option"]').count() > 0:
                        return 'ok'
                except Exception:
                    pass
        return 'เปิดดรอปดาวน์แล้วไม่มีตัวเลือกโผล่'

    async def _b_choose(self, page, how, key, what, want, exact=True, label=None, nth=1):
        """ตั้งค่าดรอปดาวน์/ช่องในการ์ด แล้วอ่านกลับมาตรวจว่าเข้าจริง

        ทุกขั้นตอนหาช่องใหม่ทุกครั้ง ไม่เก็บตัวชี้ไว้ข้ามคำสั่ง
        เพราะเว็บเป็น React — render ใหม่เมื่อไหร่ ตัวชี้เก่าใช้ไม่ได้ทันที
        """
        who = label or str(key)

        def _match(v):
            v = (v or '').strip()
            return (v == want) if exact else (want in v)

        await self._b_settle(page)
        r = await page.evaluate(JS_BUNDLE_ACT, [how, str(key), what, 'read', '', TIERS, nth])
        if not str(r).startswith('ok'):
            self.log('   ✗ %s (%s): %s' % (who, what, r), 'ERR')
            return False
        if _match(str(r).split('|', 1)[1] if '|' in str(r) else ''):
            return True

        # <select> ธรรมดา ตั้งได้เลย
        r = await page.evaluate(JS_BUNDLE_ACT, [how, str(key), what, 'set', want, TIERS, nth])
        if r == 'ok':
            self.log('   · %s %s = %s' % (who, what, want), 'INFO')
            return True
        why = str(r)

        # ช่องแบบ <select> ห้ามกดเปิด — ตัวเลือกของมันเป็นหน้าต่างของเบราว์เซอร์
        # ไม่ได้อยู่ในหน้าเว็บ กดแล้วจะค้างและหาตัวเลือกไม่เจอสักอัน
        tag = await page.evaluate(JS_BUNDLE_ACT, [how, str(key), what, 'tag', '', TIERS, nth])
        if str(tag).endswith('SELECT') and why != 'combobox':
            self.log('   ✗ %s (%s): %s' % (who, what, why), 'ERR')
            return False

        opened = await self._b_open_dd(page, how, key, what, nth)
        if opened != 'ok':
            self.log('   ✗ %s (%s) เปิดดรอปดาวน์ไม่ได้: %s (ตอนตั้งค่าตรงๆ: %s)'
                     % (who, what, opened, why), 'ERR')
            return False

        opt = None
        seen = []
        opts = page.locator('[role="option"]')
        try:
            n = await opts.count()
        except Exception:
            n = 0
        for i in range(n):
            try:
                txt = (await opts.nth(i).inner_text()).strip()
            except Exception:
                continue
            seen.append(txt)
            if _match(txt):
                opt = opts.nth(i)
                break
        if opt is None:
            self.log('   ✗ %s (%s): ไม่มีตัวเลือก "%s" (เจอ: %s)'
                     % (who, what, want, ', '.join(seen[:6]) or 'ไม่มีเลย'), 'ERR')
            await self._b_settle(page)
            return False
        try:
            await opt.click(timeout=6000)
            await page.wait_for_timeout(400)
        except Exception as ex:
            self.log('   ✗ %s (%s) กดตัวเลือกไม่ได้: %s' % (who, what, str(ex)[:60]), 'ERR')
            await self._b_settle(page)
            return False

        await self._b_settle(page)
        r = await page.evaluate(JS_BUNDLE_ACT, [how, str(key), what, 'read', '', TIERS, nth])
        got = str(r).split('|', 1)[1] if '|' in str(r) else ''
        if not _match(got):
            self.log('   ✗ %s %s ตั้งเป็น "%s" ไม่สำเร็จ (ตอนนี้ "%s")'
                     % (who, what, want, got), 'ERR')
            return False
        self.log('   · %s %s = %s' % (who, what, want), 'OK')
        return True

    async def _b_finish_wallet(self, page, name, is_fame):
        """เก็บงานการ์ด famepoint / exp ให้ครบ

        ทั้งคู่ "ต้องมี Tier" ไม่งั้นเว็บจะกดสร้างไม่ผ่าน
        และ famepoint ต้องเป็นเครดิตเรียลไทม์ ไม่ใช่เครดิตธรรมดา
        """
        await self._b_expand(page)
        ok = True
        if is_fame:
            if not await self._b_choose(page, 'name', name, 'kind', FAME_CREDIT_TYPE,
                                        exact=False, label=name):
                self.log('   ⚠ %s ยังไม่ใช่เครดิตเรียลไทม์ — ต้องแก้เองบนเว็บก่อนกดสร้าง'
                         % name, 'ERR')
                ok = False
        if not await self._b_choose(page, 'name', name, 'tier', DEFAULT_TIER, label=name):
            self.log('   ⚠ %s ยังไม่ได้เลือก Tier — เว็บจะกดสร้างไม่ผ่าน' % name, 'ERR')
            ok = False
        return ok

    async def _b_set_row(self, page, item_id, qty, tier, nth=1):
        """ตั้งจำนวน + Tier ของการ์ด "ของไอเทมนี้ ใบที่ nth"

        ไอเทมตัวเดียวกันใส่ซ้ำหลายใบได้ ถ้าไม่บอกว่าใบที่เท่าไหร่
        มันจะไปตั้งใบแรกซ้ำๆ แล้วใบหลังๆ ไม่ได้ Tier (เว็บจะกดสร้างไม่ผ่าน)
        """
        who = 'ไอเทม %s' % item_id + (' (ใบที่ %d)' % nth if nth > 1 else '')
        r = await page.evaluate(JS_BUNDLE_ACT,
                                ['id', str(item_id), 'qty', 'set', str(qty), TIERS, nth])
        if r != 'ok':
            self.log('   ! ตั้งจำนวนของ%s ไม่ได้ (%s)' % (who, r), 'WARN')
        if tier:
            await self._b_choose(page, 'id', item_id, 'tier', tier, label=who, nth=nth)

    async def _b_pick(self, page, label, value):
        r = await page.evaluate(JS_PICK_TYPE, [value, label])
        if r == 'ok':
            self.log('   · %s = %s' % (label, value), 'INFO')
            return True
        try:
            trig = page.locator('[role="combobox"]').first
            await trig.click(timeout=4000)
            await page.wait_for_timeout(400)
            opt = page.locator('[role="option"]:has-text("%s")' % value).first
            if await opt.count() > 0:
                await opt.click(timeout=4000)
                return True
        except Exception:
            pass
        self.log('   ! ตั้ง %s = %s ไม่ได้' % (label, value), 'WARN')
        return False

    # ------------------------------------------------------------------
    #  แท็บตรวจระบบ
    # ------------------------------------------------------------------
    def _build_check(self):
        p = self.tab_check

        # ---- แถบตั้งค่าที่เก็บ log ----
        box = tk.LabelFrame(p, text='  ที่เก็บบันทึกการใช้งาน  ', bg=C['bg'], fg=C['dim'],
                            font=('Segoe UI', 9, 'bold'), bd=1, relief='solid')
        box.pack(fill='x', padx=14, pady=(12, 0))
        inner = tk.Frame(box, bg=C['bg'])
        inner.pack(fill='x', padx=12, pady=10)

        r1 = tk.Frame(inner, bg=C['bg'])
        r1.pack(fill='x')
        tk.Label(r1, text='บันทึกในเครื่อง', bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9),
                 width=16, anchor='w').pack(side='left')
        self.lbl_local = tk.Label(r1, text='', bg=C['bg'], fg=C['fg'],
                                  font=('Consolas', 9), anchor='w')
        self.lbl_local.pack(side='left')
        self._btn(r1, '📂  เปิด', self.open_log_dir).pack(side='right', padx=(8, 0), ipadx=8, ipady=2)
        self._btn(r1, '↺  คืนค่าเดิม', self.reset_local).pack(side='right', padx=(8, 0), ipadx=8, ipady=2)
        self._btn(r1, '📁  เลือกโฟลเดอร์', self.pick_local).pack(side='right', ipadx=8, ipady=2)

        r2 = tk.Frame(inner, bg=C['bg'])
        r2.pack(fill='x', pady=(8, 0))
        tk.Label(r2, text='โฟลเดอร์กลางของทีม', bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9),
                 width=16, anchor='w').pack(side='left')
        self.lbl_central = tk.Label(r2, text='', bg=C['bg'], fg=C['fg'],
                                    font=('Consolas', 9), anchor='w')
        self.lbl_central.pack(side='left')
        self._btn(r2, '✕  ล้าง', self.clear_central).pack(side='right', padx=(8, 0), ipadx=8, ipady=2)
        self._btn(r2, '📁  เลือกโฟลเดอร์', self.pick_central).pack(side='right', ipadx=8, ipady=2)

        r3 = tk.Frame(inner, bg=C['bg'])
        r3.pack(fill='x', pady=(8, 0))
        tk.Label(r3, text='โฟลเดอร์โปรแกรม', bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9),
                 width=16, anchor='w').pack(side='left')
        tk.Label(r3, text=DATA_DIR + '   (เก็บ login ของ Chrome ไว้ที่นี่ ย้ายไม่ได้)',
                 bg=C['bg'], fg=C['dim'], font=('Consolas', 9), anchor='w').pack(side='left')
        self._btn(r3, '📂  เปิด', self.open_data_dir).pack(side='right', ipadx=8, ipady=2)

        self.lbl_todo = tk.Label(inner, text='', bg=C['bg'], fg=C['warn'],
                                 font=('Segoe UI', 9), wraplength=980, justify='left', anchor='w')
        self.lbl_todo.pack(fill='x', pady=(9, 0))
        self._refresh_central()

        bar = tk.Frame(p, bg=C['bg'])
        bar.pack(fill='x', padx=14, pady=12)
        self.btn_chk_logic = self._btn(bar, '⚡  ตรวจตรรกะ (เร็ว ไม่ต้องต่อเน็ต)', self.run_check_logic)
        self.btn_chk_logic.pack(side='left', ipadx=10, ipady=5)
        self.btn_chk_full = self._btn(bar, '🌐  ตรวจเต็ม (เปิดเว็บจริง)', self.run_check_full, primary=True)
        self.btn_chk_full.pack(side='left', padx=8, ipadx=10, ipady=4)
        self._btn(bar, '📋  คัดลอกผล', self.copy_check).pack(side='right', ipadx=10, ipady=5)
        self.lbl_chk = tk.Label(bar, text='', bg=C['bg'], fg=C['dim'], font=FB)
        self.lbl_chk.pack(side='right', padx=14)

        tk.Label(p, text='ใช้ตรวจว่าเครื่องมือยังทำงานถูกต้องไหม — "ตรวจเต็ม" จะเปิดเว็บจริงแล้วไล่เช็กทีละจุด '
                         'ถ้าเว็บเปลี่ยนหน้าตา จะบอกได้เลยว่าพังตรงไหนและต้องแก้ค่าตัวไหน',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9), wraplength=1000,
                 justify='left').pack(anchor='w', padx=14)

        wrap = tk.Frame(p, bg=C['bg'])
        wrap.pack(fill='both', expand=True, padx=14, pady=(10, 14))
        cols = ('st', 'name', 'detail')
        self.chk_tree = ttk.Treeview(wrap, columns=cols, show='headings', style='TR.Treeview')
        for c, t, w in (('st', 'ผล', 60), ('name', 'รายการตรวจ', 420), ('detail', 'รายละเอียด', 480)):
            self.chk_tree.heading(c, text=t)
            self.chk_tree.column(c, width=w, anchor='w')
        self.chk_tree.tag_configure('ok', foreground=C['ok'])
        self.chk_tree.tag_configure('bad', foreground=C['err'])
        sb = ttk.Scrollbar(wrap, orient='vertical', command=self.chk_tree.yview)
        self.chk_tree.configure(yscrollcommand=sb.set)
        self.chk_tree.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')
        self.check_rows = []

    def open_log_dir(self):
        try:
            os.startfile(LOCAL_LOG_DIR)
        except Exception:
            messagebox.showinfo('ที่เก็บ log', LOCAL_LOG_DIR)

    def pick_local(self):
        path = filedialog.askdirectory(title='เลือกที่เก็บบันทึกการใช้งานในเครื่องนี้')
        if not path:
            return
        moved, err = set_local_log_dir(path)
        self._refresh_central()
        if err:
            messagebox.showerror('ย้ายที่เก็บ log', 'ย้ายไม่สำเร็จ: ' + err)
            return
        log_event('set_local_log_dir', path=path, moved=moved)
        self.log(f'ย้ายที่เก็บ log ไปที่ {path} (ย้ายไฟล์เดิม {moved} ไฟล์)', 'OK')
        messagebox.showinfo('ย้ายที่เก็บ log',
                            f'เรียบร้อย\n\nที่ใหม่: {path}\nย้ายประวัติเดิมตามไปให้ {moved} ไฟล์')

    def reset_local(self):
        if os.path.abspath(LOCAL_LOG_DIR) == os.path.abspath(DATA_DIR):
            return
        moved, err = set_local_log_dir(DATA_DIR)
        self._refresh_central()
        self.log(f'คืนค่าที่เก็บ log กลับเป็นค่าเริ่มต้น (ย้ายไฟล์ {moved})', 'OK')

    def _refresh_central(self):
        default = os.path.abspath(LOCAL_LOG_DIR) == os.path.abspath(DATA_DIR)
        self.lbl_local.config(text=LOCAL_LOG_DIR + ('   (ค่าเริ่มต้น)' if default else ''),
                              fg=C['dim'] if default else C['fg'])
        self.lbl_central.config(
            text=CENTRAL_LOG_DIR or '— ยังไม่ได้ตั้ง (เก็บเฉพาะในเครื่องนี้) —',
            fg=C['fg'] if CENTRAL_LOG_DIR else C['dim'])
        todo = pending_items()
        self.lbl_todo.config(text=('⚠  ' + todo[0]) if todo else
                             '✓  ตั้งค่าครบแล้ว — log จะถูกเขียนทั้งในเครื่องและโฟลเดอร์กลาง',
                             fg=C['warn'] if todo else C['ok'])

    def open_data_dir(self):
        try:
            os.startfile(DATA_DIR)
        except Exception:
            messagebox.showinfo('ที่เก็บข้อมูล', DATA_DIR)

    def pick_central(self):
        path = filedialog.askdirectory(title='เลือกโฟลเดอร์กลางของทีม (ที่ทุกคนเข้าถึงได้)')
        if not path:
            return
        set_central_dir(path)
        self._refresh_central()
        log_event('set_central_dir', path=path)
        self.log(f'ตั้งโฟลเดอร์กลางเป็น {path}', 'OK')
        messagebox.showinfo('โฟลเดอร์กลาง',
                            'ตั้งค่าเรียบร้อย\n\nจากนี้ log จะถูกเขียนทั้งในเครื่องและในโฟลเดอร์นี้\n'
                            'ให้คนอื่นในทีมตั้งโฟลเดอร์เดียวกัน แล้วจะรวมกันเอง')

    def clear_central(self):
        if not CENTRAL_LOG_DIR:
            return
        set_central_dir('')
        self._refresh_central()
        self.log('ล้างโฟลเดอร์กลางแล้ว — กลับไปเก็บเฉพาะในเครื่อง', 'WARN')

    def _show_check(self, results, append=False):
        if not append:
            self.chk_tree.delete(*self.chk_tree.get_children())
            self.check_rows = []
        self.check_rows.extend(results)
        for x in results:
            detail = x['detail']
            if not x['ok'] and x['fix']:
                detail = (detail + '   →  ' + x['fix']).strip(' →')
            self.chk_tree.insert('', 'end', values=('ผ่าน' if x['ok'] else 'พัง', x['name'], detail),
                                 tags=('ok' if x['ok'] else 'bad',))
        bad = sum(1 for x in self.check_rows if not x['ok'])
        total = len(self.check_rows)
        self.lbl_chk.config(text=f'ผ่าน {total - bad}/{total}' + ('  ✓ ปกติดี' if not bad else f'  ✗ พัง {bad} จุด'),
                            fg=C['ok'] if not bad else C['err'])

    def copy_check(self):
        if not self.check_rows:
            return messagebox.showinfo('ตรวจระบบ', 'ยังไม่ได้ตรวจ')
        lines = [f'TR Plus Ultra v{APP_VERSION} — ผลตรวจระบบ {datetime.now():%Y-%m-%d %H:%M}', '']
        for x in self.check_rows:
            mark = 'OK  ' if x['ok'] else 'FAIL'
            lines.append(f'[{mark}] {x["name"]}  |  {x["detail"]}'
                         + (f'  ->  {x["fix"]}' if (not x['ok'] and x['fix']) else ''))
        bad = sum(1 for x in self.check_rows if not x['ok'])
        lines += ['', f'สรุป: ผ่าน {len(self.check_rows) - bad}/{len(self.check_rows)}']
        self.root.clipboard_clear()
        self.root.clipboard_append('\n'.join(lines))
        self.log('คัดลอกผลตรวจแล้ว ส่งให้คนดูแลเครื่องมือได้เลย', 'OK')

    def run_check_logic(self):
        self.nb.select(self.tab_check)
        try:
            res = run_logic_tests()
            self._show_check(res)
            bad = sum(1 for x in res if not x['ok'])
            log_event('selfcheck', mode='logic', total=len(res), failed=bad,
                      failed_items=[x['name'] for x in res if not x['ok']][:20])
            self.log('ตรวจตรรกะเสร็จ', 'OK')
        except Exception as ex:
            messagebox.showerror('ตรวจระบบ', str(ex))

    def run_check_full(self):
        if self.running:
            return messagebox.showinfo('กำลังทำงาน', 'รอให้งานปัจจุบันเสร็จก่อนนะ')
        self.nb.select(self.tab_check)
        self._show_check(run_logic_tests())
        self.running = True
        self.btn_chk_full.config(state='disabled', text='กำลังตรวจ...')
        threading.Thread(target=self._check_thread, daemon=True).start()

    def _check_thread(self):
        try:
            asyncio.run(self._check_web())
        except Exception as ex:
            msg = str(ex)[:120]
            self.root.after(0, lambda m=msg: self._show_check(
                [_t(False, 'เปิดเบราว์เซอร์ได้', m)], append=True))
        finally:
            self.running = False
            self.root.after(0, lambda: self.btn_chk_full.config(
                state='normal', text='🌐  ตรวจเต็ม (เปิดเว็บจริง)'))

    async def _check_web(self):
        self.log('เริ่มตรวจเว็บจริง...', 'STEP')
        async with async_playwright() as pw:
            browser = await pw.chromium.launch_persistent_context(**launch_kwargs(True))
            page = browser.pages[0] if browser.pages else await browser.new_page()
            try:
                res = await run_web_tests(page, log=lambda m: self.log(m))
            finally:
                try:
                    await browser.close()
                except Exception:
                    pass
        self.root.after(0, lambda: self._show_check(res, append=True))
        bad = sum(1 for x in res if not x['ok'])
        log_event('selfcheck', mode='web', total=len(res), failed=bad,
                  failed_items=[x['name'] for x in res if not x['ok']][:20])
        self.log(f'ตรวจเว็บเสร็จ — พัง {bad} จุด' if bad else 'ตรวจเว็บเสร็จ — ปกติดีทุกจุด',
                 'WARN' if bad else 'OK')

    # ---------- แจ้งบั๊ก ----------
    def bug_open_sheet(self):
        """เปิดชีทแจ้งบั๊กของทีม — ไปพิมพ์ในชีทได้เลย"""
        try:
            import webbrowser
            webbrowser.open(BUG_SHEET_URL)
            self.log('เปิดหน้าแจ้งบั๊ก — %s · %s · v%s'
                     % (TR_USER, TR_MACHINE, APP_VERSION), 'OK')
        except Exception as ex:
            self.log('เปิดหน้าแจ้งบั๊กไม่ได้: %s' % str(ex)[:80], 'ERR')
            messagebox.showerror('เปิดไม่ได้', str(ex))

    # ==================================================================
    #  แท็บ WR Master — สร้าง Item Code จากชีท Master Code
    # ==================================================================
    def _build_wr(self):
        p = self.tab_wr
        self.wq = []                       # คิวโค้ดที่จะสร้าง
        self.w_running = False
        self.w_cancel = False

        s1 = self._card(p, 'นำเข้าจากไฟล์ต้นฉบับ')
        bar = tk.Frame(s1, bg=C['bg'])
        bar.pack(fill='x')
        self._btn(bar, '📂  นำเข้าไฟล์ต้นฉบับ (.xlsx)', self.w_import,
                  primary=True).pack(side='left', ipadx=10, ipady=4)
        self._btn(bar, '✏  แก้ไข', self.w_edit).pack(side='left', padx=(8, 0),
                                                     ipadx=8, ipady=4)
        self._btn(bar, '🗑  ลบที่เลือก', self.w_del).pack(side='left', padx=(8, 0),
                                                          ipadx=8, ipady=4)
        self._btn(bar, 'ล้างคิว', self.w_clear).pack(side='left', padx=(8, 0),
                                                     ipadx=8, ipady=4)
        self.w_count = tk.Label(bar, text='คิวว่าง', bg=C['bg'], fg=C['dim'], font=FM)
        self.w_count.pack(side='right')
        tk.Label(s1, text='อ่านเฉพาะชีทแบบ Master Code (มีแถว CODE) — '
                          'ชีทบอก Limit เท่าไหร่ จะเผื่อให้อีก %d เสมอ  ·  '
                          'ดับเบิลคลิกเพื่อแก้รายตัว' % WR_SPARE,
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), anchor='w',
                 justify='left').pack(fill='x', pady=(8, 0))

        s2 = self._card(p, 'คิว Item Code ที่จะสร้าง', grow=True)
        tw = tk.Frame(s2, bg=C['bg'])
        tw.pack(fill='both', expand=True)
        wc = ('n', 'code', 'wname', 'start', 'end', 'cap', 'bundle')
        self.w_tree = ttk.Treeview(tw, columns=wc, show='headings', height=7,
                                   style='TR.Treeview', selectmode='extended')
        for c, t, w in (('n', '#', 38), ('code', 'CODE', 120),
                        ('wname', 'ชื่อ Item Code', 220), ('start', 'เริ่มใช้งาน', 140),
                        ('end', 'สิ้นสุด', 140), ('cap', 'จำกัดจำนวน', 90),
                        ('bundle', 'Bundle', 80)):
            self.w_tree.heading(c, text=t)
            self.w_tree.column(c, width=w, anchor='w')
        self.w_tree.tag_configure('warn', background='#3a3018', foreground='#e3b341')
        wsb = ttk.Scrollbar(tw, orient='vertical', command=self.w_tree.yview)
        self.w_tree.configure(yscrollcommand=wsb.set)
        self.w_tree.pack(side='left', fill='both', expand=True)
        wsb.pack(side='right', fill='y')
        self.w_tree.bind('<Double-1>', lambda e: self.w_edit())
        self.w_warn = tk.Label(s2, text='', bg=C['bg'], fg=C['warn'],
                               font=('Segoe UI', 9), anchor='w', justify='left',
                               wraplength=940)
        self.w_warn.pack(fill='x', pady=(8, 0))

        s3 = self._card(p, 'ลงมือสร้าง')
        self.wv_do = tk.BooleanVar(value=False)
        tk.Checkbutton(s3, variable=self.wv_do, command=self._w_do_changed,
                       text='กดปุ่ม “สร้าง Item Code” จริง',
                       bg=C['bg'], fg=C['fg'], selectcolor=C['input'],
                       activebackground=C['bg'], activeforeground=C['fg'],
                       font=FB, bd=0, highlightthickness=0).grid(row=0, column=0,
                                                                 sticky='w')
        self.w_mode = tk.Label(s3, text='', bg=C['bg'], fg=C['warn'], font=FM)
        self.w_mode.grid(row=0, column=1, sticky='w', padx=(10, 0))
        self.wv_hold = tk.StringVar(value=self.prefs.get('w_hold', '3'))
        tk.Label(s3, text='โหมดทดสอบ: กรอกเสร็จแล้วค้างหน้าไว้ให้ดู (วินาที)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 9)
                 ).grid(row=1, column=0, sticky='w', pady=(8, 0))
        eh = self._entry(s3, width=6)
        eh.config(textvariable=self.wv_hold)
        eh.grid(row=1, column=1, sticky='w', padx=(10, 0), pady=(8, 0), ipady=3)
        run = tk.Frame(s3, bg=C['bg'])
        run.grid(row=2, column=0, columnspan=3, sticky='w', pady=(12, 0))
        self.w_btn_run = self._btn(run, '▶  เริ่มทำงาน', self.w_start, primary=True)
        self.w_btn_run.pack(side='left', ipadx=18, ipady=5)
        self.w_btn_stop = self._btn(run, '■  ยกเลิก', self.w_stop)
        self.w_btn_stop.config(state='disabled')
        self.w_btn_stop.pack(side='left', padx=(8, 0), ipadx=12, ipady=5)

        # ---- ผลลัพธ์ ----
        self.made_codes = []
        ib = tk.Frame(p, bg=C['bg'])
        ib.pack(fill='x', padx=14, pady=(10, 4))
        iw = tk.Frame(p, bg=C['bg'])
        iw.pack(fill='both', expand=True, padx=14, pady=(0, 12))
        tk.Label(ib, text='Item Code ที่สร้างแล้ว', bg=C['bg'], fg=C['dim'],
                 font=('Segoe UI', 9, 'bold')).pack(side='left')
        self.lbl_wmade = tk.Label(ib, text='ยังไม่ได้สร้าง', bg=C['bg'], fg=C['dim'],
                                  font=('Segoe UI', 9))
        self.lbl_wmade.pack(side='left', padx=10)
        self._btn(ib, '📋  คัดลอก CODE', self.copy_made_codes).pack(
            side='right', ipadx=8, ipady=2)
        tk.Label(ib, text='(เลือกแถวที่ต้องการก่อน · ไม่เลือก = เอาทั้งหมด)',
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8)).pack(side='right', padx=8)
        mc = ('no', 'code', 'cid', 'cname', 'at', 'note')
        self.tree_code = ttk.Treeview(iw, columns=mc, show='headings', height=5,
                                      style='TR.Treeview', selectmode='extended')
        for c, t, w in (('no', '#', 42), ('code', 'CODE', 130),
                        ('cid', 'Item Code Id', 100), ('cname', 'ชื่อ Item Code', 230),
                        ('at', 'เวลา', 78), ('note', 'หมายเหตุ', 150)):
            self.tree_code.heading(c, text=t)
            self.tree_code.column(c, width=w, anchor='w')
        self.tree_code.tag_configure('bad', background='#3a1f1e', foreground='#f0736a')
        csb2 = ttk.Scrollbar(iw, orient='vertical', command=self.tree_code.yview)
        self.tree_code.configure(yscrollcommand=csb2.set)
        self.tree_code.pack(side='left', fill='both', expand=True)
        csb2.pack(side='right', fill='y')

        self._w_do_changed()
        self._w_refresh()

    def _w_do_changed(self):
        if self.wv_do.get():
            self.w_mode.config(text='⚠  จะสร้าง Item Code จริงบนเว็บ', fg=C['err'])
        else:
            self.w_mode.config(text='โหมดทดสอบ — กรอกฟอร์มให้ดูเฉยๆ ไม่กดสร้าง', fg=C['ok'])

    def _w_refresh(self):
        self.w_tree.delete(*self.w_tree.get_children())
        nobd = 0
        for i, d in enumerate(self.wq, 1):
            bad = not d.get('bundle')
            if bad:
                nobd += 1
            self.w_tree.insert('', 'end', tags=('warn',) if bad else (), values=(
                i, d['code'], d['name'], d.get('start', '—'), d.get('end', '—'),
                d.get('cap') or 'ไม่จำกัด', d.get('bundle') or '⚠ ไม่มี'))
        self.w_count.config(text='คิวว่าง' if not self.wq
                            else 'ในคิว %d โค้ด' % len(self.wq))
        self.w_warn.config(
            text=('⚠  มี %d โค้ดที่ยังไม่มีเลข Bundle ในชีท (แถวเหลือง) — '
                  'ดับเบิลคลิกใส่เองก่อน ไม่งั้นตัวนั้นจะเลือก Bundle ไม่ได้' % nobd)
            if nobd else '', fg=C['warn'])

    def _w_sel(self):
        s = self.w_tree.selection()
        return self.w_tree.index(s[0]) if s else None

    def w_import(self):
        if self.w_running:
            return
        path = filedialog.askopenfilename(
            title='เลือกไฟล์ต้นฉบับ',
            filetypes=[('Excel', '*.xlsx *.xlsm'), ('ทุกไฟล์', '*.*')])
        if not path:
            return
        dlg = CodeImportDialog(self.root, path)
        if not dlg.result:
            return
        have = {d['code'] for d in self.wq}
        add = [d for d in dlg.result if d['code'] not in have]
        self.wq.extend(add)
        self._w_refresh()
        self.log('นำเข้า Item Code %d โค้ด (รวม %d)' % (len(add), len(self.wq)), 'OK')
        for w in (dlg.warns or [])[:20]:
            self.log('  ⚠ ' + w, 'WARN')
        if len(add) < len(dlg.result):
            self.log('  ข้ามโค้ดที่อยู่ในคิวอยู่แล้ว %d' % (len(dlg.result) - len(add)),
                     'INFO')
        log_event('wr_import', count=len(add), total=len(self.wq),
                  file=os.path.basename(path))

    def w_edit(self):
        if self.w_running:
            return
        i = self._w_sel()
        if i is None:
            return messagebox.showinfo('แก้ไข', 'เลือกแถวในคิวก่อนนะ')
        d = dict(self.wq[i])
        if self._w_form(d, 'แก้ไข Item Code #%d' % (i + 1)):
            self.wq[i] = d
            self._w_refresh()

    def w_del(self):
        i = self._w_sel()
        if i is None or self.w_running:
            return
        del self.wq[i]
        self._w_refresh()

    def w_clear(self):
        if self.w_running or not self.wq:
            return
        if messagebox.askyesno('ล้างคิว', 'ลบทั้ง %d โค้ดออกจากคิว?' % len(self.wq)):
            self.wq = []
            self._w_refresh()

    def _w_form(self, d, title):
        """หน้าต่างแก้ไข Item Code หนึ่งตัว — คืน True ถ้ากดบันทึก"""
        top = tk.Toplevel(self.root)
        top.title(title)
        top.configure(bg=C['bg'])
        top.transient(self.root)
        top.grab_set()
        try:
            self.root.update_idletasks()
            x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - 620) // 2)
            y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - 430) // 2)
            top.geometry('620x430+%d+%d' % (x, y))
        except Exception:
            top.geometry('620x430')
        ok = {'v': False}
        body = tk.Frame(top, bg=C['bg'])
        body.pack(fill='both', expand=True, padx=18, pady=14)
        fields = (('CODE', 'code', 26), ('ชื่อ Item Code (ไทย/อังกฤษ)', 'name', 44),
                  ('เวลาเริ่มใช้งาน', 'start', 26), ('เวลาสิ้นสุด', 'end', 26),
                  ('จำนวนการใช้งานต่อ 1 User', 'per_user', 10),
                  ('จำกัดจำนวน (เว้นว่าง = ไม่จำกัด)', 'cap', 12),
                  ('เลข Bundle', 'bundle', 12))
        vs = {}
        for r, (lbl, key, w) in enumerate(fields):
            tk.Label(body, text=lbl, bg=C['bg'], fg=C['dim'],
                     font=('Segoe UI', 9)).grid(row=r, column=0, sticky='w', pady=5)
            v = tk.StringVar(value=str(d.get(key, '') or ''))
            vs[key] = v
            e = self._entry(body, width=w)
            e.config(textvariable=v)
            e.grid(row=r, column=1, sticky='w', padx=(12, 0), ipady=3)
        tk.Label(body, text='เวลาใส่แบบ 2026-10-01 13:00:00  ·  '
                            'ชีทบอก Limit เท่าไหร่ โปรแกรมเผื่อให้อีก %d แล้ว' % WR_SPARE,
                 bg=C['bg'], fg=C['dim'], font=('Segoe UI', 8), justify='left'
                 ).grid(row=len(fields), column=0, columnspan=2, sticky='w', pady=(10, 0))
        foot = tk.Frame(top, bg=C['card'], height=58)
        foot.pack(fill='x', side='bottom')
        foot.pack_propagate(False)

        def _save():
            for key, v in vs.items():
                d[key] = v.get().strip()
            d['unlimited'] = not d.get('cap')
            if not d['code'] or not d['name']:
                return messagebox.showwarning('กรอกไม่ครบ', 'ต้องมี CODE และชื่อ Item Code',
                                              parent=top)
            ok['v'] = True
            top.destroy()
        self._btn(foot, 'บันทึก', _save, primary=True).pack(side='right', padx=18,
                                                            pady=12, ipadx=20, ipady=4)
        self._btn(foot, 'ยกเลิก', top.destroy).pack(side='right', pady=12,
                                                    ipadx=14, ipady=4)
        self.root.wait_window(top)
        return ok['v']

    def add_made_code(self, d, cid, notes='', ok=True):
        row = {'code': d.get('code', ''), 'name': d.get('name', ''),
               'id': str(cid or '').strip() or '-', 'notes': notes, 'ok': bool(ok),
               'at': datetime.now().strftime('%H:%M:%S')}
        self.made_codes.append(row)
        row['no'] = len(self.made_codes)

        def _do():
            self.tree_code.insert('', 'end', tags=() if row['ok'] else ('bad',),
                                  values=(row['no'], row['code'], row['id'],
                                          row['name'], row['at'], row['notes']))
            self.tree_code.yview_moveto(1)
            got = sum(1 for m in self.made_codes if m['id'] != '-')
            self.lbl_wmade.config(
                text='สร้างแล้ว %d โค้ด%s' % (
                    len(self.made_codes),
                    '' if got == len(self.made_codes) else ' (ได้เลข %d)' % got),
                fg=C['ok'])
        self.root.after(0, _do)

    def copy_made_codes(self):
        """คัดลอกเฉพาะ CODE ของแถวที่เลือกไว้ — ไม่เลือกก็เอาทั้งหมด"""
        rows = list(self.tree_code.selection())
        picked = bool(rows)
        if not picked:
            rows = list(self.tree_code.get_children())
        if not rows:
            messagebox.showinfo('ยังไม่มี', 'ยังไม่ได้สร้าง Item Code ในรอบนี้')
            return
        out = []
        for w in rows:
            v = self.tree_code.item(w, 'values')
            if len(v) > 1 and str(v[1]).strip():
                out.append(str(v[1]).strip())
        if not out:
            messagebox.showinfo('ไม่มีโค้ด', 'แถวที่เลือกไม่มี CODE')
            return
        self.root.clipboard_clear()
        self.root.clipboard_append('\n'.join(out))
        self.log('คัดลอก CODE %d รายการแล้ว (%s)'
                 % (len(out), 'เฉพาะที่เลือก' if picked else 'ทั้งหมด'), 'OK')

    def w_stop(self):
        self.w_cancel = True
        self.log('กำลังยกเลิกการสร้าง Item Code...', 'WARN')

    def w_start(self):
        if self.w_running or self.running or self.c_running or self.b_running:
            return messagebox.showinfo('กำลังทำงาน', 'รอให้งานปัจจุบันเสร็จก่อนนะ')
        if not self.wq:
            return messagebox.showwarning('คิวว่าง',
                                          'ยังไม่มีโค้ดในคิว — นำเข้าไฟล์ต้นฉบับก่อน')
        do = self.wv_do.get()
        nobd = sum(1 for d in self.wq if not d.get('bundle'))
        if do:
            msg = 'จะสร้าง Item Code จริงบนเว็บ %d โค้ด\n' % len(self.wq)
            if nobd:
                msg += ('\n⚠  มี %d โค้ดที่ยังไม่มีเลข Bundle — ตัวนั้นจะเลือก Bundle '
                        'ไม่ได้\n' % nobd)
            msg += '\nตรวจชื่อ/เวลา/จำนวนในคิวเรียบร้อยแล้วใช่ไหม?'
            if not messagebox.askyesno('ยืนยัน', msg):
                return
        self.save_now()
        self.w_running = True
        self.w_cancel = False
        self.w_btn_run.config(state='disabled')
        self.w_btn_stop.config(state='normal')
        self.nb.select(self.tab_log)
        self.log('=' * 46, 'STEP')
        self.log(('เริ่มสร้าง Item Code จริง ' if do else 'เริ่มทดสอบกรอกฟอร์ม ')
                 + '%d โค้ด' % len(self.wq), 'STEP')
        log_event('wr_start', count=len(self.wq), commit=bool(do), no_bundle=nobd)
        threading.Thread(target=self._w_thread, args=(list(self.wq), do),
                         daemon=True).start()

    def _w_thread(self, rows, do):
        try:
            asyncio.run(self._w_work(rows, do))
        except Exception as ex:
            log_event('error', where='wr', message=str(ex)[:300])
            self.log('ผิดพลาด: ' + str(ex), 'ERR')
            self.log(traceback.format_exc(), 'ERR')
        finally:
            self.w_running = False

            def _rst():
                self.w_btn_run.config(state='normal')
                self.w_btn_stop.config(state='disabled')
                if self.made_codes:
                    self.nb.select(self.tab_wr)
            self.root.after(0, _rst)

    async def _w_work(self, rows, do):
        async with async_playwright() as pw:
            browser = await pw.chromium.launch_persistent_context(**launch_kwargs(False))
            page = browser.pages[0] if browser.pages else await browser.new_page()
            try:
                try:
                    hold = max(0, int(float(self.wv_hold.get() or 0)))
                except Exception:
                    hold = 3
                okc = errc = 0
                for i, d in enumerate(rows, 1):
                    if self.w_cancel:
                        self.log('ยกเลิกแล้ว', 'WARN')
                        break
                    self.set_progress(i - 1, len(rows), d['code'])
                    self.log('[%d/%d] %s  (%s)' % (i, len(rows), d['code'], d['name']),
                             'STEP')
                    try:
                        if await self._w_one(page, d, do, hold):
                            okc += 1
                        else:
                            errc += 1
                    except Exception as ex:
                        errc += 1
                        self.log('   ✗ ' + str(ex)[:160], 'ERR')
                        log_event('error', where='wr_one', code=d.get('code'),
                                  message=str(ex)[:200])
                self.set_progress(len(rows), len(rows), 'เสร็จ')
                self.log('จบ — สำเร็จ %d · ไม่ผ่าน %d' % (okc, errc),
                         'OK' if not errc else 'WARN')
                log_event('wr_done', ok=okc, fail=errc, commit=bool(do))
            finally:
                try:
                    await browser.close()
                except Exception:
                    pass

    def _build_log(self):
        self.log_box = scrolledtext.ScrolledText(
            self.tab_log, bg=C['input'], fg=C['fg'], bd=0, font=('Consolas', 9),
            insertbackground=C['fg'], wrap='word')
        self.log_box.pack(fill='both', expand=True, padx=14, pady=12)
        for tag, col in (('INFO', C['dim']), ('STEP', C['accent']), ('OK', C['ok']),
                         ('WARN', C['warn']), ('ERR', C['err'])):
            self.log_box.tag_config(tag, foreground=col)

    # ---------- utility ----------
    def log(self, msg, kind='INFO'):
        def _do():
            self.log_box.insert('end', f'[{datetime.now():%H:%M:%S}] {msg}\n', kind)
            self.log_box.see('end')
        self.root.after(0, _do)

    def set_progress(self, done, total, note=''):
        def _do():
            self.progress['value'] = (done / total * 100) if total else 0
            self.lbl_stat.config(text=f'{done}/{total}  {note}' if total else note)
        self.root.after(0, _do)

    def add_result(self, row, notes='', seq=0, req=''):
        if row.get('id') and any(r['id'] == row['id'] for r in self.results):
            return
        row = dict(row)
        row['notes'] = notes
        row['seq'] = seq                 # ลำดับในรายการที่สั่งค้น — ผูกผลลัพธ์กลับไปหาต้นทาง
        row['req'] = req                 # ItemKind ที่สั่งให้หา
        row['miss'] = False
        row['dup'] = False
        self.results.append(row)

        def _do():
            self._ins_result(row)
            self.lbl_count.config(text=f'พบ {len(self.results)} รายการ')
        self.root.after(0, _do)

    def add_made_item(self, d, aztek_id, notes='', pic=False, ok=True):
        """จดไว้ว่าสร้างไอเทมอะไรไปแล้ว ได้ Aztek Item Id อะไร — โชว์ในแท็บสร้าง Item"""
        row = {'name': d.get('name', ''), 'kind': str(d.get('kind') or ''),
               'id': str(aztek_id or '').strip() or '-', 'notes': notes,
               'pic': '🖼' if pic else '—', 'ok': bool(ok),
               'at': datetime.now().strftime('%H:%M:%S')}
        self.made_items.append(row)
        row['no'] = len(self.made_items)

        def _do():
            self.tree_item.insert('', 'end', tags=() if row['ok'] else ('bad',),
                                  values=(row['no'], row['id'], row['kind'],
                                          row['name'], row['pic'], row['at'],
                                          row['notes']))
            self.tree_item.yview_moveto(1)
            got = sum(1 for m in self.made_items if m['id'] != '-')
            self.lbl_item.config(
                text='สร้างแล้ว %d ไอเทม%s' % (
                    len(self.made_items),
                    '' if got == len(self.made_items)
                    else ' (ได้เลข %d)' % got), fg=C['ok'])
        self.root.after(0, _do)

    def copy_made_items(self):
        """คัดลอกเฉพาะ "เลข Aztek Item Id" ของแถวที่เลือกไว้ — ไม่เลือกก็เอาทั้งหมด"""
        rows = list(self.tree_item.selection())
        picked = bool(rows)
        if not picked:
            rows = list(self.tree_item.get_children())
        if not rows:
            messagebox.showinfo('ยังไม่มี', 'ยังไม่ได้สร้างไอเทมในรอบนี้')
            return
        ids, blank = [], 0
        for w in rows:
            v = self.tree_item.item(w, 'values')
            aid = str(v[1]).strip() if len(v) > 1 else ''
            if aid and aid != '-':
                ids.append(aid)
            else:
                blank += 1
        if not ids:
            messagebox.showinfo('ไม่มีเลข',
                                'แถวที่เลือกยังไม่ได้เลข Aztek Item Id (ขึ้น - อยู่)')
            return
        self.root.clipboard_clear()
        self.root.clipboard_append('\n'.join(ids))
        self.log('คัดลอก Aztek Item Id %d รายการแล้ว (%s)%s' % (
            len(ids), 'เฉพาะที่เลือก' if picked else 'ทั้งหมด',
            '  ข้ามที่ยังไม่ได้เลข %d' % blank if blank else ''), 'OK')

    def add_made_bundle(self, name, bid, n_items):
        """จดไว้ว่าสร้างบันเดิลอะไรไปแล้ว ได้เลขอะไร — โชว์ในแท็บผลลัพธ์"""
        row = {'name': name, 'id': bid or '-', 'items': n_items,
               'at': datetime.now().strftime('%H:%M:%S')}
        self.made.append(row)
        row['no'] = len(self.made)

        def _do():
            self.tree_made.insert('', 'end', values=(
                row['no'], row['id'], row['name'], row['items'], row['at']))
            self.tree_made.yview_moveto(1)
            got = sum(1 for m in self.made if m['id'] != '-')
            self.lbl_made.config(
                text='สร้างแล้ว %d บันเดิล%s' % (
                    len(self.made),
                    '' if got == len(self.made) else ' (ได้เลข %d)' % got),
                fg=C['ok'])
        self.root.after(0, _do)

    def copy_made(self):
        """คัดลอกเฉพาะ "เลข" Bundle ของแถวที่เลือกไว้ — ไม่เลือกก็เอาทั้งหมด"""
        rows = list(self.tree_made.selection())
        picked = bool(rows)
        if not picked:
            rows = list(self.tree_made.get_children())
        if not rows:
            messagebox.showinfo('ยังไม่มี', 'ยังไม่ได้สร้างบันเดิลในรอบนี้')
            return
        ids, blank = [], 0
        for w in rows:
            v = self.tree_made.item(w, 'values')
            bid = str(v[1]).strip() if len(v) > 1 else ''
            if bid and bid != '-':
                ids.append(bid)
            else:
                blank += 1
        if not ids:
            messagebox.showinfo('ไม่มีเลข',
                                'แถวที่เลือกยังไม่ได้เลข Bundle (ขึ้น - อยู่)')
            return
        self.root.clipboard_clear()
        self.root.clipboard_append('\n'.join(ids))
        self.log('คัดลอกเลข Bundle %d รายการแล้ว (%s)%s' % (
            len(ids), 'เฉพาะที่เลือก' if picked else 'ทั้งหมด',
            '  ข้ามที่ยังไม่ได้เลข %d' % blank if blank else ''), 'OK')

    def _ins_result(self, r):
        tag = ('miss',) if r.get('miss') else (('dup',) if r.get('dup') else ())
        self.tree.insert('', 'end', tags=tag,
                         values=(r.get('seq') or '', r.get('id') or '—', r['name'],
                                 r['type'], r['kind'], r['notes']))

    def finalize_results(self, criteria):
        """จัดผลลัพธ์ให้อ่านรู้เรื่อง — ทำหลังค้นเสร็จ
        1. เรียงตามลำดับที่สั่งค้น (ไม่ใช่ตามลำดับที่เว็บคืนมา)
        2. ลำดับไหนได้หลายตัว ทำเครื่องหมายว่าซ้ำ
        3. ลำดับไหนไม่เจอ ใส่แถวบอกไว้เลย จะได้เห็นครบว่าอะไรเป็นอะไร
        """
        cnt = {}
        for r in self.results:
            cnt[r.get('seq')] = cnt.get(r.get('seq'), 0) + 1
        for r in self.results:
            r['dup'] = cnt.get(r.get('seq'), 0) > 1

        why = {m['seq']: m['why'] for m in self.not_found if m.get('seq')}
        got = {r.get('seq') for r in self.results}
        for i, c in enumerate(criteria, 1):
            if i in got:
                continue
            self.results.append({
                'id': '', 'name': c.get('disp') or c.get('name') or '',
                'type': '', 'kind': c.get('kind') or '',
                'notes': '✗ ' + why.get(i, 'ไม่เจอ'),
                'seq': i, 'req': c.get('kind') or '', 'miss': True, 'dup': False,
                # เก็บแถวต้นทางไว้ด้วย เผื่อสั่งส่งเข้าคิวสร้าง Item ต่อ
                'src': dict(c)})

        self.results.sort(key=lambda r: (r.get('seq') or 0, r.get('id') or ''))
        n_ok = sum(1 for r in self.results if not r.get('miss'))
        n_miss = sum(1 for r in self.results if r.get('miss'))
        n_dup = sum(1 for r in self.results if r.get('dup'))

        def _do():
            self.tree.delete(*self.tree.get_children())
            for r in self.results:
                self._ins_result(r)
            txt = f'สั่งค้น {len(criteria)} · เจอ {n_ok}'
            if n_dup:
                txt += f' · ซ้ำ {n_dup}'
            if n_miss:
                txt += f' · ไม่เจอ {n_miss}'
            self.lbl_count.config(text=txt)
        self.root.after(0, _do)

    def clear_results(self):
        self.results.clear()
        self.not_found.clear()
        self.tree.delete(*self.tree.get_children())
        self.lbl_count.config(text='พบ 0 รายการ')

    def copy_ids(self):
        if not self.results:
            return messagebox.showinfo('ผลลัพธ์', 'ยังไม่มีผลลัพธ์')
        ids = ', '.join(r['id'] for r in self.results if r.get('id'))
        self.root.clipboard_clear()
        self.root.clipboard_append(ids)
        self.log(f'คัดลอก {len(self.results)} ID แล้ว', 'OK')

    def picked_results(self):
        """แถวที่เลือกไว้ในตารางผลค้นหา — ไม่ได้เลือกก็คืนทั้งหมด"""
        sel = list(self.tree.selection())
        if not sel:
            return list(self.results), False
        out = []
        for iid in sel:
            i = self.tree.index(iid)
            if 0 <= i < len(self.results):
                out.append(self.results[i])
        return out, True

    def miss_to_queue(self):
        """ตัวที่ค้นแล้วไม่เจอ -> ส่งเข้าคิวสร้าง Item เลย ไม่ต้องนำเข้าไฟล์ใหม่"""
        if self.c_running:
            return messagebox.showinfo('กำลังทำงาน', 'ตอนนี้กำลังสร้างไอเทมอยู่ รอให้เสร็จก่อนนะ')
        if not self.results:
            return messagebox.showinfo('ยังไม่มีผลลัพธ์', 'ค้นหาก่อนแล้วค่อยส่งเข้าคิวนะ')
        rows, picked = self.picked_results()
        miss = [r for r in rows if r.get('miss')]
        if not miss:
            return messagebox.showinfo(
                'ไม่มีตัวที่ไม่เจอ',
                'แถวที่เลือกไว้เจอบนเว็บหมดแล้ว — ไม่ต้องสร้างใหม่'
                if picked else 'รอบนี้เจอครบทุกตัว ไม่มีอะไรต้องสร้างใหม่')

        t = self.cv_type.get().strip() or DEFAULT_TYPE
        sfx = self.cv_suffix.get().strip()
        pr = self.cv_price.get().strip()
        ml = self.cv_mail.get().strip()
        have = {str(d.get('kind') or '').strip() for d in self.cq}
        add, dup = [], []
        for r in miss:
            it = miss_to_item(r, t, sfx, pr, ml)
            k = str(it.get('kind') or '').strip()
            if k and k in have:
                dup.append(k)
                continue
            if k:
                have.add(k)
            add.append(it)

        if not add:
            return messagebox.showinfo(
                'อยู่ในคิวอยู่แล้ว',
                'ตัวที่ไม่เจอ %d รายการ อยู่ในคิวสร้าง Item อยู่แล้วทั้งหมด' % len(miss))

        msg = ('ส่งเข้าคิวสร้าง Item %d รายการ\n(%s)\n\n'
               % (len(add), 'เฉพาะที่เลือก' if picked else 'แถวแดงทั้งหมด'))
        msg += '\n'.join('· %s  (ItemKind %s)' % (d['name'][:44], d['kind'])
                         for d in add[:8])
        if len(add) > 8:
            msg += '\n· ... อีก %d รายการ' % (len(add) - 8)
        if dup:
            msg += '\n\n(ข้าม %d รายการที่อยู่ในคิวแล้ว)' % len(dup)
        msg += '\n\nชื่อกับค่าต่างๆ ใช้กฎเดียวกับตอนนำเข้าไฟล์ — แก้เพิ่มทีหลังได้'
        if not messagebox.askyesno('ส่งเข้าคิว', msg):
            return

        self.cq.extend(add)
        self._c_refresh()
        self.nb.select(self.tab_create)
        self.log('ส่งตัวที่ไม่เจอเข้าคิวสร้าง Item %d รายการ (รวมในคิว %d)%s'
                 % (len(add), len(self.cq),
                    '  ข้ามที่อยู่ในคิวแล้ว %d' % len(dup) if dup else ''), 'OK')
        log_event('miss_to_queue', count=len(add), skipped=len(dup),
                  total=len(self.cq), picked=bool(picked))

    def export(self, kind):
        if not self.results:
            return messagebox.showinfo('ผลลัพธ์', 'ยังไม่มีผลลัพธ์')
        ext = '.xlsx' if kind == 'xlsx' else '.csv'
        path = filedialog.asksaveasfilename(
            defaultextension=ext, filetypes=[(kind.upper(), '*' + ext)],
            initialfile=f'TR_Items_{datetime.now():%Y%m%d_%H%M}{ext}')
        if not path:
            return
        header = ['#', 'Aztek Item Id', 'ชื่อ', 'ประเภท', 'ItemKind', 'สถานะ', 'หมายเหตุ']
        rows = [[r.get('seq') or '', r.get('id') or '', r['name'], r['type'], r['kind'],
                 ('ไม่เจอ' if r.get('miss') else ('ซ้ำ - เลือกเอง' if r.get('dup') else 'ok')),
                 r['notes']] for r in self.results]
        if kind == 'xlsx' and XLSX_OK:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Results'
            ws.append(header)
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', fgColor='1F6FEB')
            for r in rows:
                ws.append(r)
            for i, w in enumerate([6, 14, 46, 12, 12, 14, 26], 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
            wb.save(path)
        else:
            with open(path, 'w', newline='', encoding='utf-8-sig') as f:
                w = csv.writer(f)
                w.writerow(header)
                w.writerows(rows)
        self.log(f'บันทึกไฟล์แล้ว: {path}', 'OK')
        log_event('export', format=kind, rows=len(rows), file=os.path.basename(path))

    def download_template(self):
        path = filedialog.asksaveasfilename(
            defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx'), ('CSV', '*.csv')],
            initialfile='TR_Item_Template.xlsx')
        if not path:
            return
        try:
            build_template(path)
            self.log(f'สร้าง template แล้ว: {path}', 'OK')
        except Exception as ex:
            messagebox.showerror('Template', str(ex))

    # ---------- รายการที่จะค้นหา ----------
    def toggle_imp(self):
        self.imp_open = not self.imp_open
        if self.imp_open:
            # กางแล้วต้องเห็นจริง — ถ้าหน้าต่างเตี้ยไป ขยายให้เอง
            try:
                self.root.update_idletasks()
                want = min(950, self.root.winfo_screenheight() - 80)
                if self.root.winfo_height() < want:
                    self.root.geometry('%dx%d' % (max(1080, self.root.winfo_width()), want))
            except Exception:
                pass
            self.imp_box.pack(fill='both', expand=True, padx=14, pady=(12, 0),
                              before=self.s_tail)
            self.btn_imp.config(text='▾  ซ่อนรายการที่จะค้นหา')
        else:
            self.imp_box.pack_forget()
            self.btn_imp.config(text='▸  ดูรายการที่จะค้นหา')

    def _imp_why(self):
        s = self.imp_tree.selection()
        if not s:
            return
        i = self.imp_tree.index(s[0])
        w = self.imported[i].get('warn') if 0 <= i < len(self.imported) else None
        self.imp_why.config(text=('⚠  แถว #%d : %s' % (i + 1, ' · '.join(w))) if w else '')

    def refresh_imp(self):
        """วาดรายการที่จะค้นหาใหม่ — เรียกทุกครั้งที่ self.imported เปลี่ยน"""
        nwarn = mark_suspects(self.imported)
        self.imp_tree.delete(*self.imp_tree.get_children())
        for i, r in enumerate(self.imported, 1):
            dur = r.get('dur', '')
            dur = 'ถาวร' if dur == '' else ('— ไม่กรอง —' if dur == 'any' else str(dur) + ' วัน')
            trade = {'yes': 'ได้', 'no': 'ไม่ได้', 'any': '— ไม่กรอง —'}.get(
                r.get('trade', 'any'), r.get('trade', ''))
            name = r.get('disp') or r.get('name') or ''
            self.imp_tree.insert('', 'end', tags=('warn',) if r.get('warn') else (),
                                 values=(i, '⚠' if r.get('warn') else '',
                                         r.get('kind', ''), name, dur, trade,
                                         r.get('qty') or '—'))
        n = len(self.imported)
        self.imp_cnt.config(text=f'{n} รายการ' + (f'  ·  น่าสงสัย {nwarn}' if nwarn else ''))
        self.btn_imp.config(state='normal' if n else 'disabled')
        self.btn_multi.config(state='normal' if n else 'disabled')
        self.lbl_imp.config(text=(f'⚠  มี {nwarn} รายการหน้าตาไม่เหมือนไอเทมทั่วไป — กางดูได้'
                                  if nwarn else ''))
        self.imp_why.config(text='')

    def imp_del(self):
        sel = sorted((self.imp_tree.index(s) for s in self.imp_tree.selection()), reverse=True)
        if not sel:
            return messagebox.showinfo('รายการ', 'เลือกแถวที่จะเอาออกก่อนนะ')
        for i in sel:
            if 0 <= i < len(self.imported):
                del self.imported[i]
        self.refresh_imp()
        self.lbl_file.config(text=f'{self.imp_src} — {len(self.imported)} รายการ')
        self.log(f'เอาออกจากรายการ {len(sel)} ตัว (เหลือ {len(self.imported)})', 'INFO')

    def imp_copy(self):
        if not self.imported:
            return
        ids = ', '.join(str(r.get('kind', '')) for r in self.imported if r.get('kind'))
        self.root.clipboard_clear()
        self.root.clipboard_append(ids)
        self.log(f'คัดลอก {len(self.imported)} เลขไอเทมแล้ว', 'OK')

    def pick_file(self):
        path = filedialog.askopenfilename(filetypes=[('Excel / CSV', '*.xlsx *.xlsm *.csv')])
        if not path:
            return
        try:
            if path.lower().endswith('.csv'):
                rows = read_template(path)
                src = os.path.basename(path)
            else:
                if not XLSX_OK:
                    raise RuntimeError('เครื่องนี้อ่านไฟล์ Excel ไม่ได้ ลองใช้ .csv แทน')
                dlg = ImportDialog(self.root, path)
                if dlg.result is None:
                    return
                rows = dlg.result
                names = getattr(dlg, 'sheet_names', None) or (
                    [dlg.sheet_name] if dlg.sheet_name else [])
                if len(names) > 3:
                    where = '%d ชีท' % len(names)
                else:
                    where = ', '.join(names) or '-'
                src = f'{os.path.basename(path)}  ›  {where}'
            if not rows:
                messagebox.showwarning('นำเข้า', 'ไม่พบรายการที่ใช้ได้ในไฟล์นี้')
                return
            self.imported = rows
            self.imp_src = src
            self.lbl_file.config(text=f'{src} — {len(rows)} รายการ')
            self.refresh_imp()
            self.log(f'นำเข้าจาก {src}: {len(rows)} รายการ', 'OK')
            log_event('import', file=os.path.basename(path), source=src, rows=len(rows))
        except Exception as ex:
            log_event('error', where='import', message=str(ex)[:300],
                      file=os.path.basename(path))
            messagebox.showerror('อ่านไฟล์ไม่ได้', str(ex))

    def save_now(self):
        self.prefs.update({
            'headless': self.v_headless.get(),
            'exact': self.v_exact.get(),
            'c_type': self.cv_type.get().strip(),
            'c_suffix': self.cv_suffix.get().strip(),
            'c_price': self.cv_price.get().strip(),
            'c_mail': self.cv_mail.get().strip(),
            'c_hold': self.cv_hold.get().strip(),
            'w_hold': self.wv_hold.get().strip(),
        })
        # ค่าเก่าของแผง Deep Check ที่เอาออกไปแล้ว — ล้างทิ้ง ไม่ให้ย้อนมาหลอนทีหลัง
        for dead in ('deep', 'dur', 'trade', 'qty'):
            self.prefs.pop(dead, None)
        save_prefs(self.prefs)

    def deep_criteria(self):
        if not self.v_deep.get():
            return {'dur': 'any', 'trade': 'any', 'qty': ''}
        return {'dur': dur_from_text(self.v_dur.get()), 'trade': self.v_trade.get(),
                'qty': self.v_qty.get().strip()}

    def do_cancel(self):
        self.cancel = True
        self.log('กำลังยกเลิก...', 'WARN')

    # ---------- login ----------
    def open_login(self):
        if self.running:
            return messagebox.showinfo('กำลังทำงาน', 'รอให้งานปัจจุบันเสร็จก่อนนะ')
        log_event('login_open')
        threading.Thread(target=lambda: asyncio.run(self._login()), daemon=True).start()

    async def _login(self):
        self.log('เปิด Chrome ให้ล็อกอิน — ล็อกอินเสร็จแล้วปิดหน้าต่างได้เลย', 'STEP')
        try:
            async with async_playwright() as pw:
                b = await pw.chromium.launch_persistent_context(**launch_kwargs(False))
                page = b.pages[0] if b.pages else await b.new_page()
                await page.goto(ITEM_LIST_URL, wait_until='domcontentloaded', timeout=60000)
                while True:
                    await asyncio.sleep(1)
                    if not b.pages:
                        break
        except Exception as ex:
            self.log('ปิดหน้าต่าง login แล้ว (' + str(ex)[:80] + ')', 'INFO')
        self.log('พร้อมใช้งาน', 'OK')

    # ---------- run ----------
    def run_single(self):
        kind = self.v_kind.get().strip()
        name = self.v_name.get().strip()
        if not kind and not name:
            return messagebox.showwarning('ค้นหา', 'กรอก ItemKind หรือชื่อไอเทมอย่างน้อยหนึ่งช่อง')
        c = {'kind': kind, 'name': name}
        c.update(self.deep_criteria())
        self.start([c])

    def run_multi(self):
        if not self.imported:
            return
        n_deep = sum(1 for c in self.imported if has_deep(c))
        if n_deep:
            self.log('ไฟล์มีเงื่อนไข (ระยะเวลา/แลกเปลี่ยน/จำนวน) %d รายการ → '
                     'เข้าไปดูค่าจริงในหน้ารายละเอียดให้อัตโนมัติ' % n_deep, 'INFO')
        self.start([dict(c) for c in self.imported])

    def start(self, criteria):
        if self.running:
            return
        self.save_now()
        self.running = True
        self.cancel = False
        self.clear_results()
        self.btn_single.config(state='disabled')
        self.btn_multi.config(state='disabled')
        self.btn_cancel.config(state='normal')
        self.nb.select(self.tab_log)
        self.log('=' * 46, 'STEP')
        self.log(f'เริ่มค้นหา {len(criteria)} รายการ', 'STEP')
        self._run_started = datetime.now()
        d = self.deep_criteria()
        log_event('search_start', count=len(criteria), deep=bool(self.v_deep.get()),
                  dur=d['dur'], trade=d['trade'], qty=d['qty'],
                  kinds=[c.get('kind') for c in criteria][:50])
        threading.Thread(target=self._thread, args=(criteria,), daemon=True).start()

    def _thread(self, criteria):
        try:
            asyncio.run(self._work(criteria))
        except Exception as ex:
            log_event('error', where='search', message=str(ex)[:300])
            self.log('ผิดพลาด: ' + str(ex), 'ERR')
            self.log(traceback.format_exc(), 'ERR')
        finally:
            self.running = False

            def _rst():
                self.btn_single.config(state='normal')
                self.btn_multi.config(state='normal' if self.imported else 'disabled')
                self.btn_cancel.config(state='disabled')
                if self.results:
                    self.nb.select(self.tab_search)
            self.root.after(0, _rst)

    # ---------- automation ----------
    async def _work(self, criteria):
        async with async_playwright() as pw:
            if not find_chrome_exe():
                self.log('ไม่เจอ Chrome ในเครื่อง — ใช้ Chromium ของ Playwright แทน', 'WARN')
            browser = await pw.chromium.launch_persistent_context(
                **launch_kwargs(self.v_headless.get()))
            page = browser.pages[0] if browser.pages else await browser.new_page()
            try:
                await self._search_all(page, criteria)
            finally:
                try:
                    await browser.close()
                except Exception:
                    pass

    async def _goto_list(self, page):
        if not page.url.rstrip('/').endswith('/shop/items'):
            await page.goto(ITEM_LIST_URL, wait_until='domcontentloaded', timeout=45000)
            await page.wait_for_timeout(1500)

    async def _apply_search(self, page, keyword, dur_filter=None):
        await page.wait_for_selector(SEL['search_input'], timeout=20000)
        await page.fill(SEL['search_input'], keyword or '')
        f = page.locator(SEL['filter_duration'])
        if await f.count() > 0:
            await f.first.fill('' if dur_filter is None else str(dur_filter))
        await page.wait_for_timeout(250)
        await page.locator(SEL['search_button']).first.click()
        await page.wait_for_timeout(1400)

    async def _read_all_pages(self, page):
        all_rows, seen, pnum = [], set(), 1
        total = await page.evaluate(JS_TOTAL_PAGES)
        while True:
            if self.cancel:
                break
            rows = await page.evaluate(JS_READ_ROWS, COL)
            fresh = [r for r in rows if r['id'] not in seen]
            for r in fresh:
                seen.add(r['id'])
            all_rows.extend(fresh)
            self.log(f'  หน้า {pnum}/{total}: {len(rows)} แถว (ใหม่ {len(fresh)})')
            if pnum >= total or pnum > 200:
                break
            first = rows[0]['id'] if rows else None
            nxt = page.locator(f'button:has-text("{SEL["next_text"]}")')
            if await nxt.count() == 0 or await nxt.first.is_disabled():
                break
            await nxt.first.click()
            pnum += 1
            await page.wait_for_timeout(900)
            try:
                await page.wait_for_function(
                    """(prev) => {
                        const td = document.querySelector('table tbody tr td');
                        return td && (td.innerText || '').trim() !== prev;
                    }""", arg=first, timeout=8000)
            except Exception:
                break
        return all_rows

    async def _open_detail(self, page, item_id):
        ok = await page.evaluate(JS_CLICK_ROW, str(item_id))
        if not ok:
            await self._apply_search(page, str(item_id))
            ok = await page.evaluate(JS_CLICK_ROW, str(item_id))
            if not ok:
                raise RuntimeError(f'หาแถวของ ID {item_id} ไม่เจอ')
        await page.wait_for_selector(SEL['d_name'], timeout=20000)
        await page.wait_for_timeout(400)

    async def _read_detail(self, page):
        return await page.evaluate(JS_READ_DETAIL, {
            'name': SEL['d_name'], 'kind': SEL['d_kind'], 'price': SEL['d_price'],
            'duration': SEL['d_duration'], 'qty': SEL['d_qty'],
            'tradeLabel': SEL['d_trade_label'],
        })

    async def _search_all(self, page, criteria):
        await self._goto_list(page)
        for i, c in enumerate(criteria):
            if self.cancel:
                break
            shown = c.get('disp') or c.get('name') or ''
            label = f"#{i + 1} Kind={c.get('kind') or '-'}" + (f"  {shown}" if shown else '')
            self.log(f'── {label}', 'STEP')

            use_filter = bool(duration_only_numeric(c))
            await self._goto_list(page)
            await self._apply_search(page, c.get('kind') or c.get('name'),
                                     c['dur'] if use_filter else None)
            if use_filter:
                self.log(f"  ใช้ฟิลเตอร์ DurationIndex={c['dur']} (ข้าม deep check)")

            rows = await self._read_all_pages(page)

            # ช่องค้นหาของเว็บเป็นการค้นแบบ "มีคำนี้อยู่" เลยติดตัวที่ไม่ได้สั่งหามาด้วย
            # ถ้าสั่งหาด้วย ItemKind ก็เอาเฉพาะแถวที่ ItemKind ตรงเป๊ะ
            want_kind = str(c.get('kind') or '').strip()
            if want_kind and self.v_exact.get():
                before = len(rows)
                rows = [r for r in rows if str(r.get('kind', '')).strip() == want_kind]
                if before != len(rows):
                    self.log(f'  ตัดตัวที่ ItemKind ไม่ตรง {want_kind} ออก: '
                             f'{before} → {len(rows)}', 'INFO')

            if c.get('name'):
                before = len(rows)
                want = norm_name(c['name'])
                rows = [r for r in rows if want in norm_name(r['name'])]
                self.log(f"  กรองชื่อ {c['name']!r}: {before} → {len(rows)}")
            self.log(f'  พบ {len(rows)} รายการ', 'INFO' if rows else 'WARN')
            if len(rows) > 1:
                self.log('  ! เว็บมี %d ตัวที่ตรงเงื่อนไขนี้ (Aztek Id: %s) — '
                         'จะขึ้นสีเหลืองไว้ให้เลือกเอง'
                         % (len(rows), ', '.join(r['id'] for r in rows)), 'WARN')
            if not rows:
                self.not_found.append({'seq': i + 1, 'label': label,
                                       'why': 'ไม่พบแถวที่ตรงเงื่อนไข'})
                continue

            if not (has_deep(c) and not use_filter):
                for r in rows:
                    self.add_result(r, f"{c['dur']} วัน" if use_filter else '',
                                    seq=i + 1, req=want_kind)
                log_event('search_item', kind=c.get('kind'), name=c.get('disp') or c.get('name'),
                          found=len(rows), passed=len(rows), deep=False,
                          ids=[r['id'] for r in rows][:20])
                self.set_progress(i + 1, len(criteria), 'เสร็จ')
                continue

            self.log(f'  Deep check {len(rows)} รายการ...', 'STEP')
            passed = 0
            for k, r in enumerate(rows):
                if self.cancel:
                    break
                self.set_progress(k + 1, len(rows), f"deep: {r['id']}")
                try:
                    await self._open_detail(page, r['id'])
                    d = await self._read_detail(page)
                    ok, why = match_deep(d, c)
                    if ok:
                        self.add_result(r, why, seq=i + 1, req=want_kind)
                        passed += 1
                        self.log(f"  ✓ {r['id']}  {r['name']}  [{why}]", 'OK')
                    else:
                        self.log(f"  ✗ {r['id']}  {r['name']}  ({why})")
                except Exception as ex:
                    self.log(f"  ! {r['id']} ตรวจไม่ได้: {ex}", 'WARN')
                await page.go_back(wait_until='domcontentloaded')
                await page.wait_for_timeout(500)
            log_event('search_item', kind=c.get('kind'), name=c.get('disp') or c.get('name'),
                      found=len(rows), passed=passed, deep=True,
                      want_dur=c.get('dur'), want_trade=c.get('trade'), want_qty=c.get('qty'),
                      ids=[r['id'] for r in self.results][-passed:][:20] if passed else [])
            if not passed and not self.cancel:
                self.not_found.append({'seq': i + 1, 'label': label,
                                       'why': f'เจอ {len(rows)} แต่ไม่ผ่าน deep check'})
            self.set_progress(i + 1, len(criteria), 'เสร็จ')

        self.finalize_results(criteria)
        self.log('=' * 46, 'STEP')
        if self.cancel:
            self.log('ยกเลิกโดยผู้ใช้', 'WARN')
        found = [r for r in self.results if not r.get('miss')]
        dups = sorted({r['seq'] for r in self.results if r.get('dup')})
        self.log('สั่งค้น %d · เจอ %d · ไม่เจอ %d'
                 % (len(criteria), len(found), len(self.results) - len(found)),
                 'OK' if found else 'WARN')
        if found:
            self.log('IDs: ' + ', '.join(r['id'] for r in found), 'OK')
        if dups:
            self.log('ลำดับที่เจอมากกว่า 1 ตัว (แถวเหลือง เลือกเอง): '
                     + ', '.join('#%d' % d for d in dups), 'WARN')
        for m in self.not_found:
            self.log('  ✗ %s  (%s)' % (m['label'], m['why']), 'WARN')
        secs = 0
        try:
            secs = round((datetime.now() - self._run_started).total_seconds(), 1)
        except Exception:
            pass
        log_event('search_done', criteria=len(criteria), found=len(found),
                  not_found=len(self.not_found), dups=len(dups), seconds=secs,
                  cancelled=bool(self.cancel), ids=[r['id'] for r in found][:200])

    def run(self):
        self.root.mainloop()


if __name__ == '__main__':
    App().run()
